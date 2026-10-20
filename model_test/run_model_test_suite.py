"""Unified runner for the five model tests used by this project.

The runner keeps the existing Excel/result conventions while making the tested
model and the reviewing model explicit configuration values.  It intentionally
does not execute requested tools; Function Call tests only record the model's
requested tool name and arguments.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import time
from collections import OrderedDict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
MODEL_TEST_ROOT = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import config11
from model_test import functiontestformat as function_call_impl
from model_test import run_dual_chat_excel as dialog_impl
from model_test import run_model_timeweast_excel as latency_impl
from model_test import run_qwen_score_excel as score_impl
from model_test import suite_config


TEST_NAMES = ("speed", "dialog", "function_call", "latency", "score")
TEST_DESCRIPTIONS = {
    "speed": "5 non-stream and 5 stream compatibility/speed calls",
    "dialog": "3 sessions x 100 alternating DeepSeek/Doubao dialogue rounds",
    "function_call": "Function Call cases with two attempts and highlighted failures",
    "latency": "10 streamed calls with TTFT/total/token statistics",
    "score": "Three-dimension Qwen review of a generated dialogue workbook",
}


@dataclass
class SuiteContext:
    test_model: str
    review_model: str
    output_dir: Path
    work_dir: Path
    print_live: bool
    dry_run: bool
    results: list[dict[str, Any]] = field(default_factory=list)
    manifest_path: Path | None = None

    def record(
        self,
        name: str,
        status: str,
        *,
        output: Path | None = None,
        error: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> None:
        item: dict[str, Any] = {"name": name, "status": status}
        if output is not None:
            item["output"] = str(output)
        if error:
            item["error"] = error
        if metadata:
            item["metadata"] = metadata
        self.results.append(item)
        self.write_manifest()

    def write_manifest(self) -> None:
        if self.manifest_path is None:
            return
        payload = {
            "schema_version": 1,
            "updated_at": datetime.now().isoformat(timespec="seconds"),
            "test_model": self.test_model,
            "review_model": self.review_model,
            "results": self.results,
        }
        self.manifest_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the unified model test suite.")
    parser.add_argument(
        "--tests",
        nargs="+",
        choices=("all", *TEST_NAMES),
        default=["all"],
        help="Tests to run. Default: all.",
    )
    parser.add_argument(
        "--test-model",
        default=suite_config.TEST_MODEL,
        help="Model under test. Default comes from suite_config.py.",
    )
    parser.add_argument(
        "--review-model",
        default=suite_config.REVIEW_MODEL,
        help="Model used by the scoring stage. Default comes from suite_config.py.",
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=None,
        help="Existing dialogue Excel for standalone score testing.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory. Default: model_test/outputs/model_suite/<timestamp>.",
    )
    parser.add_argument(
        "--print-live",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Print every generated dialogue line. Default: false.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate configs and local documents without making API calls.",
    )
    parser.add_argument(
        "--list-tests",
        action="store_true",
        help="List available tests and exit.",
    )
    return parser.parse_args()


def selected_tests(raw: list[str]) -> list[str]:
    if "all" in raw:
        return list(TEST_NAMES)
    result: list[str] = []
    for name in raw:
        normalized = "function_call" if name == "function-call" else name
        if normalized not in result:
            result.append(normalized)
    return result


def default_output_dir() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return MODEL_TEST_ROOT / "outputs" / "model_suite" / stamp


def error_text(exc: Exception) -> str:
    return f"{type(exc).__name__}: {exc}"


def configure_model_network() -> None:
    """Bypass a stale local proxy only for configured model-service hosts."""
    if not suite_config.BYPASS_ENV_PROXY_FOR_MODEL_HOSTS:
        return

    urls = [
        suite_config.TEST_BASE_URL,
        suite_config.REVIEW_BASE_URL,
        suite_config.DIALOG_USER_BASE_URL,
    ]
    hosts = {urlparse(url).hostname for url in urls}
    hosts.discard(None)
    existing = {
        item.strip().lower()
        for item in os.environ.get("NO_PROXY", "").split(",")
        if item.strip()
    }
    added_hosts = sorted(str(host).lower() for host in hosts if str(host).lower() not in existing)
    if not added_hosts:
        return

    merged = ",".join([os.environ.get("NO_PROXY", "").strip(), *added_hosts]).strip(",")
    os.environ["NO_PROXY"] = merged
    os.environ["no_proxy"] = merged
    print(f"[INFO] bypassing environment proxy for: {', '.join(added_hosts)}")


def build_test_client() -> OpenAI:
    return OpenAI(
        api_key=suite_config.TEST_API_KEY,
        base_url=suite_config.TEST_BASE_URL,
        timeout=suite_config.REQUEST_TIMEOUT_SEC,
        max_retries=0,
    )


def build_review_client() -> OpenAI:
    return OpenAI(
        api_key=suite_config.REVIEW_API_KEY,
        base_url=suite_config.REVIEW_BASE_URL,
        timeout=float(getattr(score_impl.score_config, "REQUEST_TIMEOUT_SEC", 120.0)),
        max_retries=0,
    )


def validate_model(model: str) -> None:
    if not model.strip():
        raise ValueError("The test model ID cannot be empty.")
    available = list(getattr(config11, "DOUBAO_AVAILABLE_MODELS", []))
    if available and model not in available:
        print(
            f"[WARN] test model {model!r} is not listed in config11.DOUBAO_AVAILABLE_MODELS; continuing.",
            file=sys.stderr,
        )


def percentile(sorted_values: list[float], ratio: float) -> float | None:
    if not sorted_values:
        return None
    index = int(round((len(sorted_values) - 1) * ratio))
    return sorted_values[max(0, min(index, len(sorted_values) - 1))]


def safe_windows_filename(value: str) -> str:
    """Return a non-empty filename stem that is valid on Windows."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(". ")
    return cleaned or "model"


def total_workbook_path(ctx: SuiteContext) -> Path:
    return ctx.output_dir / f"{safe_windows_filename(ctx.test_model)}_总表.xlsx"


def safe_sheet_title(workbook: Workbook, desired: str) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', "_", desired).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate in workbook.sheetnames:
        tail = f"_{suffix}"
        candidate = f"{cleaned[:31 - len(tail)]}{tail}"
        suffix += 1
    return candidate


def copy_worksheet_into(workbook: Workbook, source_ws: Any, desired_title: str) -> Any:
    """Copy values and visible worksheet formatting into a new workbook sheet."""
    target_ws = workbook.create_sheet(safe_sheet_title(workbook, desired_title))
    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = target_ws.cell(source_cell.row, source_cell.column, source_cell.value)
            # A StyleArray stores workbook-local IDs, so assigning _style
            # across workbooks corrupts the target style table. Copy the
            # concrete style components instead.
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.protection = copy(source_cell.protection)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for column_key, source_dimension in source_ws.column_dimensions.items():
        target_dimension = target_ws.column_dimensions[column_key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
        target_dimension.outlineLevel = source_dimension.outlineLevel
    for row_key, source_dimension in source_ws.row_dimensions.items():
        target_dimension = target_ws.row_dimensions[row_key]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.auto_filter.ref = source_ws.auto_filter.ref
    return target_ws


def append_table(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 60)
        ws.column_dimensions[column_cells[0].column_letter].width = max(width, 12)


def result_output(ctx: SuiteContext, name: str) -> Path | None:
    for item in ctx.results:
        if item["name"] == name and item["status"] == "success" and item.get("output"):
            return Path(item["output"])
    return None


def read_latency_metrics(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "summary" not in workbook.sheetnames:
            return "未找到延迟汇总"
        ws = workbook["summary"]
        headers = [cell.value for cell in ws[1]]
        values = [cell.value for cell in ws[2]] if ws.max_row >= 2 else []
        row = dict(zip(headers, values))
        return (
            f"成功 {row.get('成功次数', 0)}/{row.get('总运行次数', 0)}；"
            f"TTFT均值 {row.get('首包均值(ms)', '')} ms；"
            f"总耗时均值 {row.get('总耗时均值(ms)', '')} ms"
        )
    finally:
        workbook.close()


def read_function_call_metrics(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            return "未找到 Function Call 结果"
        ws = workbook.worksheets[0]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        cases = [row for row in rows if row and row[0] and row[0] != function_call_impl.SUMMARY_LABEL]
        summary = next((row for row in rows if row and row[0] == function_call_impl.SUMMARY_LABEL), None)
        red_attempts = summary[1] if summary and len(summary) > 1 else ""
        return f"用例 {len(cases)} 条；标红尝试 {red_attempts} 次"
    finally:
        workbook.close()


def read_dialog_metrics(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            return "未找到长对话结果"
        ws = workbook.worksheets[0]
        session_ids = {row[0] for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None}
        return f"会话 {len(session_ids)} 个；消息 {max(ws.max_row - 1, 0)} 条"
    finally:
        workbook.close()


def read_score_metrics(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "summary" not in workbook.sheetnames:
            return "未找到评分汇总"
        ws = workbook["summary"]
        headers = [cell.value for cell in ws[1]]
        scores: list[str] = []
        for index, header in enumerate(headers):
            if not isinstance(header, str) or not header.endswith("_final_score"):
                continue
            values = [row[index] for row in ws.iter_rows(min_row=2, values_only=True) if isinstance(row[index], (int, float))]
            if values:
                dimension = header[: -len("_final_score")]
                scores.append(f"{dimension} {round(sum(values) / len(values), 2)}")
        return "；".join(scores) if scores else "未找到评分"
    finally:
        workbook.close()


def build_total_workbook(
    ctx: SuiteContext,
    requested_tests: list[str],
    *,
    external_dialog_input: Path | None = None,
) -> Path:
    """Combine successful stage outputs into the one user-facing workbook."""
    output = total_workbook_path(ctx)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "测试概览"
    overview["A1"] = "统一模型测试总表"
    overview["A1"].font = Font(bold=True, size=16)
    overview["A3"] = "生成时间"
    # 总表中的生成时间只展示日期；运行清单仍保留精确时间用于审计。
    overview["B3"] = datetime.now().strftime("%Y-%m-%d")
    overview["A4"] = "被测模型"
    overview["B4"] = ctx.test_model
    overview["A5"] = "评审模型"
    overview["B5"] = ctx.review_model
    overview["A6"] = "本次请求测试项"
    overview["B6"] = ", ".join(requested_tests)
    if external_dialog_input is not None:
        overview["A7"] = "外部评分对话输入"
        overview["B7"] = str(external_dialog_input)

    outputs = {name: result_output(ctx, name) for name in TEST_NAMES}
    metrics: dict[str, str] = {}
    speed_path = outputs["speed"]
    if speed_path and speed_path.exists():
        speed_data = json.loads(speed_path.read_text(encoding="utf-8"))
        speed_rows = []
        for mode in ("non_stream", "stream"):
            summary = speed_data.get("summary", {}).get(mode, {})
            speed_rows.append(
                [
                    "非流式" if mode == "non_stream" else "流式",
                    summary.get("success"), summary.get("failed"),
                    summary.get("total_mean_ms"), summary.get("ttft_mean_ms"),
                    summary.get("total_p90_ms"), summary.get("ttft_p90_ms"),
                ]
            )
        speed_summary = workbook.create_sheet(safe_sheet_title(workbook, "速度汇总"))
        append_table(speed_summary, ["模式", "成功", "失败", "平均耗时(ms)", "平均TTFT(ms)", "P90耗时(ms)", "P90 TTFT(ms)"], speed_rows)
        speed_details = workbook.create_sheet(safe_sheet_title(workbook, "速度明细"))
        records = speed_data.get("records", [])
        headers = ["模式", "运行序号", "状态", "TTFT(ms)", "总耗时(ms)", "输出字符数", "错误"]
        detail_rows = [[record.get(key, "") for key in ("mode", "run_index", "status", "ttft_ms", "total_ms", "output_chars", "error")] for record in records]
        append_table(speed_details, headers, detail_rows)
        metrics["speed"] = "；".join(f"{row[0]} 成功 {row[1]}/{row[1] + row[2]}，平均 {row[3]} ms，TTFT {row[4]} ms" for row in speed_rows)

    dialog_path = outputs["dialog"] or external_dialog_input
    if dialog_path and dialog_path.exists():
        dialog_book = load_workbook(dialog_path, read_only=False, data_only=False)
        try:
            copy_worksheet_into(workbook, dialog_book.worksheets[0], "长对话")
        finally:
            dialog_book.close()
        metrics["dialog"] = read_dialog_metrics(dialog_path)

    function_path = outputs["function_call"]
    if function_path and function_path.exists():
        function_book = load_workbook(function_path, read_only=False, data_only=False)
        try:
            for source_ws in function_book.worksheets:
                copy_worksheet_into(workbook, source_ws, "函数调用")
        finally:
            function_book.close()
        metrics["function_call"] = read_function_call_metrics(function_path)

    latency_path = outputs["latency"]
    if latency_path and latency_path.exists():
        latency_book = load_workbook(latency_path, read_only=False, data_only=False)
        try:
            for source_name, target_name in (("summary", "延迟汇总"), ("raw", "延迟明细")):
                if source_name in latency_book.sheetnames:
                    copy_worksheet_into(workbook, latency_book[source_name], target_name)
        finally:
            latency_book.close()
        metrics["latency"] = read_latency_metrics(latency_path)

    score_path = outputs["score"]
    if score_path and score_path.exists():
        score_book = load_workbook(score_path, read_only=False, data_only=False)
        try:
            for source_ws in score_book.worksheets:
                target_name = "评分汇总" if source_ws.title == "summary" else f"评分-{source_ws.title}"
                copy_worksheet_into(workbook, source_ws, target_name)
        finally:
            score_book.close()
        metrics["score"] = read_score_metrics(score_path)

    overview.append([])
    append_table(overview, ["测试项", "状态", "关键结果", "说明"], [])
    status_by_name = {item["name"]: item for item in ctx.results}
    display_names = {
        "speed": "速度", "dialog": "长对话", "function_call": "Function Call", "latency": "延迟", "score": "评分",
    }
    for name in TEST_NAMES:
        item = status_by_name.get(name)
        if item is None:
            status = "未运行"
            detail = "本次未请求"
        elif item["status"] == "success":
            status = "成功"
            detail = "已汇总到本工作簿"
        else:
            status = "失败"
            detail = item.get("error", "")
        overview.append([display_names[name], status, metrics.get(name, ""), detail])
    for column in ("A", "B", "C", "D"):
        overview.column_dimensions[column].width = {"A": 18, "B": 12, "C": 90, "D": 70}[column]
    overview.freeze_panes = "A9"
    workbook.save(output)
    workbook.close()
    return output


def finalize_total_workbook(ctx: SuiteContext, requested_tests: list[str], external_dialog_input: Path | None) -> Path:
    """Create the final workbook, preserve an interrupted Function Call checkpoint, and clean temporary outputs."""
    total_output = build_total_workbook(ctx, requested_tests, external_dialog_input=external_dialog_input)
    checkpoint_name = "03_function_call_running.xlsx"
    temporary_checkpoint = ctx.work_dir / checkpoint_name
    persistent_checkpoint = ctx.output_dir / checkpoint_name
    if "function_call" in requested_tests:
        function_failed = any(
            item["name"] == "function_call" and item["status"] == "failed"
            for item in ctx.results
        )
        if function_failed and temporary_checkpoint.exists():
            shutil.copy2(temporary_checkpoint, persistent_checkpoint)
        elif persistent_checkpoint.exists():
            persistent_checkpoint.unlink()

    for item in ctx.results:
        if item["status"] == "success":
            item["output"] = str(total_output)
            item["metadata"] = {"combined_workbook": True}
    ctx.write_manifest()
    if ctx.work_dir.exists():
        shutil.rmtree(ctx.work_dir)
    return total_output


def run_speed(ctx: SuiteContext) -> Path:
    client = build_test_client()
    output = ctx.work_dir / "01_speed.json"
    records: list[dict[str, Any]] = []
    for mode in ("non_stream", "stream"):
        for run_index in range(1, suite_config.SPEED_RUNS + 1):
            started = time.perf_counter()
            first_output: float | None = None
            parts: list[str] = []
            try:
                request: dict[str, Any] = {
                    "model": ctx.test_model,
                    "messages": [
                        {"role": "system", "content": suite_config.SPEED_SYSTEM_PROMPT},
                        {"role": "user", "content": suite_config.SPEED_PROMPT},
                    ],
                    "temperature": suite_config.SPEED_TEMPERATURE,
                    "max_tokens": suite_config.SPEED_MAX_TOKENS,
                    "extra_body": {"thinking": {"type": "disabled"}},
                }
                if mode == "stream":
                    request["stream"] = True
                    response = client.chat.completions.create(**request)
                    for chunk in response:
                        choices = getattr(chunk, "choices", None) or []
                        if not choices:
                            continue
                        delta = getattr(choices[0], "delta", None)
                        if delta is None:
                            continue
                        content = getattr(delta, "content", None)
                        reasoning = getattr(delta, "reasoning_content", None)
                        tool_calls = getattr(delta, "tool_calls", None)
                        if first_output is None and (content or reasoning or tool_calls):
                            first_output = time.perf_counter()
                        if content:
                            parts.append(str(content))
                    if hasattr(response, "close"):
                        response.close()
                else:
                    response = client.chat.completions.create(**request)
                    content = getattr(response.choices[0].message, "content", None)
                    if content:
                        parts.append(str(content))
                finished = time.perf_counter()
                records.append(
                    {
                        "mode": mode,
                        "run_index": run_index,
                        "status": "success",
                        "ttft_ms": round((first_output - started) * 1000, 3)
                        if first_output is not None
                        else None,
                        "total_ms": round((finished - started) * 1000, 3),
                        "output_chars": len("".join(parts).strip()),
                    }
                )
            except Exception as exc:  # noqa: BLE001
                records.append(
                    {
                        "mode": mode,
                        "run_index": run_index,
                        "status": "failed",
                        "ttft_ms": None,
                        "total_ms": round((time.perf_counter() - started) * 1000, 3),
                        "output_chars": 0,
                        "error": error_text(exc),
                    }
                )

    summary: dict[str, Any] = {}
    for mode in ("non_stream", "stream"):
        successful = [r for r in records if r["mode"] == mode and r["status"] == "success"]
        totals = sorted(float(r["total_ms"]) for r in successful)
        ttfts = sorted(float(r["ttft_ms"]) for r in successful if r["ttft_ms"] is not None)
        summary[mode] = {
            "success": len(successful),
            "failed": suite_config.SPEED_RUNS - len(successful),
            "total_mean_ms": round(sum(totals) / len(totals), 3) if totals else None,
            "total_p50_ms": percentile(totals, 0.5),
            "total_p90_ms": percentile(totals, 0.9),
            "ttft_mean_ms": round(sum(ttfts) / len(ttfts), 3) if ttfts else None,
            "ttft_p50_ms": percentile(ttfts, 0.5),
            "ttft_p90_ms": percentile(ttfts, 0.9),
        }
    output.write_text(
        json.dumps(
            {"model": ctx.test_model, "prompt": suite_config.SPEED_PROMPT, "records": records, "summary": summary},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    if not any(record["status"] == "success" for record in records):
        raise RuntimeError(
            "Speed smoke test received no successful model response. "
            f"See result file: {output}"
        )
    return output


def build_long_dialog_config(test_model: str) -> dict[str, Any]:
    loaded: dict[str, Any] = {
        "dialog": {
            "loop_count": suite_config.DIALOG_LOOP_COUNT,
            "rounds_per_loop": suite_config.DIALOG_ROUNDS_PER_LOOP,
            "opening_user": suite_config.DIALOG_OPENING_USER,
            "request_timeout_sec": suite_config.DIALOG_REQUEST_TIMEOUT_SEC,
            "retry_times": suite_config.DIALOG_RETRY_TIMES,
            "retry_backoff_sec": suite_config.DIALOG_RETRY_BACKOFF_SEC,
            "worksheet_name": suite_config.DIALOG_WORKSHEET_NAME,
            "context_window": {
                "enabled": suite_config.DIALOG_CONTEXT_WINDOW_ENABLED,
                "window_messages": suite_config.DIALOG_CONTEXT_WINDOW_MESSAGES,
            },
        },
        "function_call": {
            "enabled": suite_config.DIALOG_FUNCTION_CALL_ENABLED,
            "source": "config11",
            "apply_to": suite_config.DIALOG_FUNCTION_CALL_APPLY_TO,
            "tool_choice": suite_config.DIALOG_FUNCTION_CALL_TOOL_CHOICE,
            "deduplicate_by_name": suite_config.DIALOG_FUNCTION_CALL_DEDUPLICATE_BY_NAME,
            "tools_config11_descs": list(suite_config.DIALOG_FUNCTION_CALL_TOOLS),
        },
        "memobase": {"enabled": suite_config.DIALOG_MEMOBASE_ENABLED},
        "deepseek_user": {
            "base_url": suite_config.DIALOG_USER_BASE_URL,
            "api_key": suite_config.DIALOG_USER_API_KEY,
            "model": suite_config.DIALOG_USER_MODEL,
            "temperature": suite_config.DIALOG_USER_TEMPERATURE,
            "system_prompt": suite_config.USER_PROMPT,
            "thinking_control": {
                "enabled": suite_config.DIALOG_USER_THINKING_ENABLED,
                "payload": dict(suite_config.DIALOG_USER_THINKING_PAYLOAD),
            },
        },
        "doubao_assistant": {
            "base_url": suite_config.TEST_BASE_URL,
            "api_key": suite_config.TEST_API_KEY,
            "model": test_model,
            "temperature": suite_config.DIALOG_ASSISTANT_TEMPERATURE,
            "system_prompt": suite_config.ASSISTANT_PROMPT_ACTIVE,
            "thinking_control": {
                "enabled": suite_config.DIALOG_ASSISTANT_THINKING_ENABLED,
                "payload": dict(suite_config.DIALOG_ASSISTANT_THINKING_PAYLOAD),
            },
        },
    }
    return dialog_impl.validate_and_normalize_config(loaded)


def run_dialog(ctx: SuiteContext) -> Path:
    config = build_long_dialog_config(ctx.test_model)
    dialog_cfg = config["dialog"]
    deepseek_cfg = config["deepseek_user"]
    doubao_cfg = config["doubao_assistant"]
    function_cfg = config["function_call"]
    memobase_cfg = config["memobase"]
    timeout = float(dialog_cfg["request_timeout_sec"])
    deepseek_client = dialog_impl.build_client(deepseek_cfg, timeout)
    doubao_client = dialog_impl.build_client(doubao_cfg, timeout)
    memobase_user = dialog_impl.build_memobase_user() if memobase_cfg["enabled"] else None
    context_window_cfg = dialog_cfg["context_window"]
    retries = dialog_cfg["retry_times"]
    backoff = dialog_cfg["retry_backoff_sec"]
    tools = function_cfg["tools"] if function_cfg["enabled"] else None
    tool_choice = function_cfg["tool_choice"] if tools else None
    if not function_cfg["enabled"]:
        # Long-dialog testing must not attach any Function Call tools.  The
        # standalone function_call stage remains responsible for tool tests.
        tools = None
        tool_choice = None
        print("[INFO] long-dialog Function Call disabled; tools/tool_choice omitted.")

    session_ids: list[int] = []
    lines: list[str] = []
    user_events: list[str] = []
    for session_id in range(1, dialog_cfg["loop_count"] + 1):
        history: list[dict[str, str]] = []
        tool_flags: list[bool] = []
        session_lines: list[str] = []
        session_events: list[str] = []
        first_msg = {"role": "user", "content": dialog_cfg["opening_user"]}
        history.append(first_msg)
        tool_flags.append(False)
        session_lines.append(json.dumps(first_msg, ensure_ascii=False))
        session_events.append(
            dialog_impl.fetch_user_event_with_retry(
                memobase_user=memobase_user,
                user_text=first_msg["content"],
                retries=retries,
                backoff_sec=backoff,
            )
        )
        for round_index in range(1, dialog_cfg["rounds_per_loop"] + 1):
            call_index = (session_id - 1) * dialog_cfg["rounds_per_loop"] + round_index
            if round_index % 2:
                role = "assistant"
                source_history = dialog_impl.apply_history_window(
                    history,
                    enabled=context_window_cfg["enabled"],
                    window_messages=context_window_cfg["window_messages"],
                )
                if memobase_cfg["enabled"]:
                    source_history = dialog_impl.build_assistant_prompted_history(
                        source_history, session_events[-1]
                    )
                reply = dialog_impl.create_reply_with_retry(
                    client=doubao_client,
                    model_cfg=doubao_cfg,
                    history=source_history,
                    mapped=False,
                    retries=retries,
                    backoff_sec=backoff,
                    turn=call_index,
                    role=role,
                    tools=tools if function_cfg["apply_to"] in {"assistant", "both"} else None,
                    tool_choice=tool_choice,
                )
            else:
                role = "user"
                source_history, source_flags = dialog_impl.apply_history_window_with_flags(
                    history,
                    tool_flags,
                    enabled=context_window_cfg["enabled"],
                    window_messages=context_window_cfg["window_messages"],
                )
                reply = dialog_impl.create_reply_with_retry(
                    client=deepseek_client,
                    model_cfg=deepseek_cfg,
                    history=dialog_impl.build_history_for_user_view(source_history, source_flags),
                    mapped=True,
                    retries=retries,
                    backoff_sec=backoff,
                    turn=call_index,
                    role=role,
                    tools=tools if function_cfg["apply_to"] in {"user", "both"} else None,
                    tool_choice=tool_choice,
                )
            record = {"role": role, "content": reply.text.strip()}
            history.append(record)
            tool_flags.append(role == "assistant" and reply.is_tool_call)
            session_lines.append(json.dumps(record, ensure_ascii=False))
            if role == "user":
                session_events.append(
                    dialog_impl.fetch_user_event_with_retry(
                        memobase_user=memobase_user,
                        user_text=record["content"],
                        retries=retries,
                        backoff_sec=backoff,
                    )
                )
            else:
                session_events.append("")
            if ctx.print_live:
                print(f"[S{session_id}] round={round_index} {record}", flush=True)
        dialog_impl.validate_generated_history(history, dialog_cfg["rounds_per_loop"] + 1)
        session_ids.extend([session_id] * len(session_lines))
        lines.extend(session_lines)
        user_events.extend(session_events)

    output = ctx.work_dir / "02_dialog.xlsx"
    dialog_impl.write_excel_lines(
        output,
        dialog_cfg["worksheet_name"],
        session_ids,
        lines,
        user_events,
    )
    return output


def run_function_call(ctx: SuiteContext) -> Path:
    source = suite_config.FUNCTION_CALL_SOURCE
    if not source.exists():
        raise FileNotFoundError(f"Function Call source workbook not found: {source}")
    source_workbook = function_call_impl.load_workbook(source, data_only=False)
    if function_call_impl.SOURCE_SHEET not in source_workbook.sheetnames:
        raise ValueError(f"Missing sheet {function_call_impl.SOURCE_SHEET!r} in {source}")
    source_ws = source_workbook[function_call_impl.SOURCE_SHEET]
    totals = function_call_impl.validate_source(source_ws)
    original_models = function_call_impl.MODELS
    function_call_impl.MODELS = (ctx.test_model,)
    output = ctx.work_dir / "03_function_call.xlsx"
    partial = output.with_name(f"{output.stem}_running{output.suffix}")
    persisted_partial = ctx.output_dir / partial.name
    if not partial.exists() and persisted_partial.exists():
        # A prior interrupted unified run stores its checkpoint beside the
        # total workbook. Restore a working copy before resuming it.
        shutil.copy2(persisted_partial, partial)
    try:
        # Reuse a checkpoint when the same --output-dir is run again after an
        # interruption.  The legacy implementation validates the sheet shape
        # before resuming, so a stale or incompatible workbook is rejected.
        workbook = function_call_impl.load_or_create_result(source_ws, partial)
        function_call_impl.process(source_ws, workbook, partial)
        function_call_impl.append_red_summaries(workbook)
        function_call_impl.save_checkpoint(workbook, partial)
        os.replace(partial, output)
        failed, red = function_call_impl.verify_result(output)
    finally:
        function_call_impl.MODELS = original_models
    print(f"[INFO] Function Call cases: {totals}; failed attempts: {failed}, red attempts: {red}")
    return output


def run_latency(ctx: SuiteContext) -> Path:
    client = build_test_client()
    model_cfg = {
        "alias": ctx.test_model,
        "provider": "doubao",
        "model": ctx.test_model,
        "params": {
            "temperature": suite_config.LATENCY_TEMPERATURE,
            "max_tokens": suite_config.LATENCY_MAX_TOKENS,
            "top_p": suite_config.LATENCY_TOP_P,
        },
        "system_prompt": suite_config.LATENCY_SYSTEM_PROMPT,
        "disable_thinking_payload": {"thinking": {"type": "disabled"}},
    }
    original_prompt = latency_impl.TEST_PROMPT
    latency_impl.TEST_PROMPT = suite_config.LATENCY_PROMPT
    try:
        records = [
            latency_impl._run_with_retry(client, model_cfg, index)
            for index in range(1, suite_config.LATENCY_RUNS + 1)
        ]
    finally:
        latency_impl.TEST_PROMPT = original_prompt
    output = ctx.work_dir / "04_latency.xlsx"
    latency_impl._write_excel(output, records, latency_impl._build_summary(records))
    return output


def run_score(ctx: SuiteContext, input_path: Path) -> Path:
    if not input_path.exists():
        raise FileNotFoundError(f"Score input workbook not found: {input_path}")
    dimensions = score_impl.validate_dimensions()
    for dimension in dimensions:
        dimension["model"] = ctx.review_model
    # The unified suite must evaluate against the same assistant persona that
    # generated the long-dialog workbook. The legacy standalone scorer keeps
    # its own score_dialog_config prompt for backwards compatibility, but its
    # old Nana persona must not leak into unified-suite evaluations.
    _, packed_template = score_impl.validate_shared_prompt_config()
    assistant_prompt = suite_config.ASSISTANT_PROMPT_ACTIVE.strip()
    if not assistant_prompt:
        raise ValueError("suite_config.ASSISTANT_PROMPT_ACTIVE must be non-empty.")
    client = build_review_client()
    output_dir = ctx.work_dir / "05_score"
    output = output_dir / f"{input_path.stem}_score.xlsx"
    score_impl.process_one_excel(
        input_path=input_path,
        output_path=output,
        client=client,
        dimensions=dimensions,
        assistant_prompt=assistant_prompt,
        packed_template=packed_template,
    )
    return output


def validate_only(names: list[str], input_path: Path | None, test_model: str) -> None:
    build_long_dialog_config(test_model)
    source = suite_config.FUNCTION_CALL_SOURCE
    if "function_call" in names:
        if not source.exists():
            raise FileNotFoundError(f"Missing local Function Call workbook: {source}")
        source_workbook = function_call_impl.load_workbook(source, read_only=True, data_only=False)
        if function_call_impl.SOURCE_SHEET not in source_workbook.sheetnames:
            raise ValueError(f"Missing sheet {function_call_impl.SOURCE_SHEET!r} in {source}")
        function_call_impl.validate_source(source_workbook[function_call_impl.SOURCE_SHEET])
    if "score" in names:
        score_impl.validate_dimensions()
        score_impl.validate_shared_prompt_config()
        if not isinstance(suite_config.ASSISTANT_PROMPT_ACTIVE, str) or not suite_config.ASSISTANT_PROMPT_ACTIVE.strip():
            raise ValueError("suite_config.ASSISTANT_PROMPT_ACTIVE must be non-empty.")
        if input_path is not None and not input_path.exists():
            raise FileNotFoundError(f"Missing score input workbook: {input_path}")


def main() -> int:
    args = parse_args()
    if args.list_tests:
        for name in TEST_NAMES:
            print(f"{name}: {TEST_DESCRIPTIONS[name]}")
        return 0

    names = selected_tests(args.tests)
    validate_model(args.test_model)
    configure_model_network()
    output_dir = (args.output_dir or default_output_dir()).expanduser().resolve()
    work_dir = output_dir / "_temporary"
    output_dir.mkdir(parents=True, exist_ok=True)
    ctx = SuiteContext(
        test_model=args.test_model.strip(),
        review_model=args.review_model.strip(),
        output_dir=output_dir,
        work_dir=work_dir,
        print_live=args.print_live,
        dry_run=args.dry_run,
        manifest_path=output_dir / "run_manifest.json",
    )
    ctx.write_manifest()
    print(f"[INFO] tests={names}")
    print(f"[INFO] test_model={ctx.test_model}")
    print(f"[INFO] review_model={ctx.review_model}")
    print(f"[INFO] output_dir={output_dir}")

    try:
        input_path = args.input.expanduser().resolve() if args.input else None
        if args.dry_run:
            validate_only(names, input_path, ctx.test_model)
            print("[DRY-RUN] configuration and local documents are valid")
            return 0

        ctx.work_dir.mkdir(parents=True, exist_ok=True)
        dialog_output: Path | None = None
        for name in names:
            try:
                if name == "speed":
                    path = run_speed(ctx)
                elif name == "dialog":
                    path = run_dialog(ctx)
                    dialog_output = path
                elif name == "function_call":
                    path = run_function_call(ctx)
                elif name == "latency":
                    path = run_latency(ctx)
                elif name == "score":
                    score_input = input_path or dialog_output
                    if score_input is None:
                        raise ValueError("Score testing requires --input or a successful dialog stage.")
                    path = run_score(ctx, score_input)
                else:  # pragma: no cover
                    raise ValueError(f"Unknown test: {name}")
                ctx.record(name, "success", output=path)
                print(f"[DONE] {name}: {path}")
            except Exception as exc:  # noqa: BLE001
                ctx.record(name, "failed", error=error_text(exc))
                print(f"[ERROR] {name}: {error_text(exc)}", file=sys.stderr)

        total_output = finalize_total_workbook(ctx, names, input_path)
        print(f"[DONE] total workbook: {total_output}")
        failed = [item for item in ctx.results if item["status"] == "failed"]
        return 1 if failed else 0
    except Exception as exc:  # noqa: BLE001
        ctx.record("suite", "failed", error=error_text(exc))
        print(f"[ERROR] suite: {error_text(exc)}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
