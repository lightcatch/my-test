"""Run the desktop Function Call test set against the configured Doubao model.

The script deliberately does not execute any tools.  It records the model's
natural-language message and requested tool calls, then highlights attempts
whose tool-call behavior disagrees with the test case expectation.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from config11 import (  # noqa: E402
    DOUBAO_API_KEY,
    DOUBAO_API_URL,
    MAX_TOKENS,
    SYSTEM_PROMPT_TEMPLATE,
    TEMPERATURE,
)


# Keep the legacy entry point self-contained as well.  The unified runner uses
# the same project-local workbook directly.
TEST_CASE_DIRECTORY = PROJECT_ROOT / "model_test" / "test_data"
DEFAULT_SOURCE = TEST_CASE_DIRECTORY / "functioncall测试用例.xlsx"
DEFAULT_OUTPUT_DIRECTORY = PROJECT_ROOT / "model_test" / "outputs" / "function_call"
SOURCE_SHEET = "Sheet1"

MODELS = (
    "doubao-seed-2-1-lite-260915",
)
# 中文编码检查：Function Call 工具描述必须保存为 UTF-8 中文原文。

# Keep the Function Call contract local to this standalone test program.  The
# names, parameter types, and required fields match the test workbook; the
# descriptions deliberately spell out trigger boundaries so model behavior is
# less dependent on another module's evolving configuration.
play_music_function_desc = {
    "type": "function",
    "function": {
        "name": "play_music",
        "description": """这是播放音乐工具，只在用户当前消息明确要求开始播放或聆听时调用。
        触发词包括播放、放一首、听音乐、想听、播一下、整点音乐、放首歌、继续播放，以及明确指定歌曲、歌手或音乐类型并要求播放的表达；
        例如‘播放《两只老虎》’、‘我想听周杰伦的歌’、‘随便放点旋律’都必须调用。
        用户只是在问歌词、歌手、歌曲背景、歌曲评价、音乐知识，想学唱或学乐器、看MV、下载或购买唱片、设置闹钟，或只是提到歌曲而没有播放动作时绝对不要调用；
        ‘我还没听过这首歌’也不是播放请求。明确要求播放但没有具体曲目时调用并将 song_name 填为 random；有歌曲或歌手时原样填写用户指定目标，不得改调用其他工具，不要因回复角色需要而省略工具调用。""",
        "parameters": {
            "type": "object",
            "properties": {
                "song_name": {
                    "type": "string",
                    "description": """要播放的歌曲名称、歌手或音乐类型。明确播放但没有具体目标时必须填写 random；用户指定了歌曲、歌手或组合时原样填写，不要自行编造歌名。只有明确播放意图才填写此参数。""",
                }
            },
            "required": ["song_name"],
        },
    },
}

GET_WEATHER_FUNCTION_DESC = {
    "type": "function",
    "function": {
        "name": "get_weather",
        "description": """这是实时天气查询工具。只要用户明确要知道真实地点当前、今天、明天或未来的天气、气温、下雨、风力或天气预报，就必须调用；
        地点可以是城市、省份、景点或国家，例如‘杭州天气怎么样’、‘今天天气如何’、‘洛杉矶冷不冷’、‘What's the weather like in London?’。未提供地点但明确询问天气时也调用并把 location 留空。
        省份使用省会，景点按所在城市填写；英文问题 lang 用 en_US，中文默认 zh_CN。只谈天气这个词、天气预报原理、电影《天气之子》、天气比喻或感受、历史天气知识、天气应用推荐、照片拍摄时天气，或对天气分类做增删时不要调用。
        不要把一般知识问题当实时查询。""",
        "parameters": {
            "type": "object",
            "properties": {
                "location": {
                    "type": "string",
                    "description": """实时天气地点，如杭州、北京、洛杉矶或London。用户未提供地点时留空，不要猜测。""",
                },
                "lang": {
                    "type": "string",
                    "description": """回复语言代码。中文用户填写 zh_CN，英文用户填写 en_US；默认 zh_CN。""",
                },
            },
            "required": ["lang"],
        },
    },
}

get_time_function_desc = {
    "type": "function",
    "function": {
        "name": "get_time",
        "description": """这是获取系统当前日期和时间的工具。用户明确询问现在几点、当前时间、今天几号、今天星期几、当前年份、北京时间、精确分秒，或要求先取得当前日期来判断节气时必须调用；
        例如‘告诉我今天是周几’和‘今天是哪一年’都要调用。用户谈历史年份、时间一词、任务耗时、闹钟、倒计时、计时器、网络延迟、设备运行时间、歌曲或一般时间知识时不要调用。
        此函数没有参数，调用时只传空参数对象。""",
        "parameters": {"type": "object", "properties": {}, "required": []},
    },
}

exit_function_desc = {
    "type": "function",
    "function": {
        "name": "exit_function",
        "description": """这是结束当前对话的工具。用户明确表示结束、退出、注销会话、登出、离开当前聊天，或使用‘再见’、‘拜拜’、‘晚安’、‘不聊了’、‘结束当前对话’等明确告别且没有继续聊天意图时必须调用；
        即使只有两个字的‘再见’或‘拜拜’也要调用。‘再见，我的任务完成了’、‘拜拜，我得走了’同样是结束。
        用户只是去洗手、吃饭、找人、看电视、暂时离开、等一分钟、想想、换话题，明确说等会回来，或只是感谢而未结束时绝对不要调用。
        调用时 say_goodbye 写简短友好告别语；工具调用优先于角色化长回复。""",
        "parameters": {
            "type": "object",
            "properties": {
                "say_goodbye": {
                    "type": "string",
                    "description": """简短友好的结束语，必须表达对话已结束，例如‘再见’或‘下次需要时再来’。""",
                }
            },
            "required": ["say_goodbye"],
        },
    },
}

search_diary_function_desc = {
    "type": "function",
    "function": {
        "name": "search_diary",
        "description": """这是检索已经保存内容的日记查询工具。用户明确要求查询、找找、搜索、念一下、复述、重温或找回已保存的日记、信件、文章、诗歌或历史记录时必须调用；‘查询日记’、‘那封信里写了什么’、‘找找上周的所有记录’都属于触发。
        只评论或批评一篇内容、准备写新日记、写作文、买日记本、丢了日记本、没写过日记、查看容量、保存/删除/修改内容、查询写日记方法，或提到别人的文章但不要求检索时绝对不要调用。
        没有日期但明确要求找回已存内容可以调用并留空日期字段；
        单日填写 search_date，日期范围填写 search_start_date 和 search_end_date，不能编造日期，也不要同时填互斥字段。""",
        "parameters": {
            "type": "object",
            "properties": {
                "search_date": {
                    "type": "string",
                    "description": """用户明确指定的单个日记日期，如2024年5月20日；没有单日日期时留空。""",
                },
                "search_start_date": {
                    "type": "string",
                    "description": """用户明确给出的查询范围开始日期；只有范围查询时填写，不要凭空猜日期。""",
                },
                "search_end_date": {
                    "type": "string",
                    "description": """用户明确给出的查询范围结束日期；只有范围查询时填写，不要凭空猜日期。""",
                },
            },
        },
    },
}

search_news_function_desc = {
    "type": "function",
    "function": {
        "name": "search_news",
        "description": """这是实时新闻检索工具。用户要求今天或最近的新闻、新闻播报、最新战报、国际大事、热点消息、某事件最新进展、人物当前动态，或说‘听新闻’、‘随便给我听点消息’时必须调用。
        用户未给具体标题但要浏览新闻时调用，news_title 留空；用户指定事件或主题时填入原话，random_search 设为 false；只有随机听最新新闻时 random_search 才为 true。
        新闻定义、历史知识、电视剧评价、专业学习、辨别假新闻、专业排名、删除或保存新闻、用户自己说‘我有个新闻’等非实时请求绝对不要调用；
        ‘不用查最新’是明确的禁止信号。不要因为出现‘新闻’二字就触发。""",
        "parameters": {
            "type": "object",
            "properties": {
                "news_title": {
                    "type": "string",
                    "description": """要查询的实时新闻主题或事件名称；没有具体主题时留空。""",
                },
                "news_type": {
                    "type": "string",
                    "description": """新闻分类，如社会、科技、财经、体育、国际；无法判断时留空。""",
                },
                "random_search": {
                    "type": "boolean",
                    "description": """随机浏览最新新闻时为 true；指定主题或事件时为 false。""",
                },
            },
            "required": ["news_title", "random_search"],
        },
    },
}

TOOLS = (
    play_music_function_desc,
    GET_WEATHER_FUNCTION_DESC,
    get_time_function_desc,
    exit_function_desc,
    search_diary_function_desc,
    search_news_function_desc,
)
FUNCTION_NAMES = tuple(tool["function"]["name"] for tool in TOOLS)

# The persona prompt is retained for realistic responses, but tool selection
# must be decided by the user's explicit intent before role-play prose.  This
# short protocol prevents the model from replacing an expected tool call with
# a long in-character reply or from treating a shared keyword as a trigger.
FUNCTION_CALL_SYSTEM_SUFFIX = """

<Function Call 测试判定协议>
本轮只做工具意图识别，不执行任何工具。先判断用户最后一条消息的明确请求，再决定是否调用工具；工具调用优先于角色化回复。每条消息最多调用一个工具，必须使用与意图完全匹配的工具；没有明确触发条件时不要调用任何工具。
退出：再见、拜拜、晚安、不聊了、退出、注销会话、登出、结束当前对话等明确结束当前聊天的表达必须调用 exit_function；只是去洗手、吃饭、暂时离开、等会回来、想想或换话题不得调用。
日记：明确查询、搜索、找回、重温、念出已保存日记、信件、文章或记录时调用 search_diary；评论、批评、写新内容、保存/删除/修改、容量和写作建议不得调用。
新闻：请求今天/最近的新闻、播报、热点、最新进展、听新闻时调用 search_news；历史知识、新闻定义、电视剧、专业学习、假新闻辨别、删除保存或“我有个新闻”不得调用。
时间：询问现在几点、今天日期/星期/年份、当前北京时间或要求先取当前日期时调用 get_time；历史年份、闹钟、倒计时、任务耗时、设备运行时间和网络延迟不得调用。
天气：询问真实地点当前或未来天气、温度、下雨、风力或预报时调用 get_weather；天气原理、电影、比喻、照片天气、天气应用推荐和分类操作不得调用。
音乐：明确要求播放、听、放一首或继续播放时调用 play_music；歌词、歌手/背景/评价、学唱学乐器、MV、下载购买、闹钟或“还没听过”不得调用。
调用时严格按工具 schema 填参数，不要凭空编造地点或日期；边界场景即使含有对应关键词也保持不调用。
</Function Call 测试判定协议>
"""

REQUIRED_HEADERS = (
    "用例编号",
    "关联算法设计",
    "用例标题",
    "用例类型",
    "用例优先级",
    "前置条件",
    "用户对话",
    "期望回复结果",
    "重复执行次数",
    "自动化",
    "人工复核",
    "维护人",
    "备注",
)
REPLY_HEADERS = ("LLM回复1", "函数调用1", "LLM回复2", "函数调用2")
SUMMARY_LABEL = "标红数量（回复轮次）"

UNIQUE_RED_CASES_LABEL = "唯一标红用例数"

FAIL_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
FAIL_FONT = Font(color="9C0006")
NORMAL_FILL = PatternFill(fill_type=None)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="批量执行 Function Call 测试用例")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument(
        "--output",
        type=Path,
        help="结果工作簿路径；省略时自动使用带启动时间的文件名",
    )
    parser.add_argument(
        "--partial",
        type=Path,
        help="断点续跑文件路径；省略时依据结果文件名自动生成",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只验证输入工作簿和用例映射，不调用模型或写结果文件",
    )
    return parser.parse_args()


def build_timestamped_output_path() -> Path:
    """Create a collision-resistant project-local result path for a new batch run."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = DEFAULT_OUTPUT_DIRECTORY / f"functioncall测试结果_{timestamp}.xlsx"
    suffix = 1
    while candidate.exists():
        candidate = TEST_CASE_DIRECTORY / f"functioncall测试结果_{timestamp}_{suffix}.xlsx"
        suffix += 1
    return candidate


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def header_columns(ws: Any) -> dict[str, int]:
    return {
        cell_text(ws.cell(1, column).value): column
        for column in range(1, ws.max_column + 1)
        if cell_text(ws.cell(1, column).value)
    }


def expected_function(expected: str) -> str:
    # Chinese text immediately preceding a function name is also a ``\w``
    # character to Python, so a word-boundary expression would miss strings
    # such as "调用exit_function".
    matches = [name for name in FUNCTION_NAMES if name in expected]
    if len(matches) != 1:
        raise ValueError(f"无法从期望结果识别唯一目标 function：{expected}")
    return matches[0]


def validate_source(ws: Any) -> dict[str, int]:
    headers = header_columns(ws)
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"输入用例集缺少列：{', '.join(missing)}")

    data_rows = [row for row in range(2, ws.max_row + 1) if cell_text(ws.cell(row, headers["用例编号"]).value)]
    if len(data_rows) != 120:
        raise ValueError(f"输入用例集应有120条非空用例，实际为{len(data_rows)}条")

    case_ids = [cell_text(ws.cell(row, headers["用例编号"]).value) for row in data_rows]
    if len(set(case_ids)) != len(case_ids):
        raise ValueError("输入用例集存在重复的用例编号")

    totals = {"功能测试": 0, "边界测试": 0}
    for row in data_rows:
        case_type = cell_text(ws.cell(row, headers["用例类型"]).value)
        expected = cell_text(ws.cell(row, headers["期望回复结果"]).value)
        dialogue = cell_text(ws.cell(row, headers["用户对话"]).value)
        if not dialogue:
            raise ValueError(f"第{row}行用户对话为空")
        if case_type not in totals:
            raise ValueError(f"第{row}行用例类型不支持：{case_type}")
        totals[case_type] += 1
        if case_type == "功能测试":
            expected_function(expected)
        elif "不应调用任何function" not in expected:
            raise ValueError(f"第{row}行边界用例未明确要求不调用 function")
    return totals


def clone_sheet(source_ws: Any, target_ws: Any) -> None:
    """Copy values and visible worksheet formatting without retaining extra sheets."""
    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = target_ws.cell(source_cell.row, source_cell.column, source_cell.value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)
            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)
            if source_cell.fill:
                target_cell.fill = copy(source_cell.fill)
            if source_cell.font:
                target_cell.font = copy(source_cell.font)
            if source_cell.border:
                target_cell.border = copy(source_cell.border)

    for key, dimension in source_ws.column_dimensions.items():
        target_ws.column_dimensions[key].width = dimension.width
        target_ws.column_dimensions[key].hidden = dimension.hidden
    for index, dimension in source_ws.row_dimensions.items():
        target_ws.row_dimensions[index].height = dimension.height
        target_ws.row_dimensions[index].hidden = dimension.hidden

    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))
    target_ws.freeze_panes = source_ws.freeze_panes

    source_headers = header_columns(source_ws)
    reference_column = source_headers["备注"]
    for offset, header in enumerate(REPLY_HEADERS, start=1):
        column = source_ws.max_column + offset
        target = target_ws.cell(1, column, header)
        source_header = source_ws.cell(1, reference_column)
        target._style = copy(source_header._style)
        target.alignment = copy(source_header.alignment)
        target.font = copy(source_header.font)
        target.fill = copy(source_header.fill)
        target.border = copy(source_header.border)
        target.protection = copy(source_header.protection)
        target.number_format = source_header.number_format
        target_ws.column_dimensions[get_column_letter(column)].width = 30

    target_ws.auto_filter.ref = f"A1:{get_column_letter(source_ws.max_column + len(REPLY_HEADERS))}{source_ws.max_row}"


def new_result_workbook(source_ws: Any) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for model in MODELS:
        sheet = workbook.create_sheet(title=model)
        clone_sheet(source_ws, sheet)
    return workbook


def load_or_create_result(source_ws: Any, partial_path: Path) -> Workbook:
    if partial_path.exists():
        workbook = load_workbook(partial_path)
        if workbook.sheetnames != list(MODELS):
            raise ValueError(f"执行中工作簿工作表不匹配：{workbook.sheetnames}")
        for sheet in workbook.worksheets:
            if sheet.max_column != len(REQUIRED_HEADERS) + len(REPLY_HEADERS):
                raise ValueError(f"执行中工作簿列数异常：{sheet.title}")
        print(f"恢复执行中工作簿：{partial_path}")
        return workbook
    return new_result_workbook(source_ws)


def user_message(dialogue: str) -> str:
    return re.sub(r"^\s*user\s*:\s*", "", dialogue, count=1, flags=re.IGNORECASE).strip()


def normalize_tool_calls(tool_calls: Optional[Iterable[Any]]) -> list[dict[str, str]]:
    normalized: list[dict[str, str]] = []
    for tool_call in tool_calls or []:
        function = getattr(tool_call, "function", None)
        name = cell_text(getattr(function, "name", ""))
        arguments = cell_text(getattr(function, "arguments", ""))
        normalized.append({"name": name, "arguments": arguments})
    return normalized


def call_model(client: OpenAI, model: str, content: str) -> tuple[str, list[dict[str, str]], Optional[str]]:
    """Return reply, requested function calls, and a sanitized final failure reason."""
    last_error = ""
    for attempt in range(1, 4):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT_TEMPLATE + FUNCTION_CALL_SYSTEM_SUFFIX},
                    {"role": "user", "content": content},
                ],
                tools=list(TOOLS),
                temperature=TEMPERATURE,
                max_tokens=MAX_TOKENS,
                extra_body={"thinking": {"type": "disabled"}},
            )
            message = response.choices[0].message
            return cell_text(message.content), normalize_tool_calls(message.tool_calls), None
        except Exception as error:  # Retry all client/API failures; final text is deliberately short.
            last_error = type(error).__name__
            if attempt < 3:
                time.sleep(attempt * 2)
    return f"API调用失败：{last_error}", [], last_error


def is_pass(case_type: str, expected: str, calls: list[dict[str, str]], api_error: Optional[str]) -> bool:
    if api_error:
        return False
    names = [call["name"] for call in calls]
    if case_type == "功能测试":
        return names == [expected_function(expected)]
    if case_type == "边界测试":
        return not names
    raise ValueError(f"未知用例类型：{case_type}")


def set_attempt_cells(
    ws: Any,
    row: int,
    attempt: int,
    reply: str,
    calls: list[dict[str, str]],
    api_error: Optional[str],
    passed: bool,
) -> None:
    reply_column = len(REQUIRED_HEADERS) + (attempt - 1) * 2 + 1
    calls_column = reply_column + 1
    reply_cell = ws.cell(row, reply_column, reply)
    calls_cell = ws.cell(
        row,
        calls_column,
        "API调用失败" if api_error else json.dumps(calls, ensure_ascii=False),
    )
    if passed:
        return
    for cell in (reply_cell, calls_cell):
        cell.fill = copy(FAIL_FILL)
        cell.font = copy(FAIL_FONT)


def attempt_is_complete(ws: Any, row: int, attempt: int) -> bool:
    reply_column = len(REQUIRED_HEADERS) + (attempt - 1) * 2 + 1
    calls_column = reply_column + 1
    return ws.cell(row, reply_column).value is not None and ws.cell(row, calls_column).value is not None


def save_checkpoint(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def process(source_ws: Any, workbook: Workbook, partial_path: Path) -> None:
    headers = header_columns(source_ws)
    client = OpenAI(api_key=DOUBAO_API_KEY, base_url=DOUBAO_API_URL)
    data_rows = [row for row in range(2, source_ws.max_row + 1) if cell_text(source_ws.cell(row, headers["用例编号"]).value)]

    for model in MODELS:
        ws = workbook[model]
        completed_since_checkpoint = 0
        print(f"开始模型：{model}", flush=True)
        for row in data_rows:
            case_id = cell_text(source_ws.cell(row, headers["用例编号"]).value)
            case_type = cell_text(source_ws.cell(row, headers["用例类型"]).value)
            expected = cell_text(source_ws.cell(row, headers["期望回复结果"]).value)
            content = user_message(cell_text(source_ws.cell(row, headers["用户对话"]).value))
            for attempt in (1, 2):
                if attempt_is_complete(ws, row, attempt):
                    continue
                reply, calls, api_error = call_model(client, model, content)
                passed = is_pass(case_type, expected, calls, api_error)
                set_attempt_cells(ws, row, attempt, reply, calls, api_error, passed)
                status = "通过" if passed else "标红"
                print(f"{model} | {case_id} | 第{attempt}次 | {status}", flush=True)
                completed_since_checkpoint += 1
                if completed_since_checkpoint >= 10:
                    save_checkpoint(workbook, partial_path)
                    completed_since_checkpoint = 0
        save_checkpoint(workbook, partial_path)


def append_red_summaries(workbook: Workbook) -> dict[str, dict[str, Any]]:
    """Append aggregate and per-case-type red statistics at each sheet end."""
    summaries: dict[str, dict[str, Any]] = {}
    for ws in workbook.worksheets:
        has_existing_summary = ws.cell(ws.max_row, 1).value == SUMMARY_LABEL
        data_end_row = ws.max_row - 1 if has_existing_summary else ws.max_row
        red_attempts = sum(
            1
            for row in range(2, data_end_row + 1)
            for reply_column in (len(REQUIRED_HEADERS) + 1, len(REQUIRED_HEADERS) + 3)
            if ws.cell(row, reply_column).fill.fgColor.rgb == "00FFC7CE"
        )
        red_cases = sum(
            1
            for row in range(2, data_end_row + 1)
            if any(
                ws.cell(row, reply_column).fill.fgColor.rgb == "00FFC7CE"
                for reply_column in (len(REQUIRED_HEADERS) + 1, len(REQUIRED_HEADERS) + 3)
            )
        )
        category_stats: dict[str, dict[str, int]] = {}
        for row in range(2, data_end_row + 1):
            case_type = cell_text(ws.cell(row, 4).value)
            stats = category_stats.setdefault(case_type, {"red_attempts": 0, "red_cases": 0})
            row_red_attempts = sum(
                ws.cell(row, reply_column).fill.fgColor.rgb == "00FFC7CE"
                for reply_column in (len(REQUIRED_HEADERS) + 1, len(REQUIRED_HEADERS) + 3)
            )
            stats["red_attempts"] += row_red_attempts
            if row_red_attempts:
                stats["red_cases"] += 1
        summary_row = data_end_row + 1
        label_cell = ws.cell(summary_row, 1, SUMMARY_LABEL)
        value_cell = ws.cell(summary_row, 2, red_attempts)
        unique_label_cell = ws.cell(summary_row, 3, UNIQUE_RED_CASES_LABEL)
        unique_value_cell = ws.cell(summary_row, 4, red_cases)
        label_cell.font = Font(bold=True)
        value_cell.font = Font(bold=True)
        unique_label_cell.font = Font(bold=True)
        unique_value_cell.font = Font(bold=True)
        summary_cells = (
            (5, "总用例数", data_end_row - 1),
            (7, "总尝试数", (data_end_row - 1) * 2),
            (9, "失败尝试数", red_attempts),
        )
        for column, label, value in summary_cells:
            ws.cell(summary_row, column, label).font = Font(bold=True)
            ws.cell(summary_row, column + 1, value).font = Font(bold=True)
        for offset, (case_type, stats) in enumerate(category_stats.items()):
            column = 11 + offset * 2
            if column + 1 > 17:
                break
            ws.cell(summary_row, column, case_type).font = Font(bold=True)
            ws.cell(
                summary_row,
                column + 1,
                f"标红尝试 {stats['red_attempts']} 次；标红条例 {stats['red_cases']} 条",
            ).font = Font(bold=True)
        summaries[ws.title] = {
            "total_cases": data_end_row - 1,
            "total_attempts": (data_end_row - 1) * 2,
            "failed_attempts": red_attempts,
            "red_attempts": red_attempts,
            "red_cases": red_cases,
            "category_stats": category_stats,
        }
    return summaries


def verify_result(path: Path) -> tuple[int, int]:
    workbook = load_workbook(path, data_only=False)
    if workbook.sheetnames != list(MODELS):
        raise ValueError(f"结果工作簿工作表不正确：{workbook.sheetnames}")

    failed_attempts = 0
    red_attempts = 0
    for ws in workbook.worksheets:
        if ws.max_row != 122 or ws.max_column != 17:
            raise ValueError(f"{ws.title} 结构异常：{ws.max_row}行，{ws.max_column}列")
        if ws.cell(122, 1).value != SUMMARY_LABEL:
            raise ValueError(f"{ws.title} 缺少末尾标红数量")
        headers = header_columns(ws)
        if [headers.get(header) for header in REQUIRED_HEADERS + REPLY_HEADERS] != list(range(1, 18)):
            raise ValueError(f"{ws.title} 列顺序异常")
        sheet_failed_attempts = 0
        sheet_red_cases: set[int] = set()
        for row in range(2, 122):
            case_type = cell_text(ws.cell(row, headers["用例类型"]).value)
            expected = cell_text(ws.cell(row, headers["期望回复结果"]).value)
            for attempt in (1, 2):
                reply_col = len(REQUIRED_HEADERS) + (attempt - 1) * 2 + 1
                calls_col = reply_col + 1
                calls_text = cell_text(ws.cell(row, calls_col).value)
                api_error = calls_text == "API调用失败"
                calls = [] if api_error else json.loads(calls_text)
                passed = is_pass(case_type, expected, calls, "api_error" if api_error else None)
                red = ws.cell(row, reply_col).fill.fgColor.rgb == "00FFC7CE"
                if not passed:
                    failed_attempts += 1
                    sheet_failed_attempts += 1
                    if red:
                        red_attempts += 1
                        sheet_red_cases.add(row)
                elif red:
                    raise ValueError(f"{ws.title} 第{row}行第{attempt}次通过结果被错误标红")
        if ws.cell(122, 2).value != sheet_failed_attempts:
            raise ValueError(
                f"{ws.title} 末尾标红数量异常：{ws.cell(122, 2).value}，应为{sheet_failed_attempts}"
            )
        if ws.cell(122, 3).value != UNIQUE_RED_CASES_LABEL or ws.cell(122, 4).value != len(sheet_red_cases):
            raise ValueError(
                f"{ws.title} unique red case count mismatch: {ws.cell(122, 4).value}; "
                f"expected {len(sheet_red_cases)}"
            )
    if failed_attempts != red_attempts:
        raise ValueError(f"红标数量异常：失败{failed_attempts}次，红标{red_attempts}次")
    return failed_attempts, red_attempts


def main() -> None:
    args = parse_args()
    output_path = args.output or build_timestamped_output_path()
    partial_path = args.partial or output_path.with_name(f"{output_path.stem}_执行中{output_path.suffix}")
    if not args.source.exists():
        raise FileNotFoundError(f"未找到输入用例集：{args.source}")

    source_workbook = load_workbook(args.source, data_only=False)
    if SOURCE_SHEET not in source_workbook.sheetnames:
        raise ValueError(f"输入工作簿缺少 {SOURCE_SHEET}")
    source_ws = source_workbook[SOURCE_SHEET]
    totals = validate_source(source_ws)
    print(f"输入验证通过：功能测试{totals['功能测试']}条，边界测试{totals['边界测试']}条")
    print(f"计划调用：{len(MODELS) * 120 * 2}次")
    if args.dry_run:
        return

    workbook = load_or_create_result(source_ws, partial_path)
    process(source_ws, workbook, partial_path)
    summaries = append_red_summaries(workbook)
    save_checkpoint(workbook, partial_path)
    os.replace(partial_path, output_path)
    failed_attempts, red_attempts = verify_result(output_path)
    model_summary = summaries[MODELS[0]]
    category_summary = "；".join(
        f"{case_type}：标红尝试 {stats['red_attempts']} 次，标红条例 {stats['red_cases']} 条"
        for case_type, stats in model_summary["category_stats"].items()
    )
    print(
        f"统计：总用例 {model_summary['total_cases']} 条；总尝试 {model_summary['total_attempts']} 次；"
        f"失败尝试 {failed_attempts} 次；标红尝试 {red_attempts} 次；"
        f"唯一标红条例 {model_summary['red_cases']} 条；{category_summary}"
    )
    print(f"完成：失败并标红{failed_attempts}次；结果文件：{output_path}")
    if failed_attempts != red_attempts:
        raise RuntimeError("结果验证失败")


if __name__ == "__main__":
    main()
