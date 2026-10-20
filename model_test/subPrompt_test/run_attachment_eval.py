# -*- coding: utf-8 -*-
"""Evaluate completed Nana/ai抱枕 dialogues with DeepSeek and Qwen Max.

DeepSeek judges attachment using the configured Nana prompt. Qwen Max then
reviews that judgement and assigns a single 0-10 quality score. The source
workbook is never overwritten; a timestamped result workbook is created.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# --------------------------- Editable defaults ---------------------------
PROJECT_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from config11 import (  # noqa: E402
    Ali_API_KEY,
    Ali_API_URL,
    DEEPSEEK_API_KEY,
    DEEPSEEK_API_URL,
    DEEPSEEK_AVAILABLE_MODELS,
    DOUBAO_API_KEY,
    DOUBAO_API_URL,
    DOUBAO_AVAILABLE_MODELS,
    XUNFEI_API_KEY,
    XUNFEI_API_URL,
    XUNFEI_AVAILABLE_MODELS,
)

DEFAULT_INPUT_XLSX = SCRIPT_DIR / "nana_dialogue_test_cases.xlsx"
DEFAULT_OUTPUT_DIR = SCRIPT_DIR
DEFAULT_EVALUATOR_PROMPT_PATH = Path(
    r"D:\PyCode\gitcode\pillow-ai\prompts\Nana_prompt\抱枕_奈奈_依恋度计算提示词.txt"
)

# This prompt evaluates DeepSeek's attachment judgement, not the original
# dialogue directly. Keep the JSON contract stable so the Excel summary works.
DEFAULT_JUDGE_SYSTEM_PROMPT = """
你是 ai抱枕奈奈对话评测的复核专家。请根据：
1. 依恋度计算提示词；
2. 已完成的用户与奈奈消息记录；
3. DeepSeek 对该记录给出的依恋度判定；
判断 DeepSeek 的判定是否准确、是否严格遵守提示词、是否正确区分用户消息中的深度情感互动与普通功能问答。

请只输出一个 JSON 对象，不要输出 Markdown 代码块或额外解释：
{
  "score": 0,
  "evaluation": "对 DeepSeek 判定质量的总体评价",
  "problems": "发现的问题；没有问题时填写无",
  "suggestions": "具体改进建议；没有建议时填写无"
}

score 必须是 0 到 10 的数字。10 表示判定准确、证据充分且严格遵守格式和逻辑；0 表示判定与消息记录或提示词基本无关。
""".strip()

DEFAULT_QWEN_MODEL = "qwen3-max"
DEEPSEEK_TEMPERATURE = 0.2
DEEPSEEK_MAX_TOKENS = 2048
QWEN_TEMPERATURE = 0.1
QWEN_MAX_TOKENS = 2048
REQUEST_TIMEOUT_SEC = 1200.0
REQUEST_RETRIES = 3
RETRY_BACKOFF_SEC = 2.0

# Both requests explicitly enable model reasoning. The APIs expose different
# payload names, so keep the provider-specific bodies separate.
DEEPSEEK_THINKING_BODY = {"thinking": {"type": "enabled"}}
QWEN_THINKING_BODY = {"enable_thinking": True}


class ModelConfig:
    """显式模型切换配置，改这里即可切换评判模型提供方。"""

    def __init__(self) -> None:
        # 依恋度评判模型：默认 DeepSeek Flash。
        self.used_model = DEEPSEEK_AVAILABLE_MODELS[0]
        self.used_api_key = DEEPSEEK_API_KEY
        self.used_base_url = DEEPSEEK_API_URL

        # 下面保留显式切换示例；切换时请同时修改三行。
        # self.used_model = DOUBAO_AVAILABLE_MODELS[17]
        # self.used_api_key = DOUBAO_API_KEY
        # self.used_base_url = DOUBAO_API_URL

        # self.used_model = "qwen3.5-plus"
        # self.used_api_key = Ali_API_KEY
        # self.used_base_url = Ali_API_URL

        # self.used_model = XUNFEI_AVAILABLE_MODELS[0]
        # self.used_api_key = XUNFEI_API_KEY
        # self.used_base_url = XUNFEI_API_URL

        # Qwen 复核模型：默认 qwen3-max；同样支持显式切换。
        self.judge_model = DEFAULT_QWEN_MODEL
        self.judge_api_key = Ali_API_KEY
        self.judge_base_url = Ali_API_URL

        # self.judge_model = DOUBAO_AVAILABLE_MODELS[17]
        # self.judge_api_key = DOUBAO_API_KEY
        # self.judge_base_url = DOUBAO_API_URL

        # self.judge_model = DEEPSEEK_AVAILABLE_MODELS[1]
        # self.judge_api_key = DEEPSEEK_API_KEY
        # self.judge_base_url = DEEPSEEK_API_URL


MODEL_CONFIG = ModelConfig()

RESULT_HEADERS = [
    "DeepSeek依恋度评判",
    "DeepSeekObtained",
    "Qwen评分",
    "Qwen评价",
    "Qwen问题",
    "Qwen改进建议",
    "DeepSeek状态",
    "Qwen状态",
    "DeepSeek耗时(ms)", 
    "Qwen耗时(ms)", 
    "Qwen原始输出",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="对已完成的 ai抱枕奈奈消息记录进行 DeepSeek 依恋度评判和 Qwen 复核"
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT_XLSX, help="输入 Excel")
    parser.add_argument("--output", type=Path, default=None, help="新增结果 Excel；默认带时间戳")
    parser.add_argument(
        "--prompt-path",
        type=Path,
        default=DEFAULT_EVALUATOR_PROMPT_PATH,
        help="DeepSeek 依恋度评判提示词文件",
    )
    parser.add_argument("--judge-prompt-path", type=Path, default=None, help="Qwen 复核提示词文件")
    parser.add_argument("--deepseek-model", default=MODEL_CONFIG.used_model)
    parser.add_argument("--qwen-model", default=MODEL_CONFIG.judge_model)
    parser.add_argument("--limit", type=int, default=None, help="只处理前 N 条用例")
    parser.add_argument("--start-row", type=int, default=2, help="从 Excel 行号开始处理")
    parser.add_argument("--resume", action="store_true", help="在已有输出文件上跳过成功行")
    parser.add_argument("--dry-run", action="store_true", help="只校验输入和提示词，不调用模型")
    return parser.parse_args()


def resolve_path(path: Path) -> Path:
    if path.is_absolute():
        return path
    return (PROJECT_ROOT / path).resolve()


def read_text(path: Path) -> str:
    path = resolve_path(path)
    if not path.exists():
        raise FileNotFoundError(f"提示词文件不存在: {path}")
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", raw, 0, len(raw), f"无法解码提示词: {path}")


def load_judge_prompt(path: Path | None) -> str:
    if path is None:
        return DEFAULT_JUDGE_SYSTEM_PROMPT
    return read_text(path)


def default_output_path(input_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return input_path.parent / f"ai抱枕_依恋度评测_{stamp}.xlsx"


def header_map(ws: Any) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None and str(cell.value).strip()
    }


def validate_input(ws: Any) -> tuple[dict[str, int], list[int]]:
    headers = header_map(ws)
    required = ("用例编号", "用例标题", "用户对话")
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"Sheet1 缺少列: {', '.join(missing)}")

    case_col = headers["用例编号"]
    dialogue_col = headers["用户对话"]
    rows: list[int] = []
    seen: set[str] = set()
    for row in range(2, ws.max_row + 1):
        case_id = str(ws.cell(row, case_col).value or "").strip()
        dialogue = str(ws.cell(row, dialogue_col).value or "").strip()
        if not case_id and not dialogue:
            continue
        if not case_id:
            raise ValueError(f"第 {row} 行用例编号为空")
        if not dialogue:
            raise ValueError(f"第 {row} 行用户对话为空")
        if case_id in seen:
            raise ValueError(f"用例编号重复: {case_id}")
        seen.add(case_id)
        rows.append(row)
    if not rows:
        raise ValueError("没有找到有效用例")
    return headers, rows


def append_result_columns(ws: Any) -> dict[str, int]:
    headers = header_map(ws)
    if all(header in headers for header in RESULT_HEADERS):
        return headers

    first_result_col = ws.max_column + 1
    style_source = ws.cell(1, min(ws.max_column, 13))
    for offset, header in enumerate(RESULT_HEADERS):
        col = first_result_col + offset
        cell = ws.cell(1, col, header)
        cell._style = copy(style_source._style)
        cell.font = copy(style_source.font)
        cell.fill = copy(style_source.fill)
        cell.border = copy(style_source.border)
        cell.alignment = copy(style_source.alignment)
        cell.number_format = style_source.number_format
        ws.column_dimensions[get_column_letter(col)].width = 32 if offset in (0, 3, 4, 5, 10) else 18

    for row in range(2, ws.max_row + 1):
        for offset in range(len(RESULT_HEADERS)):
            col = first_result_col + offset
            source = ws.cell(row, min(ws.max_column, 13))
            target = ws.cell(row, col)
            target._style = copy(source._style)
            target.alignment = Alignment(vertical="top", wrap_text=offset in (0, 3, 4, 5, 10))
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 240)
    return header_map(ws)


def build_client(api_key: str, base_url: str) -> OpenAI:
    if not api_key.strip() or not base_url.strip():
        raise ValueError("模型 API 配置为空，请检查 config11.py / score_dialog_config.py")
    return OpenAI(api_key=api_key, base_url=base_url, timeout=REQUEST_TIMEOUT_SEC, max_retries=0)


def extract_content(response: Any) -> str:
    message = response.choices[0].message
    content = getattr(message, "content", None)
    if isinstance(content, str) and content.strip():
        return content.strip()
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            text = getattr(item, "text", None)
            if isinstance(text, str) and text.strip():
                parts.append(text.strip())
        return "\n".join(parts).strip()
    refusal = getattr(message, "refusal", None)
    return refusal.strip() if isinstance(refusal, str) else ""


def extract_stream_content(stream: Any) -> str:
    parts: list[str] = []
    try:
        for chunk in stream:
            choices = getattr(chunk, "choices", None) or []
            if not choices:
                continue
            delta = getattr(choices[0], "delta", None)
            content = getattr(delta, "content", None) if delta else None
            if content:
                parts.append(str(content))
    finally:
        if hasattr(stream, "close"):
            stream.close()
    return "".join(parts).strip()


def call_deepseek(client: OpenAI, model: str, system_prompt: str, dialogue: str) -> tuple[str, float]:
    user_message = (
        "请根据系统提示词，对下面这段已经完成的用户与奈奈消息记录进行依恋度判定。"
        "只评估记录，不要与用户继续对话；严格遵守系统提示词要求的 JSON 输出格式。\n\n"
        f"已完成消息记录：\n{dialogue}"
    )
    request = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_message},
        ],
        "temperature": DEEPSEEK_TEMPERATURE,
        "max_tokens": DEEPSEEK_MAX_TOKENS,
        "stream": True,
        "stream_options": {"include_usage": True},
        "extra_body": DEEPSEEK_THINKING_BODY,
    }
    started = time.perf_counter()
    last_error: Exception | None = None
    for attempt in range(1, REQUEST_RETRIES + 1):
        try:
            stream = client.chat.completions.create(**request)
            text = extract_stream_content(stream)
            if not text:
                raise ValueError("DeepSeek 返回为空")
            return text, round((time.perf_counter() - started) * 1000, 2)
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt < REQUEST_RETRIES:
                time.sleep(RETRY_BACKOFF_SEC * attempt)
    raise RuntimeError(f"DeepSeek 调用失败: {type(last_error).__name__}: {last_error}") from last_error


def extract_json_object(text: str) -> dict[str, Any]:
    candidate = text.strip()
    if candidate.startswith("```"):
        candidate = re.sub(r"^```[\w-]*\s*", "", candidate, count=1)
        candidate = re.sub(r"\s*```$", "", candidate, count=1).strip()
    start, end = candidate.find("{"), candidate.rfind("}")
    if start >= 0 and end > start:
        candidate = candidate[start : end + 1]
    parsed = json.loads(candidate)
    if not isinstance(parsed, dict):
        raise ValueError("JSON 根节点不是对象")
    return parsed


def normalize_score(value: Any) -> float:
    if isinstance(value, bool):
        raise ValueError("评分不能是布尔值")
    score = float(value)
    if not 0 <= score <= 10:
        raise ValueError(f"评分超出 0-10 范围: {score}")
    return round(score, 2)


def parse_judge(text: str) -> dict[str, Any]:
    parsed = extract_json_object(text)
    score = parsed.get("score")
    if score is None:
        score = parsed.get("final_score")
    if score is None:
        raise ValueError("Qwen 输出缺少 score")
    return {
        "score": normalize_score(score),
        "evaluation": str(parsed.get("evaluation", "")).strip(),
        "problems": str(parsed.get("problems", "无")).strip() or "无",
        "suggestions": str(parsed.get("suggestions", "无")).strip() or "无",
    }


def parse_obtained(text: str) -> str:
    try:
        parsed = extract_json_object(text)
    except Exception:  # noqa: BLE001
        return "未知"
    value = parsed.get("Obtained", parsed.get("obtained"))
    if isinstance(value, bool):
        return "yes" if value else "no"
    value_text = str(value or "").strip().lower()
    if value_text in {"yes", "true", "是", "有"}:
        return "yes"
    if value_text in {"no", "false", "否", "无"}:
        return "no"
    return "未知"


def call_qwen(
    client: OpenAI,
    model: str,
    judge_prompt: str,
    evaluator_prompt: str,
    dialogue: str,
    deepseek_result: str,
) -> tuple[dict[str, Any], str, float]:
    user_message = (
        "产品：ai抱枕\n\n"
        "依恋度计算提示词：\n"
        f"{evaluator_prompt}\n\n"
        "已完成消息记录：\n"
        f"{dialogue}\n\n"
        "DeepSeek 的依恋度评判：\n"
        f"{deepseek_result}"
    )
    request = {
        "model": model,
        "messages": [
            {"role": "system", "content": judge_prompt},
            {"role": "user", "content": user_message},
        ],
        "temperature": QWEN_TEMPERATURE,
        "max_tokens": QWEN_MAX_TOKENS,
        "extra_body": QWEN_THINKING_BODY,
    }
    started = time.perf_counter()
    last_error: Exception | None = None
    for attempt in range(1, REQUEST_RETRIES + 1):
        try:
            response = client.chat.completions.create(**request)
            raw = extract_content(response)
            if not raw:
                raise ValueError("Qwen 返回为空")
            return parse_judge(raw), raw, round((time.perf_counter() - started) * 1000, 2)
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt < REQUEST_RETRIES:
                time.sleep(RETRY_BACKOFF_SEC * attempt)
    raise RuntimeError(f"Qwen 调用失败: {type(last_error).__name__}: {last_error}") from last_error


def set_cell(ws: Any, row: int, col: int, value: Any) -> None:
    ws.cell(row, col, value)


def populate_summary(ws: Any, input_path: Path, output_path: Path, prompt_path: Path, stats: dict[str, Any]) -> None:
    ws.delete_rows(1, ws.max_row)
    rows = [
        ("ai抱枕依恋度评测汇总", ""),
        ("输入文件", str(input_path)),
        ("输出文件", str(output_path)),
        ("DeepSeek提示词", str(prompt_path)),
        ("DeepSeek模型", stats["deepseek_model"]),
        ("Qwen模型", stats["qwen_model"]),
        ("总用例数", stats["total"]),
        ("DeepSeek成功数", stats["deepseek_success"]),
        ("DeepSeek失败数", stats["deepseek_failed"]),
        ("Qwen成功数", stats["qwen_success"]),
        ("Qwen失败数", stats["qwen_failed"]),
        ("有效评分数", len(stats["scores"])),
        ("Qwen平均分", round(sum(stats["scores"]) / len(stats["scores"]), 2) if stats["scores"] else "无"),
        ("Qwen最低分", min(stats["scores"]) if stats["scores"] else "无"),
        ("Qwen最高分", max(stats["scores"]) if stats["scores"] else "无"),
        ("DeepSeek Obtained=yes", stats["obtained_yes"]),
        ("DeepSeek Obtained=no", stats["obtained_no"]),
        ("生成时间", datetime.now().isoformat(timespec="seconds")),
    ]
    for row_idx, (key, value) in enumerate(rows, start=1):
        ws.cell(row_idx, 1, key)
        ws.cell(row_idx, 2, value)
    ws.cell(20, 1, "评分分布")
    ws.cell(20, 2, "数量")
    for score in range(11):
        ws.cell(21 + score, 1, score)
        ws.cell(21 + score, 2, stats["score_distribution"].get(score, 0))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws["A1"].font = Font(bold=True, size=14)


def process(args: argparse.Namespace) -> Path:
    input_path = resolve_path(args.input)
    if not input_path.exists():
        raise FileNotFoundError(f"输入 Excel 不存在: {input_path}")
    prompt_path = resolve_path(args.prompt_path)
    evaluator_prompt = read_text(prompt_path)
    judge_prompt = load_judge_prompt(args.judge_prompt_path)
    output_path = resolve_path(args.output) if args.output else default_output_path(input_path)

    if args.resume:
        if not output_path.exists():
            raise FileNotFoundError(f"--resume 指定的输出文件不存在: {output_path}")
        workbook = load_workbook(output_path)
    else:
        workbook = load_workbook(input_path)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("输入工作簿缺少 Sheet1")
    source_ws = workbook["Sheet1"]
    headers, rows = validate_input(source_ws)
    result_headers = append_result_columns(source_ws)
    if "Sheet2" not in workbook.sheetnames:
        workbook.create_sheet("Sheet2")
    summary_ws = workbook["Sheet2"]

    if args.limit is not None:
        if args.limit <= 0:
            raise ValueError("--limit 必须为正数")
        rows = rows[: args.limit]
    rows = [row for row in rows if row >= args.start_row]
    if args.dry_run:
        print(f"输入校验通过: {input_path}")
        print(f"提示词: {prompt_path} ({len(evaluator_prompt)} chars)")
        print(f"待处理用例: {len(rows)}")
        return output_path

    deepseek_client = build_client(MODEL_CONFIG.used_api_key, MODEL_CONFIG.used_base_url)
    qwen_client = build_client(MODEL_CONFIG.judge_api_key, MODEL_CONFIG.judge_base_url)
    case_col = headers["用例编号"]
    dialogue_col = headers["用户对话"]
    stats: dict[str, Any] = {
        "total": len(rows),
        "deepseek_success": 0,
        "deepseek_failed": 0,
        "qwen_success": 0,
        "qwen_failed": 0,
        "scores": [],
        "score_distribution": {},
        "obtained_yes": 0,
        "obtained_no": 0,
        "deepseek_model": args.deepseek_model,
        "qwen_model": args.qwen_model,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    for row in rows:
        case_id = str(source_ws.cell(row, case_col).value or "").strip()
        dialogue = str(source_ws.cell(row, dialogue_col).value or "").strip()
        status_col = result_headers["DeepSeek状态"]
        qwen_status_col = result_headers["Qwen状态"]
        deepseek_already_done = args.resume and source_ws.cell(row, status_col).value == "成功"
        qwen_already_done = args.resume and source_ws.cell(row, qwen_status_col).value == "成功"
        if deepseek_already_done:
            existing_obtained = str(source_ws.cell(row, result_headers["DeepSeekObtained"]).value or "未知")
            stats["deepseek_success"] += 1
            if existing_obtained == "yes":
                stats["obtained_yes"] += 1
            elif existing_obtained == "no":
                stats["obtained_no"] += 1
        if qwen_already_done:
            try:
                score = normalize_score(source_ws.cell(row, result_headers["Qwen评分"]).value)
                stats["scores"].append(score)
                stats["score_distribution"][int(score)] = stats["score_distribution"].get(int(score), 0) + 1
                stats["qwen_success"] += 1
            except Exception:  # noqa: BLE001
                pass
            if deepseek_already_done:
                continue

        deepseek_text = str(source_ws.cell(row, result_headers["DeepSeek依恋度评判"]).value or "").strip()
        if deepseek_already_done and deepseek_text:
            print(f"[{case_id}] reuse existing DeepSeek judgement", flush=True)
        else:
            deepseek_text = ""

        if not deepseek_text:
            print(f"[{case_id}] DeepSeek evaluating...", flush=True)
            try:
                deepseek_text, deepseek_ms = call_deepseek(
                    deepseek_client, args.deepseek_model, evaluator_prompt, dialogue
                )
                set_cell(source_ws, row, result_headers["DeepSeek依恋度评判"], deepseek_text)
                set_cell(source_ws, row, result_headers["DeepSeekObtained"], parse_obtained(deepseek_text))
                set_cell(source_ws, row, result_headers["DeepSeek状态"], "成功")
                set_cell(source_ws, row, result_headers["DeepSeek耗时(ms)"], deepseek_ms)
                stats["deepseek_success"] += 1
                obtained = parse_obtained(deepseek_text)
                if obtained == "yes":
                    stats["obtained_yes"] += 1
                elif obtained == "no":
                    stats["obtained_no"] += 1
            except Exception as exc:  # noqa: BLE001
                set_cell(source_ws, row, result_headers["DeepSeek状态"], f"失败: {type(exc).__name__}: {exc}")
                stats["deepseek_failed"] += 1
                workbook.save(output_path)
                continue

        print(f"[{case_id}] Qwen judging...", flush=True)
        try:
            judge, raw_qwen, qwen_ms = call_qwen(
                qwen_client, args.qwen_model, judge_prompt, evaluator_prompt, dialogue, deepseek_text
            )
            set_cell(source_ws, row, result_headers["Qwen评分"], judge["score"])
            set_cell(source_ws, row, result_headers["Qwen评价"], judge["evaluation"])
            set_cell(source_ws, row, result_headers["Qwen问题"], judge["problems"])
            set_cell(source_ws, row, result_headers["Qwen改进建议"], judge["suggestions"])
            set_cell(source_ws, row, result_headers["Qwen状态"], "成功")
            set_cell(source_ws, row, result_headers["Qwen耗时(ms)"], qwen_ms)
            set_cell(source_ws, row, result_headers["Qwen原始输出"], raw_qwen)
            score = judge["score"]
            stats["scores"].append(score)
            stats["score_distribution"][int(score)] = stats["score_distribution"].get(int(score), 0) + 1
            stats["qwen_success"] += 1
        except Exception as exc:  # noqa: BLE001
            set_cell(source_ws, row, result_headers["Qwen状态"], f"失败: {type(exc).__name__}: {exc}")
            stats["qwen_failed"] += 1
        workbook.save(output_path)

    populate_summary(summary_ws, input_path, output_path, prompt_path, stats)
    workbook.save(output_path)
    print(f"结果已写入: {output_path}")
    return output_path


def main() -> int:
    args = parse_args()
    try:
        process(args)
        return 0
    except Exception as exc:  # noqa: BLE001
        print(f"[ERROR] {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
