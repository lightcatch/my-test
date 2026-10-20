from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_ROOT = Path("D:/PyCode/test/outputs/语小萌豆包测试结果")
DIMENSIONS = ["角色扮演吸引力", "角色一致性", "对话能力"]
SUMMARY_HEADERS = [
    "分组",
    "会话数",
    "角色扮演吸引力均分",
    "角色一致性均分",
    "对话能力均分",
    "综合均分",
    "最强维度",
    "最弱维度",
    "低分句数",
    "低分主因TOP3",
]
SUBSET_ORDER = {
    "2-3女": 1,
    "2-3男": 2,
    "3-6女": 3,
    "3-6男": 4,
    "6-8女": 5,
    "6-8男": 6,
    "总体": 99,
}

CATEGORY_PATTERNS = {
    "技术术语/超龄知识": {
        "硬件/编程术语": [
            "refactor",
            "bug",
            "cpu",
            "gpu",
            "rtx",
            "intel",
            "atx",
            "u盘",
            "显卡",
            "显存",
            "帧率",
            "主板",
            "双系统",
            "小程序",
            "硬件",
            "参数",
            "装机",
            "代码",
            "程序",
            "蓝牙",
            "模块",
            "接口",
            "风扇",
            "测温",
            "机箱",
            "电脑城",
            "供电",
            "双通道",
            "对照组",
            "测量",
            "实验",
        ],
        "成人化知识科普": [
            "专业知识",
            "超龄",
            "成人化",
            "儿童认知边界",
            "低龄儿童",
            "科普",
            "知识储备",
            "知识幻觉",
            "现代硬件参数",
        ],
        "未回到语言表达主线": [
            "语言学习",
            "语言表达",
            "核心使命",
            "未引导",
            "强化不当知识领域",
            "偏离使命",
            "词汇扩展",
        ],
    },
    "系统痕迹/格式不合规": {
        "代码/函数痕迹": [
            "achievement_unlocked",
            "函数调用",
            "系统函数",
            "后台",
            "程序指令",
            "代码痕迹",
            "technical",
            "premature",
        ],
        "括号/格式外露": [
            "括号",
            "格式不合规",
            "技术性表述",
            "剥离技术标识",
        ],
        "旁白/动作描写": [
            "旁白",
            "动作描写",
            "神态描写",
            "我笑眯眯地",
            "接过",
            "蹲着爬",
        ],
    },
    "人设边界/世界观偏移": {
        "森林背景脱锚": [
            "阿巴森林",
            "森林元素",
            "背景元素",
            "世界观",
            "背景无关",
            "角色锚点",
            "角色灵魂",
            "语言灵使",
            "森林灵使",
            "魔法视角",
            "守护者",
        ],
        "过度共演/越界跟随": [
            "我就在这里等你",
            "跟着你到楼下",
            "帮我叠衣服",
            "去贴翅膀",
            "拉钩拉钩",
            "回沙子城堡",
            "我来帮你",
            "过度共演",
            "越界",
        ],
        "伙伴视角偏移": [
            "伙伴平等",
            "平等视角",
            "监护人语气",
            "自我限制",
            "坦诚限制",
            "沉浸感",
            "伙伴共情",
            "支持者",
        ],
    },
    "提问模板化/重复确认": {
        "重复已知信息": [
            "重复",
            "已明确",
            "已答",
            "已知",
            "冗余确认",
            "重复确认",
            "重复提问",
            "信息重复",
        ],
        "固定句式追问": [
            "模板化",
            "固定句式",
            "程式化",
            "提问模式较常规",
            "什么颜色",
            "是不是",
            "好不好",
        ],
        "推进力度偏弱": [
            "推进新话题",
            "逻辑延伸性微弱",
            "互动引导性弱",
            "联想深度可加强",
        ],
    },
    "创意延展不足": {
        "缺少感官联想": [
            "感官",
            "触感",
            "温度",
            "画面感",
            "具象化",
            "感受型提问",
            "触觉",
            "温度细节",
        ],
        "想象力/故事延展不足": [
            "创意",
            "想象力",
            "故事接龙",
            "联想",
            "延伸",
            "深化",
            "拓展",
            "趣味知识",
            "创编",
            "互动深度",
        ],
        "表达变化不足": [
            "词汇",
            "修辞",
            "童谣",
            "记忆点",
            "表达多样性",
            "韵律",
            "生活细节",
        ],
    },
    "情绪承接不足": {
        "共情不足": [
            "共情",
            "情绪优先",
            "情绪确认",
            "情感优先",
            "委屈",
            "安抚",
            "被看见",
        ],
        "情感深化不足": [
            "情感深化",
            "情感价值",
            "情感联结",
            "承接用户情绪",
            "先共情",
            "后续提问未延续",
        ],
        "敏感点预判不足": [
            "敏感点",
            "怕怕",
            "好可怕",
            "预判",
            "脆弱感",
        ],
    },
    "价值导向/安全表达偏差": {
        "不当比较/偏见": [
            "比妈妈好用",
            "贬低",
            "偏见",
            "中立鼓励",
            "不当比较",
            "羡慕",
            "所有小朋友都比不过",
        ],
        "说教感": [
            "说教",
            "监护人语气",
            "规则转化",
            "成人化口吻",
        ],
        "安全提醒生硬": [
            "安全提醒",
            "不能",
            "危险引导",
            "吓他一大跳",
            "小心别摔倒",
        ],
    },
    "工具使用/流程时机不当": {
        "工具触发随意": [
            "工具使用随意",
            "工具箱",
            "视觉游戏",
            "指定游戏",
            "工具使用",
        ],
        "工具触发时机": [
            "触发逻辑",
            "触发条件",
            "低频触发",
            "触发时机",
            "用户完成结尾再授予",
        ],
        "称号授予过早": [
            "称号",
            "故事小作家",
            "接龙高手",
            "premature",
            "过早",
        ],
    },
}

CATEGORY_DETAIL_TEMPLATES = {
    "创意延展不足": "这类低分通常不是基础回应错误，而是回复停留在表层承接，没有继续补充更贴合上下文的感官细节、想象联想或故事推进，所以对话会显得稳但不够生动。",
    "技术术语/超龄知识": "这类低分说明回复直接引入了硬件、编程或成人化知识，超出了角色年龄与陪伴语境，也容易把原本的语言互动拉向技术说明。",
    "人设边界/世界观偏移": "这类低分主要出在角色站位不稳，回复脱离了“语小萌”应有的森林伙伴感，或者出现越界共演、世界观脱锚、陪伴视角失衡的问题。",
    "提问模板化/重复确认": "这类低分说明回复虽然基本承接了用户内容，但追问方式偏固定，或者重复确认用户已经给出的信息，导致推进感和新鲜度不足。",
    "情绪承接不足": "这类低分集中在情绪回应层面，表现为先讲知识或工具、后接情绪，或者对委屈、害怕、期待等状态承接不够深，温度感略弱。",
    "系统痕迹/格式不合规": "这类低分反映的是输出形态问题，例如函数名、系统标识、括号痕迹或旁白动作暴露在对话里，会直接破坏角色沉浸感。",
    "价值导向/安全表达偏差": "这类低分主要与表达导向有关，包括不当比较、轻微说教或安全提醒方式偏硬，容易削弱伙伴式、支持式的交流氛围。",
    "工具使用/流程时机不当": "这类低分通常出在工具或流程插入点上，表现为称号授予过早、工具触发过于突兀，或者把系统流程感直接带进了对话正文。",
    "其他": "这类低分分布较散，通常不属于高频主问题，但仍说明局部回复存在表达、承接或角色边界上的不稳定点。",
}

FALLBACK_CATEGORY = "其他"
REASON_META_KEYWORDS = [
    "不应",
    "未",
    "违背",
    "偏离",
    "削弱",
    "破坏",
    "错失",
    "冗余",
    "重复",
    "稍显",
    "缺乏",
    "超出",
    "过度",
    "不当",
    "模板化",
    "生硬",
    "抽象",
    "成人化",
    "知识幻觉",
    "世界观",
    "角色",
    "提问",
    "情绪",
    "共情",
    "工具",
    "设定",
    "人设",
    "沉浸感",
    "低龄",
    "儿童",
    "逻辑延伸",
    "安全提醒",
    "监护人",
    "价值观",
]
EMPTY_LOW_SCORE_MARKERS = [
    "无低分回复",
    "无显著低分回复",
    "无实际低分回复",
    "无低分问题",
    "无扣分项",
    "高分示例",
    "无符合低分标准",
    "高质回复示例",
    "满分评价",
]
SUMMARY_COL_COUNT = len(SUMMARY_HEADERS)


def clean_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def to_number(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def detect_subset(file_name: str) -> str:
    match = re.search(r"(2-3|3-6|6-8)\s*[-_ ]?\s*([男女])", file_name)
    if match:
        return f"{match.group(1)}{match.group(2)}"
    return "未知"


def is_empty_low_score(text: str) -> bool:
    if not text:
        return True
    normalized = text.strip()
    if normalized in {"无", "None"}:
        return True
    return any(marker in normalized for marker in EMPTY_LOW_SCORE_MARKERS)


def strip_sentence_prefix(text: str) -> str:
    text = clean_text(text)
    text = re.sub(r"^第\d+轮assistant回复[:：]\s*", "", text, flags=re.I)
    text = re.sub(r"^\[\d+\]\s*assistant[:：]\s*", "", text, flags=re.I)
    text = re.sub(r"^assistant[:：]\s*", "", text, flags=re.I)
    text = re.sub(r"^\[\d+\]\s*", "", text)
    return text.strip(' "“”')


def extract_parenthetical_reason(text: str) -> tuple[str, str | None]:
    text = clean_text(text)
    matches = list(re.finditer(r"[（(]([^（）()]*)[）)]", text))
    if not matches:
        return text, None
    last = matches[-1]
    before = text[: last.start()].strip(" ，,;；")
    reason = last.group(1).strip()
    if before:
        return before, reason
    return text, None


def looks_like_reason(text: str) -> bool:
    text = clean_text(text)
    if not text:
        return False
    if "assistant:" in text.lower():
        return False

    score = 0
    for keyword in REASON_META_KEYWORDS:
        if keyword in text:
            score += 1
    if re.search(r"(角色|用户|提问|知识|工具|设定|人设|对话|表达|安全|话题|世界观|情绪|共情|童趣|模板化|沉浸感)", text):
        score += 2
    if re.search(r"(不应|违背|偏离|削弱|破坏|错失|冗余|重复|稍显|缺乏|超出|过度|不当)", text):
        score += 2
    if text.startswith(("角色", "提问", "用户", "知识", "安全", "共情", "表达")):
        score += 1
    if len(text) <= 80 and score >= 2:
        return True
    return score >= 4


def split_low_items(low_text: str) -> list[dict[str, str | None]]:
    text = clean_text(low_text)
    if is_empty_low_score(text):
        return []

    parts = [clean_text(part) for part in re.split(r"\*\*+", text) if clean_text(part)]
    if not parts:
        return []

    items: list[dict[str, str | None]] = []
    for part in parts:
        sentence_candidate, paren_reason = extract_parenthetical_reason(part)
        sentence_candidate = strip_sentence_prefix(sentence_candidate)

        if paren_reason:
            if sentence_candidate:
                items.append({"sentence": sentence_candidate, "explicit_reason": paren_reason})
            continue

        if looks_like_reason(part):
            if items:
                existing = items[-1].get("explicit_reason")
                items[-1]["explicit_reason"] = part if not existing else f"{existing}；{part}"
            continue

        if sentence_candidate:
            items.append({"sentence": sentence_candidate, "explicit_reason": None})

    if not items and not looks_like_reason(text):
        items.append({"sentence": strip_sentence_prefix(text), "explicit_reason": None})

    cleaned_items: list[dict[str, str | None]] = []
    for item in items:
        sentence = strip_sentence_prefix(item.get("sentence") or "")
        reason = clean_text(item.get("explicit_reason")) or None
        if not sentence:
            continue
        if is_empty_low_score(sentence):
            continue
        cleaned_items.append({"sentence": sentence, "explicit_reason": reason})
    return cleaned_items


def sentence_split(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    parts = re.split(r"(?<=[。！？；])\s*", text)
    return [part.strip() for part in parts if part.strip()]


def extract_issue_sentences(text: str) -> list[str]:
    markers = [
        "不足在于",
        "细微不足",
        "唯一不足",
        "可优化处",
        "问题：",
        "问题:",
        "未",
        "重复",
        "不应",
        "偏离",
        "破坏",
        "削弱",
        "模板化",
        "超出",
    ]
    return [sentence for sentence in sentence_split(text) if any(marker in sentence for marker in markers)]


def keyword_score(text: str, keyword: str) -> int:
    return text.lower().count(keyword.lower())


def flatten_keywords() -> dict[str, list[str]]:
    flat_map: dict[str, list[str]] = {}
    for category, sub_map in CATEGORY_PATTERNS.items():
        keywords: list[str] = []
        for sub_keywords in sub_map.values():
            keywords.extend(sub_keywords)
        flat_map[category] = keywords
    return flat_map


FLAT_CATEGORY_KEYWORDS = flatten_keywords()


def basis_quality(text: str, category: str) -> float:
    text = clean_text(text)
    if not text or text == "-":
        return -100.0

    issue_markers = ["未", "重复", "不应", "超出", "偏离", "破坏", "削弱", "模板化", "缺失", "问题", "违背", "沉浸感", "技术", "术语", "角色", "世界观", "工具", "提问", "共情", "成人化", "突兀", "生硬"]
    praise_markers = ["太厉害", "太棒", "太有意思", "真棒", "好聪明", "好有趣", "真了不起", "真精彩", "太周到", "太暖心"]

    score = 0.0
    if looks_like_reason(text):
        score += 12.0

    keyword_hits = sum(keyword_score(text, keyword) for keyword in FLAT_CATEGORY_KEYWORDS.get(category, []))
    score += min(18.0, keyword_hits * 3.0)

    length = len(text)
    score += max(-10.0, 10.0 - abs(length - 28) * 0.25)

    if any(marker in text for marker in ["建议", "改为"] + issue_markers):
        score += 8.0
    if not any(marker in text for marker in issue_markers) and any(marker in text for marker in praise_markers):
        score -= 25.0
    if text.startswith(("表达多样性上", "沟通技巧层面", "对话全程", "本次对话中", "人类相似性上", "角色灵魂鲜活", "整体完美融合")):
        score -= 20.0

    if text in {"故事小作家", "接龙高手", "小小表演家", "分类小达人"}:
        score -= 30.0
    if "achievement_unlocked" in text.lower():
        score -= 15.0
    if re.fullmatch(r"[\w\u4e00-\u9fff]{1,8}", text):
        score -= 10.0
    return score


def classify_reason(
    sentence: str,
    explicit_reason: str | None,
    logic: str,
    improvement: str,
) -> tuple[str, str | None]:
    combined = " ".join(filter(None, [sentence, explicit_reason or "", logic, improvement]))
    scores: dict[str, float] = {}
    subcategory_choice: dict[str, tuple[str | None, float]] = {}

    for category, sub_map in CATEGORY_PATTERNS.items():
        total = 0.0
        best_sub = (None, 0.0)
        for subcategory, keywords in sub_map.items():
            sub_score = 0.0
            for keyword in keywords:
                count = keyword_score(combined, keyword)
                if count:
                    sub_score += float(count)
                    if explicit_reason and keyword.lower() in explicit_reason.lower():
                        sub_score += 0.5
            if sub_score > best_sub[1]:
                best_sub = (subcategory, sub_score)
            total += sub_score
        scores[category] = total
        subcategory_choice[category] = best_sub

    sentence_lower = sentence.lower()
    if any(keyword in sentence_lower for keyword in ["achievement_unlocked", "()"]):
        scores["系统痕迹/格式不合规"] += 3
    if any(keyword in sentence_lower for keyword in ["显卡", "cpu", "gpu", "rtx", "装机", "代码", "参数", "主板", "u盘", "bug", "refactor"]):
        scores["技术术语/超龄知识"] += 3
    if any(keyword in sentence for keyword in ["什么颜色", "是不是", "好不好"]) and (
        "已" in (explicit_reason or "") or "重复" in (explicit_reason or "") or "模板化" in (explicit_reason or "")
    ):
        scores["提问模板化/重复确认"] += 2
    if any(keyword in sentence for keyword in ["等你", "跟着你", "楼下", "拉钩", "帮我叠衣服", "站岗"]):
        scores["人设边界/世界观偏移"] += 2
    if "故事小作家" in sentence and "achievement_unlocked" in sentence_lower:
        scores["工具使用/流程时机不当"] += 1.5

    ranked = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    main_category = ranked[0][0] if ranked and ranked[0][1] > 0 else FALLBACK_CATEGORY
    secondary_category = None
    if len(ranked) > 1 and ranked[1][1] > 0:
        secondary_category = ranked[1][0]

    best_subcategory = None
    if main_category in subcategory_choice and subcategory_choice[main_category][1] > 0:
        best_subcategory = subcategory_choice[main_category][0]
    return main_category, best_subcategory or secondary_category


def choose_basis(
    category: str,
    explicit_reason: str | None,
    sentence: str,
    logic: str,
    improvement: str,
) -> tuple[str, str]:
    if explicit_reason:
        return clean_text(explicit_reason)[:160], "显式"

    keywords = FLAT_CATEGORY_KEYWORDS.get(category, [])
    candidates = (
        extract_issue_sentences(improvement)
        + extract_issue_sentences(logic)
        + sentence_split(improvement)
        + sentence_split(logic)
        + sentence_split(sentence)
    )
    best_candidate = ""
    best_score = float("-inf")
    for candidate in candidates:
        score = sum(keyword_score(candidate, keyword) for keyword in keywords) * 3
        score += 6 if looks_like_reason(candidate) else 0
        score += basis_quality(candidate, category)
        if score > best_score:
            best_candidate = candidate
            best_score = score

    if not best_candidate:
        best_candidate = candidates[0] if candidates else sentence[:160]
    return clean_text(best_candidate)[:160], "推断" if best_candidate else "其他"


def top_categories(rows: list[dict[str, object]], limit: int = 3) -> list[tuple[str, int]]:
    counter = Counter(str(row["主因分类"]) for row in rows if row.get("主因分类"))
    return counter.most_common(limit)


def top_categories_text(rows: list[dict[str, object]], limit: int = 3) -> str:
    categories = top_categories(rows, limit=limit)
    if not categories:
        return "-"
    return "；".join(f"{name}({count})" for name, count in categories)


def truncate_text(text: str, limit: int) -> str:
    text = clean_text(text)
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def example_priority(row: dict[str, object]) -> float:
    category = str(row["主因分类"])
    score = 0.0
    reason_quality = basis_quality(str(row["判定依据摘要"]), category)
    if row["依据来源"] == "显式" and reason_quality > 0:
        score += 40.0
    score += reason_quality * 2
    sentence_len = len(str(row["低分句"]))
    if sentence_len < 40:
        score -= 30.0
    score += max(0.0, 40.0 - abs(sentence_len - 100) * 0.4)
    basis_len = len(str(row["判定依据摘要"]))
    score += max(0.0, 20.0 - abs(basis_len - 32) * 0.5)
    if any(keyword in str(row["判定依据摘要"]) for keyword in ["未", "重复", "超出", "不应", "模板化", "世界观", "工具", "情绪"]):
        score += 5.0
    return score


def choose_typical_basis(rows: list[dict[str, object]]) -> str:
    if not rows:
        return "-"
    category = str(rows[0]["主因分类"])
    explicit_rows = [
        row for row in rows if row["依据来源"] == "显式" and basis_quality(str(row["判定依据摘要"]), category) > 0
    ]
    pool = explicit_rows or rows
    best_row = max(
        pool,
        key=lambda row: (
            1 if row["依据来源"] == "显式" and basis_quality(str(row["判定依据摘要"]), category) > 0 else 0,
            example_priority(row),
            basis_quality(str(row["判定依据摘要"]), category),
        ),
    )
    return truncate_text(str(best_row["判定依据摘要"]), 80)


def choose_examples(rows: list[dict[str, object]], count: int = 2) -> list[dict[str, object]]:
    if not rows:
        return []

    ranked = sorted(
        rows,
        key=lambda row: (
            1 if row["依据来源"] == "显式" else 0,
            example_priority(row),
            -len(str(row["低分句"])),
        ),
        reverse=True,
    )

    selected: list[dict[str, object]] = []
    used_subsets: set[str] = set()
    for candidate in ranked:
        subset = str(candidate["子集"])
        if subset in used_subsets:
            continue
        selected.append(candidate)
        used_subsets.add(subset)
        if len(selected) == count:
            return selected

    for candidate in ranked:
        if candidate in selected:
            continue
        selected.append(candidate)
        if len(selected) == count:
            return selected
    return selected


def format_example(row: dict[str, object]) -> str:
    prefix = f"【{row['子集']}·{row['维度']}】"
    sentence = truncate_text(str(row["低分句"]), 110)
    basis = truncate_text(str(row["判定依据摘要"]), 32)
    return f"{prefix}{sentence}（原因：{basis}）"


def build_category_detail_description(category: str, count: int, dim_counter: Counter[str]) -> str:
    base = CATEGORY_DETAIL_TEMPLATES.get(category, CATEGORY_DETAIL_TEMPLATES["其他"])
    top_dims = dim_counter.most_common(2)
    if not top_dims:
        return base
    dim_summary = "；".join(f"{dim}({dim_count})" for dim, dim_count in top_dims)
    primary_dim = top_dims[0][0]
    return f"{base} 在当前模型中共出现 {count} 次，主要集中在 {dim_summary}，其中 {primary_dim} 最突出。"


def build_detail_blocks(folder_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    blocks: list[dict[str, object]] = []
    for category, count in top_categories(folder_rows, limit=3):
        category_rows = [row for row in folder_rows if row["主因分类"] == category]
        dim_counter = Counter(str(row["维度"]) for row in category_rows)
        dim_summary = "；".join(f"{dim}({dim_count})" for dim, dim_count in dim_counter.most_common(2))
        examples = choose_examples(category_rows, count=2)
        formatted_examples = [format_example(row) for row in examples]
        while len(formatted_examples) < 2:
            formatted_examples.append("样本不足")
        blocks.append(
            {
                "主因": category,
                "出现次数": count,
                "主要影响维度": dim_summary or "-",
                "详细说明": build_category_detail_description(category, count, dim_counter),
                "典型判定依据": choose_typical_basis(category_rows),
                "原文示例1": formatted_examples[0],
                "原文示例2": formatted_examples[1],
            }
        )
    return blocks


def auto_width(ws, max_width: int = 42) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), max_width)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_sheet(ws, freeze: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    auto_width(ws)


def safe_sheet_name(name: str) -> str:
    invalid = set("[]:*?/\\")
    name = "".join("_" if char in invalid else char for char in name)
    return name[:31]


def merge_summary_value_row(ws, row_idx: int, label: str, value: object) -> None:
    ws.cell(row=row_idx, column=1, value=label)
    ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=SUMMARY_COL_COUNT)
    ws.cell(row=row_idx, column=2, value=value)


def style_summary_sheet(ws, detail_title_row: int, table_header_row: int) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="EAF4E2")
    label_fill = PatternFill("solid", fgColor="F6F8FB")
    bold_font = Font(bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[detail_title_row]:
        cell.fill = section_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(detail_title_row + 1, table_header_row):
        label_cell = ws.cell(row=row_idx, column=1)
        if label_cell.value:
            label_cell.fill = label_fill
            label_cell.font = bold_font
        for col_idx in range(1, SUMMARY_COL_COUNT + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(vertical="top", wrap_text=True)

    for cell in ws[table_header_row]:
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = f"A{table_header_row + 1}"
    ws.column_dimensions["A"].width = 18
    for col_idx in range(2, SUMMARY_COL_COUNT + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16
    ws.row_dimensions[detail_title_row].height = 24


def write_summary_sheet(
    ws,
    insights: list[str],
    detail_blocks: list[dict[str, object]],
    summary_rows: list[dict[str, object]],
) -> None:
    ws.append(["摘要项", "内容"])
    for idx, insight in enumerate(insights, start=1):
        ws.append([f"结论{idx}", insight])
    ws.append([None] * SUMMARY_COL_COUNT)

    detail_title_row = ws.max_row + 1
    ws.merge_cells(start_row=detail_title_row, start_column=1, end_row=detail_title_row, end_column=SUMMARY_COL_COUNT)
    ws.cell(row=detail_title_row, column=1, value="低分主因详细展开")

    current_row = detail_title_row + 1
    field_order = ["主因", "出现次数", "主要影响维度", "详细说明", "典型判定依据", "原文示例1", "原文示例2"]
    for block in detail_blocks:
        for field in field_order:
            merge_summary_value_row(ws, current_row, field, block[field])
            current_row += 1
        ws.append([None] * SUMMARY_COL_COUNT)
        current_row += 1

    table_header_row = ws.max_row + 1
    ws.append(SUMMARY_HEADERS)
    for row in summary_rows:
        ws.append([row[col] for col in SUMMARY_HEADERS])

    style_summary_sheet(ws, detail_title_row=detail_title_row, table_header_row=table_header_row)


def output_filename_for_folder(folder_label: str) -> str:
    if "1.5" in folder_label:
        return "doubao1.5.xlsx"
    if "2.0" in folder_label:
        return "doubao2.0.xlsx"
    lower_label = folder_label.lower()
    if "千问" in folder_label or "qwen" in lower_label:
        return "qwen.xlsx"
    if "deepseek" in lower_label:
        return "deepseek.xlsx"
    return safe_sheet_name(folder_label) + ".xlsx"


def collect_model_groups(root: Path) -> list[tuple[str, list[Path]]]:
    direct_files = sorted(root.glob("qwen_score*.xlsx"))
    sub_groups: list[tuple[str, list[Path]]] = []
    for path in sorted(root.iterdir()):
        if not path.is_dir():
            continue
        files = sorted(path.glob("qwen_score*.xlsx"))
        if files:
            sub_groups.append((path.name, files))

    if direct_files and not sub_groups:
        return [(root.name, direct_files)]
    if sub_groups and not direct_files:
        if len(sub_groups) == 1:
            return [(root.name, sub_groups[0][1])]
        return sub_groups

    groups = list(sub_groups)
    if direct_files:
        groups.insert(0, (root.name, direct_files))
    return groups


def build_folder_workbook(
    folder_label: str,
    folder_summary_rows: dict[str, list[dict[str, object]]],
    folder_insights: dict[str, list[str]],
    folder_detail_blocks: dict[str, list[dict[str, object]]],
    detail_rows: list[dict[str, object]],
    summary_group_rows: list[dict[str, object]],
    category_definition_rows: list[dict[str, object]],
) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    ws_summary_sheet = wb.create_sheet("总结")
    write_summary_sheet(
        ws=ws_summary_sheet,
        insights=folder_insights[folder_label],
        detail_blocks=folder_detail_blocks[folder_label],
        summary_rows=folder_summary_rows[folder_label],
    )

    folder_detail_rows = [row for row in detail_rows if row["文件夹"] == folder_label]
    ws_detail = wb.create_sheet("低分明细")
    detail_header = [
        "文件夹",
        "子集",
        "文件名",
        "session_id",
        "维度",
        "维度分数",
        "低分句",
        "主因分类",
        "次因分类",
        "判定依据摘要",
        "依据来源",
    ]
    ws_detail.append(detail_header)
    for row in folder_detail_rows:
        ws_detail.append([row[col] for col in detail_header])
    style_sheet(ws_detail)
    ws_detail.column_dimensions["C"].width = 36
    ws_detail.column_dimensions["G"].width = 60
    ws_detail.column_dimensions["J"].width = 58

    folder_group_rows = [row for row in summary_group_rows if row["文件夹"] == folder_label]
    ws_group = wb.create_sheet("低分归类汇总")
    group_header = [
        "文件夹",
        "维度",
        "主因分类",
        "低分句数",
        "占该维度低分句比例",
        "涉及session数",
        "代表子集",
        "代表句",
        "代表依据",
    ]
    ws_group.append(group_header)
    for row in folder_group_rows:
        ws_group.append([row[col] for col in group_header])
    style_sheet(ws_group)
    ws_group.column_dimensions["H"].width = 60
    ws_group.column_dimensions["I"].width = 58

    ws_rule = wb.create_sheet("分类口径")
    rule_header = ["主因分类", "说明"]
    ws_rule.append(rule_header)
    for row in category_definition_rows:
        ws_rule.append([row[col] for col in rule_header])
    style_sheet(ws_rule)
    ws_rule.column_dimensions["B"].width = 72

    return wb


def main(root: Path) -> None:
    folder_metrics: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "scores": defaultdict(list),
            "subset_scores": defaultdict(lambda: defaultdict(list)),
            "session_ids": set(),
            "subset_sessions": defaultdict(set),
        }
    )
    detail_rows: list[dict[str, object]] = []

    for folder_label, score_files in collect_model_groups(root):
        for file_path in score_files:
            subset = detect_subset(file_path.name)
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            if "summary" not in workbook.sheetnames:
                continue
            ws = workbook["summary"]
            headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            idx = {header: i for i, header in enumerate(headers)}

            for row in ws.iter_rows(min_row=2, values_only=True):
                if all(value is None for value in row):
                    continue

                session_id = row[idx["session_id"]]
                folder_metrics[folder_label]["session_ids"].add((file_path.name, session_id))
                folder_metrics[folder_label]["subset_sessions"][subset].add((file_path.name, session_id))

                for dimension in DIMENSIONS:
                    score_key = f"{dimension}_final_score"
                    low_key = f"{dimension}_low_score"
                    logic_key = f"{dimension}_logic_analysis"
                    improvement_key = f"{dimension}_improvement_suggestion"

                    score_value = to_number(row[idx[score_key]])
                    if score_value is not None:
                        folder_metrics[folder_label]["scores"][dimension].append(score_value)
                        folder_metrics[folder_label]["subset_scores"][subset][dimension].append(score_value)

                    low_text = clean_text(row[idx[low_key]])
                    logic = clean_text(row[idx[logic_key]])
                    improvement = clean_text(row[idx[improvement_key]])

                    for item in split_low_items(low_text):
                        sentence = clean_text(item["sentence"])
                        explicit_reason = clean_text(item.get("explicit_reason")) or None
                        main_category, secondary_category = classify_reason(sentence, explicit_reason, logic, improvement)
                        basis, basis_source = choose_basis(main_category, explicit_reason, sentence, logic, improvement)
                        detail_rows.append(
                            {
                                "文件夹": folder_label,
                                "子集": subset,
                                "文件名": file_path.name,
                                "session_id": session_id,
                                "维度": dimension,
                                "维度分数": score_value,
                                "低分句": sentence,
                                "主因分类": main_category,
                                "次因分类": secondary_category or "",
                                "判定依据摘要": basis,
                                "依据来源": basis_source,
                                "显式判定依据": explicit_reason or "",
                            }
                        )

    if not detail_rows:
        raise SystemExit("未提取到任何低分句，未生成报告。")

    folder_summary_rows: dict[str, list[dict[str, object]]] = defaultdict(list)
    folder_insights: dict[str, list[str]] = {}
    folder_detail_blocks: dict[str, list[dict[str, object]]] = {}

    for folder_label, metrics in folder_metrics.items():
        score_map = metrics["scores"]
        detail_in_folder = [row for row in detail_rows if row["文件夹"] == folder_label]
        subset_scores = metrics["subset_scores"]

        overall_dimension_avgs = {
            dim: round(sum(values) / len(values), 2) if values else None
            for dim, values in score_map.items()
        }
        overall_composite = round(sum(v for v in overall_dimension_avgs.values() if v is not None) / len(DIMENSIONS), 2)
        best_dim = max(overall_dimension_avgs.items(), key=lambda item: item[1])[0]
        worst_dim = min(overall_dimension_avgs.items(), key=lambda item: item[1])[0]

        for subset, dims in subset_scores.items():
            dimension_avgs = {
                dim: round(sum(values) / len(values), 2) if values else None
                for dim, values in dims.items()
            }
            composite = round(sum(v for v in dimension_avgs.values() if v is not None) / len(DIMENSIONS), 2)
            best_subset_dim = max(dimension_avgs.items(), key=lambda item: item[1])[0]
            worst_subset_dim = min(dimension_avgs.items(), key=lambda item: item[1])[0]
            subset_detail = [row for row in detail_in_folder if row["子集"] == subset]
            folder_summary_rows[folder_label].append(
                {
                    "分组": subset,
                    "会话数": len(metrics["subset_sessions"][subset]),
                    "角色扮演吸引力均分": dimension_avgs["角色扮演吸引力"],
                    "角色一致性均分": dimension_avgs["角色一致性"],
                    "对话能力均分": dimension_avgs["对话能力"],
                    "综合均分": composite,
                    "最强维度": best_subset_dim,
                    "最弱维度": worst_subset_dim,
                    "低分句数": len(subset_detail),
                    "低分主因TOP3": top_categories_text(subset_detail),
                }
            )

        folder_summary_rows[folder_label].append(
            {
                "分组": "总体",
                "会话数": len(metrics["session_ids"]),
                "角色扮演吸引力均分": overall_dimension_avgs["角色扮演吸引力"],
                "角色一致性均分": overall_dimension_avgs["角色一致性"],
                "对话能力均分": overall_dimension_avgs["对话能力"],
                "综合均分": overall_composite,
                "最强维度": best_dim,
                "最弱维度": worst_dim,
                "低分句数": len(detail_in_folder),
                "低分主因TOP3": top_categories_text(detail_in_folder),
            }
        )
        folder_summary_rows[folder_label].sort(key=lambda row: SUBSET_ORDER.get(str(row["分组"]), 999))
        folder_detail_blocks[folder_label] = build_detail_blocks(detail_in_folder)

        subset_rows = [row for row in folder_summary_rows[folder_label] if row["分组"] != "总体"]
        best_subset = max(subset_rows, key=lambda row: row["综合均分"])
        worst_subset = min(subset_rows, key=lambda row: row["综合均分"])
        reason_counter = Counter(row["主因分类"] for row in detail_in_folder)
        top_reason_lines = "；".join(f"{name}({count})" for name, count in reason_counter.most_common(3))
        folder_insights[folder_label] = [
            f"总体综合均分 {overall_composite}，其中 {best_dim} 最强（{overall_dimension_avgs[best_dim]}），{worst_dim} 最弱（{overall_dimension_avgs[worst_dim]}）。",
            f"表现最好子集为 {best_subset['分组']}（综合均分 {best_subset['综合均分']}），最弱子集为 {worst_subset['分组']}（综合均分 {worst_subset['综合均分']}）。",
            f"低分主因 Top3：{top_reason_lines if top_reason_lines else '-'}。",
        ]

    comparison_rows: list[dict[str, object]] = []
    for folder_label in sorted(folder_metrics):
        rows = folder_summary_rows[folder_label]
        overall = next(row for row in rows if row["分组"] == "总体")
        subset_rows = [row for row in rows if row["分组"] != "总体"]
        best_subset = max(subset_rows, key=lambda row: row["综合均分"])
        worst_subset = min(subset_rows, key=lambda row: row["综合均分"])
        comparison_rows.append(
            {
                "文件夹": folder_label,
                "会话数": overall["会话数"],
                "角色扮演吸引力均分": overall["角色扮演吸引力均分"],
                "角色一致性均分": overall["角色一致性均分"],
                "对话能力均分": overall["对话能力均分"],
                "综合均分": overall["综合均分"],
                "最好子集": f"{best_subset['分组']} ({best_subset['综合均分']})",
                "最弱子集": f"{worst_subset['分组']} ({worst_subset['综合均分']})",
                "低分句数": overall["低分句数"],
                "低分主因TOP3": overall["低分主因TOP3"],
            }
        )

    summary_group_rows: list[dict[str, object]] = []
    summary_groups: dict[tuple[str, str, str], list[dict[str, object]]] = defaultdict(list)
    for row in detail_rows:
        key = (str(row["文件夹"]), str(row["维度"]), str(row["主因分类"]))
        summary_groups[key].append(row)

    for (folder_label, dimension, main_category), rows in sorted(summary_groups.items()):
        session_count = len({(row["文件名"], row["session_id"]) for row in rows})
        representative = rows[0]
        same_dimension_total = len(
            [item for item in detail_rows if item["文件夹"] == folder_label and item["维度"] == dimension]
        )
        ratio = round(len(rows) / same_dimension_total, 4) if same_dimension_total else 0
        summary_group_rows.append(
            {
                "文件夹": folder_label,
                "维度": dimension,
                "主因分类": main_category,
                "低分句数": len(rows),
                "占该维度低分句比例": ratio,
                "涉及session数": session_count,
                "代表子集": representative["子集"],
                "代表句": representative["低分句"],
                "代表依据": representative["判定依据摘要"],
            }
        )

    category_definition_rows = [{"主因分类": category, "说明": "；".join(sub_map.keys())} for category, sub_map in CATEGORY_PATTERNS.items()]

    output_paths: list[Path] = []
    for folder_label in sorted(folder_summary_rows):
        wb = build_folder_workbook(
            folder_label=folder_label,
            folder_summary_rows=folder_summary_rows,
            folder_insights=folder_insights,
            folder_detail_blocks=folder_detail_blocks,
            detail_rows=detail_rows,
            summary_group_rows=summary_group_rows,
            category_definition_rows=category_definition_rows,
        )
        output_path = root / output_filename_for_folder(folder_label)
        wb.save(output_path)
        output_paths.append(output_path)

    for output_path in output_paths:
        print(f"OUTPUT={output_path}")
    for folder_label in sorted(folder_summary_rows):
        overall = next(row for row in folder_summary_rows[folder_label] if row["分组"] == "总体")
        print(
            f"{folder_label}: 综合均分={overall['综合均分']}, "
            f"角色扮演吸引力={overall['角色扮演吸引力均分']}, "
            f"角色一致性={overall['角色一致性均分']}, "
            f"对话能力={overall['对话能力均分']}, "
            f"低分主因TOP3={overall['低分主因TOP3']}"
        )
    print(f"LOW_DETAIL_ROWS={len(detail_rows)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build summary workbooks for 语小萌 score outputs.")
    parser.add_argument(
        "--root",
        default=str(DEFAULT_ROOT),
        help="Root folder that contains qwen_score files directly or inside model subfolders.",
    )
    args = parser.parse_args()
    main(Path(args.root))
