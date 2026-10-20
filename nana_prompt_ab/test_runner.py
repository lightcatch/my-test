from __future__ import annotations

import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from config import (
    NANA_DIALOG_PROMPT_A,
    NANA_DIALOG_PROMPT_B,
    NANA_SYSTEM_PROMPT_A,
    NANA_SYSTEM_PROMPT_B,
    OPENING_USER_MESSAGE,
    RunnerConfig,
    VARIANT_PROMPTS,
)
from prompts import PILLOW_VAR_KEYS, fill_pillow_prompt, prompt_hash
from runner import SingleRunnerLock, run_campaign_once


# 每个变量值都带唯一哨兵，便于断言它没有随历史滚进后续请求。
SENTINEL_VARIABLES = {
    "worldview": "SENTINEL_WORLDVIEW",
    "story": "SENTINEL_STORY",
    "memory_chapter": "SENTINEL_MEMORY",
    "event_section": "SENTINEL_EVENT",
    "profile_section": "SENTINEL_PROFILE",
    "weather": "SENTINEL_WEATHER",
    "Volume": "SENTINEL_VOLUME",
    "charging_status": "SENTINEL_CHARGING",
    "special_scenarios": "SENTINEL_SCENARIO",
    "battery_notification": "SENTINEL_BATTERY",
}
SENTINELS = tuple(SENTINEL_VARIABLES.values())


class FakeCompletions:
    def __init__(self, prefix: str):
        self.prefix = prefix
        self.requests = []

    def create(self, **request):
        self.requests.append(request)
        turn = sum(1 for message in request["messages"] if message["role"] == "user")
        content = f"语气平淡** {self.prefix}回复{turn}"
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])


class FakeClient:
    def __init__(self, prefix: str):
        self.chat = SimpleNamespace(completions=FakeCompletions(prefix))

    @property
    def requests(self):
        return self.chat.completions.requests


def fake_fetcher(sn: str, user_input: str) -> dict:
    assert sn, "sn must be forwarded to the variables API"
    assert user_input, "user input must be forwarded to the variables API"
    return dict(SENTINEL_VARIABLES)


class PromptStructureTests(unittest.TestCase):
    def test_variant_prompts_are_wired_to_their_own_constants(self):
        # A/B 的具体内容由使用者改 config.py 决定，这里只校验接线：
        # 每组都挂着自己的 sn、system、dialog，且内容非空。
        self.assertTrue(NANA_SYSTEM_PROMPT_A.strip())
        self.assertTrue(NANA_SYSTEM_PROMPT_B.strip())
        self.assertTrue(NANA_DIALOG_PROMPT_A.strip())
        self.assertTrue(NANA_DIALOG_PROMPT_B.strip())
        for variant, prompts in VARIANT_PROMPTS.items():
            self.assertEqual(
                sorted(prompts), ["dialog_prompt", "sn", "system_prompt"], variant
            )
            self.assertTrue(prompts["sn"].strip())
        self.assertIs(VARIANT_PROMPTS["A"]["system_prompt"], NANA_SYSTEM_PROMPT_A)
        self.assertIs(VARIANT_PROMPTS["B"]["system_prompt"], NANA_SYSTEM_PROMPT_B)

    def test_both_dialog_templates_keep_all_placeholders(self):
        for template in (NANA_DIALOG_PROMPT_A, NANA_DIALOG_PROMPT_B):
            for key in PILLOW_VAR_KEYS:
                self.assertIn("{{" + key + "}}", template, key)

    def test_prompt_hash_reacts_to_either_prompt(self):
        base = VARIANT_PROMPTS["A"]
        changed_dialog = {**base, "dialog_prompt": base["dialog_prompt"] + "x"}
        changed_system = {**base, "system_prompt": base["system_prompt"] + "x"}
        fingerprint = prompt_hash(base)
        self.assertEqual(fingerprint, prompt_hash(dict(base)))
        self.assertNotEqual(fingerprint, prompt_hash(changed_dialog))
        self.assertNotEqual(fingerprint, prompt_hash(changed_system))


class FillPromptTests(unittest.TestCase):
    def test_all_placeholders_are_replaced_and_user_input_appended(self):
        filled = fill_pillow_prompt(NANA_DIALOG_PROMPT_A, dict(SENTINEL_VARIABLES), "我回来了")
        self.assertNotIn("{{", filled)
        self.assertTrue(filled.rstrip().endswith("我回来了"))
        for sentinel in SENTINELS:
            self.assertIn(sentinel, filled)

    def test_missing_keys_become_empty_and_do_not_raise(self):
        filled = fill_pillow_prompt(NANA_DIALOG_PROMPT_A, {}, "在吗")
        self.assertNotIn("{{", filled)
        self.assertTrue(filled.rstrip().endswith("在吗"))

    def test_none_values_become_empty(self):
        variables = {key: None for key in PILLOW_VAR_KEYS}
        filled = fill_pillow_prompt(NANA_DIALOG_PROMPT_A, variables, "嗯")
        self.assertNotIn("{{", filled)
        self.assertNotIn("None", filled)


class CampaignTests(unittest.TestCase):
    def _run(self, *, loop_count: int, rounds_per_loop: int):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        results = root / "results.jsonl"
        nana = FakeClient("奈奈")
        user = FakeClient("冉冉")

        ok = run_campaign_once(
            cfg=RunnerConfig(
                retry_times=1, loop_count=loop_count, rounds_per_loop=rounds_per_loop
            ),
            campaign_id="test",
            run_index=1,
            results_path=results,
            logger=logging.getLogger("test"),
            client_factory=lambda _cfg: nana,
            user_client_factory=lambda _cfg: user,
            variables_fetcher=fake_fetcher,
        )
        rows = [json.loads(line) for line in results.read_text(encoding="utf-8").splitlines()]
        return ok, rows, nana, user, root

    def test_row_count_and_role_alternation(self):
        ok, rows, _nana, _user, _root = self._run(loop_count=1, rounds_per_loop=4)
        self.assertTrue(ok)
        # 一轮 = 用户一句 + 奈奈一句，所以每段会话 2 * rounds_per_loop 行。
        expected = len(VARIANT_PROMPTS) * 1 * 2 * 4
        self.assertEqual(len(rows), expected)
        self.assertEqual([row["variant"] for row in rows], ["A"] * 8 + ["B"] * 8)
        session = rows[:8]
        self.assertEqual(
            [row["role"] for row in session], ["user", "assistant"] * 4
        )
        self.assertEqual([row["round"] for row in session], [1, 1, 2, 2, 3, 3, 4, 4])
        self.assertEqual(session[0]["raw_user_message"], OPENING_USER_MESSAGE)

    def test_assembled_prompt_is_recorded_but_never_enters_history(self):
        ok, rows, nana, user, _root = self._run(loop_count=1, rounds_per_loop=4)
        self.assertTrue(ok)

        assistant_rows = [row for row in rows if row["role"] == "assistant"]
        user_rows = [row for row in rows if row["role"] == "user"]
        # 每个 variant 一段 4 轮的会话 → 4 个奈奈轮 + 4 个用户轮。
        self.assertEqual(len(assistant_rows), 4 * len(VARIANT_PROMPTS))
        self.assertEqual(len(user_rows), 4 * len(VARIANT_PROMPTS))
        for row in assistant_rows:
            self.assertIn("SENTINEL_WORLDVIEW", row["assembled_prompt"])
            self.assertNotIn("{{", row["assembled_prompt"])
        for row in user_rows:  # 用户轮发的是原话，从不拼装
            self.assertEqual(row["assembled_prompt"], "")

        # 关键回归护栏：拼装文本只出现在每次奈奈请求的最后一条消息上，
        # 历史里的消息保持原话，绝不带哨兵。
        self.assertEqual(len(nana.requests), len(assistant_rows))
        for request in nana.requests:
            messages = request["messages"]
            self.assertEqual(messages[0]["role"], "system")
            for message in messages[:-1]:
                for sentinel in SENTINELS:
                    self.assertNotIn(sentinel, message["content"])
            self.assertIn("SENTINEL_WORLDVIEW", messages[-1]["content"])
            for sentinel in SENTINELS:
                self.assertNotIn(sentinel, user_messages_text(user.requests))

    def test_request_parameters_per_side(self):
        ok, _rows, nana, user, _root = self._run(loop_count=1, rounds_per_loop=2)
        self.assertTrue(ok)
        cfg = RunnerConfig()
        self.assertTrue(nana.requests)
        self.assertTrue(user.requests)
        for request in nana.requests:
            self.assertEqual(request["model"], cfg.model_name)
            self.assertEqual(request["temperature"], cfg.temperature)
            self.assertEqual(request["max_tokens"], cfg.max_tokens)
            self.assertEqual(request["extra_body"], {"thinking": {"type": "disabled"}})
            self.assertNotIn("tools", request)
            self.assertNotIn("tool_choice", request)
        for request in user.requests:
            self.assertEqual(request["model"], cfg.user_model_name)
            self.assertNotIn("extra_body", request)
            self.assertNotIn("tools", request)
            self.assertNotIn("tool_choice", request)

    def test_report_has_assembled_prompt_column(self):
        from openpyxl import load_workbook

        ok, _rows, _nana, _user, root = self._run(loop_count=1, rounds_per_loop=2)
        self.assertTrue(ok)
        workbook = load_workbook(root / "report.xlsx")
        self.assertEqual(workbook.sheetnames, ["逐轮记录", "汇总", "运行批次"])
        header = [cell.value for cell in workbook["逐轮记录"][1]]
        for column in ("role", "session", "round", "raw_user_message", "assembled_prompt", "reply_text"):
            self.assertIn(column, header)
        # 2 组 × 1 段 × 2 轮 × 每轮 2 行 = 8 行。
        self.assertEqual(workbook["逐轮记录"].max_row - 1, 2 * 1 * 2 * 2)


def user_messages_text(requests) -> str:
    return "\n".join(
        message["content"]
        for request in requests
        for message in request["messages"]
    )


class RunnerLockTests(unittest.TestCase):
    def test_runner_lock_prevents_second_runner(self):
        with tempfile.TemporaryDirectory() as temp:
            with SingleRunnerLock(Path(temp)):
                with self.assertRaisesRegex(RuntimeError, "another runner is active"):
                    with SingleRunnerLock(Path(temp)):
                        pass


if __name__ == "__main__":
    unittest.main()
