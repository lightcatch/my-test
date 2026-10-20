r"""Run and persist the Nana prompt A/B dialogue campaign.

Two variants (A / B) each carry their own system prompt and dialog prompt, both
declared in ``config.py``.  A session runs ``rounds_per_loop`` rounds, and one
round is a user line followed by one Nana reply:

* the user line of round 1 is the fixed ``OPENING_USER_MESSAGE`` and costs no
  model call; later rounds ask the simulated user (DeepSeek) for the next line;
* the Nana side (Doubao) answers exactly once per round.

Before every Nana reply the ``conversation-variables`` API is queried and its
variables are filled into the variant's dialog prompt.  That assembled text is
what gets sent on the wire, while ``history`` only ever stores the raw text, so
the assembly is turn-local.  Every utterance writes its own record row (two per
round), and the assembled prompt is kept in the ``assembled_prompt`` column.

PowerShell:
    # Default: 8 hours, every 30 minutes, using the defaults in config.py.
    & D:\PyCode\test\nana_prompt_ab\start_runner.ps1

    # One immediate A/B campaign, then exit.
    & D:\PyCode\test\nana_prompt_ab\start_runner.ps1 -Once

    # Override only the values that differ from the config.py defaults.
    & D:\PyCode\test\nana_prompt_ab\start_runner.ps1 -Model doubao-seed-2-1-turbo-260628 -DurationHours 2
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import time
import uuid
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from config import (
    OPENING_USER_MESSAGE,
    USER_SIMULATOR_PROMPT,
    VARIANT_PROMPTS,
    RunnerConfig,
    build_client,
    build_user_client,
    load_config,
    parser,
)
from prompts import fetch_pillow_variables, fill_pillow_prompt, prompt_hash


VARIANTS = ("A", "B")
_PARENTHESIS_RE = re.compile(r"[()（）【】\[\]]")
_TONE_PREFIX_RE = re.compile(r"^[^*\r\n]{1,40}\*\*")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def make_output_dir(root: Path | None) -> Path:
    base = root or Path(__file__).resolve().parent / "outputs"
    campaign = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:8]
    path = base / campaign
    path.mkdir(parents=True, exist_ok=False)
    return path


class SingleRunnerLock:
    """An exclusive lock shared by all runners targeting one output root."""

    def __init__(self, output_root: Path):
        self.output_root = output_root
        self.path = output_root / ".runner.lock"
        self.fd: int | None = None

    @staticmethod
    def _pid_is_running(pid: int) -> bool:
        if pid <= 0:
            return False
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return False
        except PermissionError:
            return True
        except OSError:
            return False
        return True

    def __enter__(self) -> "SingleRunnerLock":
        self.output_root.mkdir(parents=True, exist_ok=True)
        try:
            self.fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                details = json.loads(self.path.read_text(encoding="utf-8"))
                active = self._pid_is_running(int(details.get("pid", 0)))
            except (OSError, ValueError, TypeError, json.JSONDecodeError):
                active = False
            if active:
                raise RuntimeError(f"another runner is active; lock={self.path}") from None
            # A crashed process left a stale lock. It is safe to reclaim it.
            self.path.unlink(missing_ok=True)
            self.fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        assert self.fd is not None
        os.write(self.fd, json.dumps({"pid": os.getpid(), "started_at": utc_now()}).encode("utf-8"))
        return self

    def __exit__(self, exc_type: object, exc: object, traceback: object) -> None:
        if self.fd is not None:
            os.close(self.fd)
            self.fd = None
        self.path.unlink(missing_ok=True)


def setup_logger(path: Path) -> logging.Logger:
    logger = logging.getLogger(f"nana_prompt_ab.{path}")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    handler = logging.FileHandler(path, encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(handler)
    logger.addHandler(logging.StreamHandler(sys.stdout))
    return logger


def extract_response_text(response: Any) -> str:
    choices = getattr(response, "choices", None)
    if not choices:
        return ""
    message = getattr(choices[0], "message", None)
    content = getattr(message, "content", "") if message is not None else ""
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        return "".join(str(getattr(item, "text", "")) for item in content).strip()
    return str(content).strip() if content is not None else ""


def observation(text: str) -> dict[str, Any]:
    stripped = text.strip()
    return {
        "output_chars": len(text),
        "starts_with_tone_marker": bool(_TONE_PREFIX_RE.match(stripped)),
        "contains_parentheses": bool(_PARENTHESIS_RE.search(text)),
        "empty_output": not bool(stripped),
    }


def apply_history_window(
    history: list[dict[str, str]], *, enabled: bool, window_messages: int
) -> list[dict[str, str]]:
    """Trim what is sent to the model; ``history`` itself keeps everything."""
    if not enabled or len(history) <= window_messages:
        return list(history)
    return history[-window_messages:]


def map_history_for_deepseek(history: list[dict[str, str]]) -> list[dict[str, str]]:
    """Swap the roles so the user simulator sees itself as the assistant."""
    mapped: list[dict[str, str]] = []
    for message in history:
        role = message["role"]
        if role == "user":
            mapped_role = "assistant"
        elif role == "assistant":
            mapped_role = "user"
        else:
            raise ValueError(f"Unexpected role in history: {role}")
        mapped.append({"role": mapped_role, "content": message["content"]})
    return mapped


def request_once(
    client: Any,
    cfg: RunnerConfig,
    messages: list[dict[str, str]],
    *,
    user_side: bool = False,
) -> tuple[str, float]:
    """One model call. No tools / tool choice are ever sent."""
    if user_side:
        kwargs: dict[str, Any] = {"model": cfg.user_model_name, **cfg.user_request_parameters}
    else:
        kwargs = {"model": cfg.model_name, **cfg.request_parameters}
    started = time.perf_counter_ns()
    response = client.chat.completions.create(messages=messages, **kwargs)
    elapsed = (time.perf_counter_ns() - started) / 1_000_000
    return extract_response_text(response), round(elapsed, 3)


def call_with_retries(
    *,
    client: Any,
    cfg: RunnerConfig,
    messages: list[dict[str, str]],
    user_side: bool,
) -> tuple[str, list[float], int, str | None]:
    errors: list[str] = []
    latencies: list[float] = []
    text = ""
    attempts = 0
    for attempt in range(1, cfg.retry_times + 1):
        attempts = attempt
        started = time.perf_counter_ns()
        try:
            text, latency = request_once(client, cfg, messages, user_side=user_side)
            if not text:
                raise RuntimeError("empty model response")
            latencies.append(latency)
            break
        except Exception as exc:  # noqa: BLE001
            text = ""
            latencies.append(round((time.perf_counter_ns() - started) / 1_000_000, 3))
            errors.append(f"attempt {attempt}: {type(exc).__name__}: {exc}")
            if attempt < cfg.retry_times:
                time.sleep(cfg.retry_backoff_seconds)
    error_text = None if text else ("; ".join(errors) if errors else "empty model response")
    return text, latencies, attempts, error_text


def build_record(
    *,
    cfg: RunnerConfig,
    campaign_id: str,
    run_index: int,
    variant: str,
    session_index: int,
    round_index: int,
    role: str,
    raw_user_message: str,
    assembled_prompt: str,
    reply_text: str,
    latency_ms: float | None,
    attempt_latencies_ms: list[float],
    attempts: int,
    error: str | None,
    skipped: bool,
    fingerprint: str,
) -> dict[str, Any]:
    return {
        "timestamp": utc_now(),
        "campaign_id": campaign_id,
        "run_index": run_index,
        "variant": variant,
        "model": cfg.model_name if role == "assistant" else cfg.user_model_name,
        "session": session_index,
        "round": round_index,
        "role": role,
        "raw_user_message": raw_user_message,
        "assembled_prompt": assembled_prompt,
        "reply_text": reply_text,
        "latency_ms": latency_ms,
        "attempt_latencies_ms": attempt_latencies_ms,
        "attempts": attempts,
        "error": error,
        "skipped": skipped,
        **observation(reply_text),
        "prompt_hash": fingerprint,
    }


def run_session(
    *,
    client: Any,
    user_client: Any,
    cfg: RunnerConfig,
    campaign_id: str,
    run_index: int,
    variant: str,
    session_index: int,
    prompts: dict[str, str],
    write_record: Callable[[dict[str, Any]], None],
    logger: logging.Logger,
    variables_fetcher: Callable[[str, str], dict | None] = fetch_pillow_variables,
) -> bool:
    """One session: ``rounds_per_loop`` rounds, one round = user line + reply.

    Round 1's user line is the fixed ``OPENING_USER_MESSAGE`` and costs no model
    call; every later round asks the simulated user for its next line.  Each
    round therefore writes exactly two rows (user, then assistant).

    ``history`` only ever holds raw text, which is why the assembled prompt
    never leaks into a later request.
    """
    fingerprint = prompt_hash(prompts)
    history: list[dict[str, str]] = []

    def record(**kwargs: Any) -> dict[str, Any]:
        return build_record(
            cfg=cfg,
            campaign_id=campaign_id,
            run_index=run_index,
            variant=variant,
            session_index=session_index,
            fingerprint=fingerprint,
            **kwargs,
        )

    def skip_round(round_index: int) -> None:
        """Write the two skipped rows for a round that never ran."""
        for role in ("user", "assistant"):
            write_skip(round_index, role, history[-1]["content"] if history else "")

    def write_skip(round_index: int, role: str, raw: str) -> None:
        write_record(
            record(
                round_index=round_index,
                role=role,
                raw_user_message=raw,
                assembled_prompt="",
                reply_text="",
                latency_ms=None,
                attempt_latencies_ms=[],
                attempts=0,
                error="skipped_after_previous_failure",
                skipped=True,
            )
        )

    failed = False
    for round_index in range(1, cfg.rounds_per_loop + 1):
        if failed:
            skip_round(round_index)
            continue

        # ---- 第一句：第 1 轮用固定开场白，省一次模型调用 ----
        if round_index == 1:
            user_text, user_latencies, user_attempts, user_error = (
                OPENING_USER_MESSAGE, [], 0, None,
            )
        else:
            user_text, user_latencies, user_attempts, user_error = call_with_retries(
                client=user_client,
                cfg=cfg,
                messages=[
                    {"role": "system", "content": USER_SIMULATOR_PROMPT},
                    *map_history_for_deepseek(
                        apply_history_window(
                            history, enabled=True, window_messages=cfg.window_messages
                        )
                    ),
                ],
                user_side=True,
            )
        write_record(
            record(
                round_index=round_index,
                role="user",
                raw_user_message=user_text,
                assembled_prompt="",
                reply_text=user_text,
                latency_ms=round(sum(user_latencies), 3) if user_latencies else None,
                attempt_latencies_ms=user_latencies,
                attempts=user_attempts,
                error=user_error,
                skipped=False,
            )
        )
        if not user_text:
            failed = True
            logger.error(
                "campaign=%s variant=%s session=%s round=%s role=user failed: %s",
                campaign_id, variant, session_index, round_index, user_error,
            )
            # 用户轮失败，本轮补一条跳过的奈奈行，保持「每轮两行」。
            write_skip(round_index, "assistant", user_text)
            continue
        history.append({"role": "user", "content": user_text})

        # ---- 第二句：奈奈回一句，拼装结果只上 wire，history 仍存原话 ----
        variables = variables_fetcher(prompts["sn"], user_text)
        assembled = ""
        if variables:
            assembled = fill_pillow_prompt(prompts["dialog_prompt"], variables, user_text)
        windowed = apply_history_window(
            history, enabled=True, window_messages=cfg.window_messages
        )
        if assembled:
            wire = [*windowed[:-1], {"role": "user", "content": assembled}]
        else:
            wire = windowed
        messages = [{"role": "system", "content": prompts["system_prompt"]}, *wire]

        text, latencies, attempts, error_text = call_with_retries(
            client=client, cfg=cfg, messages=messages, user_side=False
        )
        succeeded = bool(text)
        if not succeeded:
            failed = True
            logger.error(
                "campaign=%s variant=%s session=%s round=%s role=assistant failed: %s",
                campaign_id, variant, session_index, round_index, error_text,
            )
        write_record(
            record(
                round_index=round_index,
                role="assistant",
                raw_user_message=user_text,
                assembled_prompt=assembled,
                reply_text=text,
                latency_ms=round(sum(latencies), 3) if latencies else None,
                attempt_latencies_ms=latencies,
                attempts=attempts,
                error=error_text,
                skipped=False,
            )
        )
        if succeeded:
            history.append({"role": "assistant", "content": text})
    return not failed


def append_jsonl(path: Path, record: dict[str, Any]) -> None:
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
        handle.flush()


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    index = (len(ordered) - 1) * q
    lower, upper = int(index), min(int(index) + 1, len(ordered) - 1)
    value = ordered[lower] + (ordered[upper] - ordered[lower]) * (index - lower)
    return round(value, 3)


def read_records(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    records: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                records.append(json.loads(line))
    return records


def _ok(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [row for row in rows if not row.get("error") and not row.get("skipped")]


def _latencies(rows: list[dict[str, Any]]) -> list[float]:
    return [float(row["latency_ms"]) for row in rows if row.get("latency_ms") is not None]


def build_report(path: Path, records: list[dict[str, Any]], cfg: RunnerConfig) -> None:
    book = Workbook()
    details = book.active
    details.title = "逐轮记录"
    fields = [
        "timestamp", "campaign_id", "run_index", "variant", "model", "session", "round", "role",
        "raw_user_message", "assembled_prompt", "reply_text", "latency_ms",
        "attempt_latencies_ms", "attempts", "error", "skipped", "output_chars",
        "starts_with_tone_marker", "contains_parentheses", "empty_output", "prompt_hash",
    ]
    details.append(fields)
    for record in records:
        details.append([
            json.dumps(record.get(field), ensure_ascii=False) if field == "attempt_latencies_ms" else record.get(field)
            for field in fields
        ])

    summary = book.create_sheet("汇总")
    summary.append(["指标", "A", "B"])

    def per_variant(fn: Callable[[list[dict[str, Any]]], Any]) -> list[Any]:
        return [fn([r for r in records if r["variant"] == variant]) for variant in VARIANTS]

    def assistants(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [row for row in rows if row.get("role") == "assistant"]

    def users(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [row for row in rows if row.get("role") == "user"]

    def ratio(rows: list[dict[str, Any]], predicate: Callable[[dict[str, Any]], bool]) -> float | None:
        return round(sum(predicate(row) for row in rows) / len(rows), 4) if rows else None

    def average_ok_latency(rows: list[dict[str, Any]]) -> float | None:
        values = _latencies(_ok(assistants(rows)))
        return round(mean(values), 3) if values else None

    for name, fn in (
        ("总行数", lambda rows: len(rows)),
        ("奈奈轮数", lambda rows: len(assistants(rows))),
        ("模拟用户轮数", lambda rows: len(users(rows))),
        ("成功率", lambda rows: ratio(rows, lambda row: not row.get("error") and not row.get("skipped"))),
        ("奈奈轮成功率", lambda rows: ratio(assistants(rows), lambda row: not row.get("error") and not row.get("skipped"))),
        ("平均延时(ms)", average_ok_latency),
        ("P50延时(ms)", lambda rows: percentile(_latencies(_ok(assistants(rows))), 0.5)),
        ("P95延时(ms)", lambda rows: percentile(_latencies(_ok(assistants(rows))), 0.95)),
        ("平均输出字符数(奈奈轮)", lambda rows: round(mean([int(r.get("output_chars", 0)) for r in _ok(assistants(rows))]), 2) if _ok(assistants(rows)) else None),
        ("语气标记出现率(奈奈轮)", lambda rows: ratio(assistants(rows), lambda row: bool(row.get("starts_with_tone_marker")))),
        ("括号出现率(奈奈轮)", lambda rows: ratio(assistants(rows), lambda row: bool(row.get("contains_parentheses")))),
        ("拼装降级轮数", lambda rows: sum(1 for r in _ok(assistants(rows)) if not r.get("assembled_prompt"))),
    ):
        summary.append([name] + per_variant(fn))
    summary.append([])
    summary.append(["配置", "值"])
    for key, value in asdict(cfg).items():
        summary.append([key, value])

    runs = book.create_sheet("运行批次")
    runs.append(["campaign_id", "run_index", "variant", "records", "success", "failed", "mean_latency_ms"])
    seen = sorted({(r["campaign_id"], r["run_index"], r["variant"]) for r in records})
    for campaign_id, run_index, variant in seen:
        rows = [r for r in records if (r["campaign_id"], r["run_index"], r["variant"]) == (campaign_id, run_index, variant)]
        latencies = _latencies(_ok(assistants(rows)))
        runs.append([
            campaign_id, run_index, variant, len(rows),
            sum(not r.get("error") and not r.get("skipped") for r in rows),
            sum(bool(r.get("error")) for r in rows),
            round(mean(latencies), 3) if latencies else None,
        ])

    for sheet in book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    temp = path.with_suffix(path.suffix + ".tmp")
    book.save(temp)
    os.replace(temp, path)


def run_campaign_once(
    *,
    cfg: RunnerConfig,
    campaign_id: str,
    run_index: int,
    results_path: Path,
    logger: logging.Logger,
    client_factory: Callable[[RunnerConfig], Any] = build_client,
    user_client_factory: Callable[[RunnerConfig], Any] = build_user_client,
    variables_fetcher: Callable[[str, str], dict | None] = fetch_pillow_variables,
) -> bool:
    client = client_factory(cfg)
    user_client = user_client_factory(cfg)

    def write(record: dict[str, Any]) -> None:
        append_jsonl(results_path, record)

    aborted = False
    for variant in VARIANTS:
        if aborted:
            break
        for session_index in range(1, cfg.loop_count + 1):
            ok = run_session(
                client=client,
                user_client=user_client,
                cfg=cfg,
                campaign_id=campaign_id,
                run_index=run_index,
                variant=variant,
                session_index=session_index,
                prompts=VARIANT_PROMPTS[variant],
                write_record=write,
                logger=logger,
                variables_fetcher=variables_fetcher,
            )
            if not ok:
                aborted = True
                break

    records = [
        r for r in read_records(results_path)
        if r.get("campaign_id") == campaign_id and r.get("run_index") == run_index
    ]
    build_report(results_path.with_name("report.xlsx"), read_records(results_path), cfg)
    # 每轮两行（用户一句 + 奈奈一句）。
    expected = len(VARIANTS) * cfg.loop_count * 2 * cfg.rounds_per_loop
    return len(records) == expected and all(
        not record.get("error") and not record.get("skipped") and record.get("reply_text")
        for record in records
    )


def run_main(args: argparse.Namespace, cfg: RunnerConfig, output_root: Path) -> int:
    output_dir = make_output_dir(output_root)
    results_path = output_dir / "results.jsonl"
    logger = setup_logger(output_dir / "runner.log")
    campaign_id = output_dir.name
    logger.info("started campaign=%s model=%s interval_minutes=%s window_hours=%s", campaign_id, cfg.model_name, cfg.interval_minutes, cfg.window_hours)
    started = time.monotonic()
    run_index = 0
    had_failure = False
    try:
        while True:
            if run_index and args.once:
                break
            if run_index and time.monotonic() - started >= cfg.window_hours * 3600:
                break
            run_index += 1
            logger.info("running campaign=%s run_index=%s", campaign_id, run_index)
            campaign_ok = run_campaign_once(
                cfg=cfg,
                campaign_id=campaign_id,
                run_index=run_index,
                results_path=results_path,
                logger=logger,
            )
            had_failure = not campaign_ok or had_failure
            if args.once:
                break
            if not campaign_ok:
                logger.error("stopping campaign after failed run; no further scheduled retries")
                break
            next_run = started + run_index * cfg.interval_minutes * 60
            delay = max(0.0, next_run - time.monotonic())
            if time.monotonic() + delay >= started + cfg.window_hours * 3600:
                break
            logger.info("next run in %.1f seconds", delay)
            time.sleep(delay)
    except KeyboardInterrupt:
        logger.info("stopped by user")
        return 130
    finally:
        logger.info("finished campaign=%s runs=%s results=%s", campaign_id, run_index, results_path)
    status = "FAIL" if had_failure else "PASS"
    print(f"{status}: campaign={campaign_id} runs={run_index} results={results_path} report={output_dir / 'report.xlsx'}")
    return 1 if had_failure else 0


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        cfg = load_config(args)
        output_root = args.output_dir or Path(__file__).resolve().parent / "outputs"
        with SingleRunnerLock(output_root):
            return run_main(args, cfg, output_root)
    except Exception as exc:  # noqa: BLE001
        print(f"FAIL: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

