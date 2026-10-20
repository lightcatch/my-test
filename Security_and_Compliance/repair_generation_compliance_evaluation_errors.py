# -*- coding: utf-8 -*-
"""重新评测生成内容合规评分工作簿中的 evaluation_error 记录。"""

from __future__ import annotations

import argparse
import sys
import time
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl import load_workbook

from audit_generation_compliance_with_deepseekpro import JUDGE_PROMPT, POLICY, parse_verdict
from audit_refusal_with_deepseekpro import (
    FAIL_FILL,
    FAIL_FONT,
    QUESTION_HEADER,
    REPLY_HEADER,
    RISK_HEADER,
    AuditError,
    configure_console,
    deepseek_config,
    remove_secret,
    text,
)
from repair_evaluation_common import (
    RepairError,
    column_map as required_column_map,
    detail_rows as workbook_detail_rows,
    find_normal_row as style_reference_row,
)


TARGET_FILE = Path(r"D:\测试用例模板\新题库\测试结果\题库精简\生成内容合规评分_DeepSeekPro_v4pro.xlsx")
MAX_RETRIES = 10
TIMEOUT_SECONDS = 90
SOURCE_ROW_HEADER = "源Excel行号"
AUDITOR_HEADER = "测评员"
CONTAINS_HEADER = "是否包含违规内容"
DECISION_HEADER = "判定"
REASON_HEADER = "判定理由"
STATUS_HEADER = "测评状态"
SUBCLASS_HEADER = "风险子类/场景子类"


@dataclass(frozen=True)
class ErrorCase:
    worksheet_row: int
    risk: str
    subclass: str
    question: str
    reply: str


def column_map(sheet) -> dict[str, int]:
    required = (
        SOURCE_ROW_HEADER,
        RISK_HEADER,
        SUBCLASS_HEADER,
        QUESTION_HEADER,
        REPLY_HEADER,
        AUDITOR_HEADER,
        CONTAINS_HEADER,
        DECISION_HEADER,
        REASON_HEADER,
        STATUS_HEADER,
    )
    return required_column_map(sheet, required, text)


def detail_rows(sheet, columns: dict[str, int]) -> list[int]:
    return workbook_detail_rows(sheet, SOURCE_ROW_HEADER, columns)


def error_cases(sheet, columns: dict[str, int], rows: list[int]) -> list[ErrorCase]:
    cases: list[ErrorCase] = []
    for row in rows:
        if text(sheet.cell(row, columns[STATUS_HEADER]).value) != "evaluation_error":
            continue
        cases.append(
            ErrorCase(
                worksheet_row=row,
                risk=text(sheet.cell(row, columns[RISK_HEADER]).value),
                subclass=text(sheet.cell(row, columns[SUBCLASS_HEADER]).value),
                question="" if sheet.cell(row, columns[QUESTION_HEADER]).value is None else str(sheet.cell(row, columns[QUESTION_HEADER]).value),
                reply="" if sheet.cell(row, columns[REPLY_HEADER]).value is None else str(sheet.cell(row, columns[REPLY_HEADER]).value),
            )
        )
    return cases


def re_evaluate(case: ErrorCase, client: OpenAI, api_key: str, model: str):
    if not case.reply.strip():
        raise RepairError(f"第 {case.worksheet_row} 行模型回复为空，不能将 evaluation_error 重评为有效结论")
    prompt = JUDGE_PROMPT.format(
        policy=POLICY,
        risk=case.risk,
        subclass=case.subclass,
        question=case.question,
        reply=case.reply,
    )
    last_error = ""
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
                max_tokens=300,
                response_format={"type": "json_object"},
            )
            return parse_verdict(response.choices[0].message.content or "")
        except Exception as exc:
            last_error = remove_secret(f"{type(exc).__name__}: {exc}", api_key)
            if attempt < MAX_RETRIES:
                time.sleep(min(2 ** (attempt - 1), 16))
    raise RepairError(f"第 {case.worksheet_row} 行重评 {MAX_RETRIES} 次后仍未得到有效 JSON：{last_error}")


def find_normal_row(sheet, columns: dict[str, int], rows: list[int]) -> int:
    return style_reference_row(
        sheet,
        columns,
        rows,
        status_header=STATUS_HEADER,
        decision_header=DECISION_HEADER,
        passing_decision="合格",
        text=text,
    )


def apply_verdict(sheet, columns: dict[str, int], case: ErrorCase, verdict, model: str, normal_row: int) -> None:
    row = case.worksheet_row
    sheet.cell(row, columns[AUDITOR_HEADER]).value = f"DeepSeekPro ({model})"
    sheet.cell(row, columns[CONTAINS_HEADER]).value = verdict.contains_violation
    sheet.cell(row, columns[DECISION_HEADER]).value = verdict.decision
    sheet.cell(row, columns[REASON_HEADER]).value = verdict.reason
    sheet.cell(row, columns[STATUS_HEADER]).value = verdict.status

    style_columns = (columns[QUESTION_HEADER], columns[DECISION_HEADER], columns[STATUS_HEADER])
    if verdict.decision == "合格":
        for column in style_columns:
            target = sheet.cell(row, column)
            source = sheet.cell(normal_row, column)
            target.fill = copy(source.fill)
            target.font = copy(source.font)
    else:
        for column in style_columns:
            target = sheet.cell(row, column)
            target.fill = FAIL_FILL
            target.font = FAIL_FONT


def recompute_summary(sheet, columns: dict[str, int], rows: list[int]) -> tuple[int, int, Counter[str], Counter[str]]:
    total = len(rows)
    decisions = Counter(text(sheet.cell(row, columns[DECISION_HEADER]).value) for row in rows)
    statuses = Counter(text(sheet.cell(row, columns[STATUS_HEADER]).value) for row in rows)
    if statuses.get("evaluation_error", 0):
        raise RepairError("仍存在 evaluation_error，拒绝写入未完成的重评结果")
    if any(decision not in ("合格", "不合格") for decision in decisions):
        raise RepairError(f"明细存在无效判定：{dict(decisions)}")

    passed = decisions["合格"]
    summary_row = next(
        (row for row in range(2, sheet.max_row + 1) if text(sheet.cell(row, 1).value) == "合规汇总"),
        None,
    )
    if summary_row is None:
        raise RepairError("找不到合规汇总区域")
    summary_values = {
        "检查总数": total,
        "合格数": passed,
        "不合格数": total - passed,
        "合格率": passed / total,
        "不合格率": (total - passed) / total,
    }
    for row in range(summary_row + 1, sheet.max_row + 1):
        label = text(sheet.cell(row, 1).value)
        if label in summary_values:
            sheet.cell(row, 2).value = summary_values[label]

    groups: dict[str, list[int]] = defaultdict(list)
    for row in rows:
        groups[text(sheet.cell(row, columns[RISK_HEADER]).value)].append(row)
    risk_header_row = summary_row + 1
    risk_count_column = None
    risk_passed_column = None
    for column in range(4, sheet.max_column + 1):
        header = text(sheet.cell(risk_header_row, column).value)
        if header == "检查数":
            risk_count_column = column
        elif header == "合格数":
            risk_passed_column = column
    if risk_count_column is None or risk_passed_column is None:
        raise RepairError("风险大类汇总缺少检查数或合格数列")
    for row in range(risk_header_row + 1, sheet.max_row + 1):
        risk = text(sheet.cell(row, 4).value)
        if risk not in groups:
            continue
        group_rows = groups[risk]
        sheet.cell(row, risk_count_column).value = len(group_rows)
        sheet.cell(row, risk_passed_column).value = sum(
            text(sheet.cell(detail_row, columns[DECISION_HEADER]).value) == "合格" for detail_row in group_rows
        )
    return total, passed, decisions, statuses


def verify_saved(path: Path) -> tuple[int, int, int]:
    # The verification needs random cell access; read-only worksheets make that O(n^2).
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        columns = column_map(sheet)
        rows = detail_rows(sheet, columns)
        decisions = Counter(text(sheet.cell(row, columns[DECISION_HEADER]).value) for row in rows)
        statuses = Counter(text(sheet.cell(row, columns[STATUS_HEADER]).value) for row in rows)
        if statuses.get("evaluation_error", 0):
            raise RepairError("保存后仍存在 evaluation_error")
        summary_row = next(
            row for row in range(2, sheet.max_row + 1) if text(sheet.cell(row, 1).value) == "合规汇总"
        )
        stored = {
            text(sheet.cell(row, 1).value): sheet.cell(row, 2).value
            for row in range(summary_row + 1, sheet.max_row + 1)
            if text(sheet.cell(row, 1).value)
        }
        expected_total = len(rows)
        expected_passed = decisions["合格"]
        if stored.get("检查总数") != expected_total or stored.get("合格数") != expected_passed:
            raise RepairError("保存后的总体汇总与明细不一致")
        if stored.get("不合格数") != expected_total - expected_passed:
            raise RepairError("保存后的不合格数与明细不一致")
        return expected_total, expected_passed, statuses.get("evaluation_error", 0)
    finally:
        workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="修复生成内容合规评分工作簿中的 evaluation_error")
    parser.add_argument("--input", default=str(TARGET_FILE), help="待修复的合规评分 Excel 路径")
    parser.add_argument("--model", default="deepseek-v4-pro", help="DeepSeek 评审模型")
    parser.add_argument("--dry-run", action="store_true", help="仅列出 evaluation_error，不调用 API 或覆盖文件")
    return parser.parse_args()


def main() -> None:
    configure_console()
    args = parse_args()
    target = Path(args.input)
    if not target.is_file():
        raise RepairError(f"找不到待修复工作簿：{target}")
    workbook = load_workbook(target)
    try:
        sheet = workbook.active
        columns = column_map(sheet)
        rows = detail_rows(sheet, columns)
        cases = error_cases(sheet, columns, rows)
        print(f"明细记录数：{len(rows)}；evaluation_error：{len(cases)}")
        if not cases:
            print("没有 evaluation_error；无需修改。")
            return
        for case in cases:
            print(f"待重评：工作表第 {case.worksheet_row} 行，{case.risk}")
        if args.dry_run:
            return
        api_key, base_url, model = deepseek_config(args.model)
        client = OpenAI(api_key=api_key, base_url=base_url, timeout=TIMEOUT_SECONDS, max_retries=0)
        verdicts = {case.worksheet_row: re_evaluate(case, client, api_key, model) for case in cases}
        normal_row = find_normal_row(sheet, columns, rows)
        for case in cases:
            apply_verdict(sheet, columns, case, verdicts[case.worksheet_row], model, normal_row)
        total, passed, _, _ = recompute_summary(sheet, columns, rows)
        workbook.save(target)
    finally:
        workbook.close()
    verified_total, verified_passed, remaining_errors = verify_saved(target)
    print(f"已覆盖更新：{target}")
    print(f"合格率：{verified_passed / verified_total:.2%}（{verified_passed}/{verified_total}）；剩余 evaluation_error：{remaining_errors}")


if __name__ == "__main__":
    try:
        main()
    except (AuditError, RepairError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(2)
