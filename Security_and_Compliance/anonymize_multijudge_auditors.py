# -*- coding: utf-8 -*-
"""将三模型人工抽检工作簿中的测评员显示名称匿名化。"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TARGET_FILE = Path(r"D:\测试用例模板\新题库\测试结果\题库精简\生成内容合规人工抽检.xlsx")
SOURCE_ROW_HEADER = "源Excel行号"
FINAL_AUDITOR_HEADER = "测评员"
FINAL_REASON_HEADER = "判定理由"

REPLACEMENTS = (
    ("豆包 Seed 2.0 Lite 260428", "测评员3"),
    ("doubao-seed-2-0-lite-260428", "测评员3"),
    ("DeepSeek v4 Pro", "测评员1"),
    ("deepseek-v4-pro", "测评员1"),
    ("Qwen3-Max", "测评员2"),
    ("qwen3-max", "测评员2"),
)
ORIGINAL_IDENTIFIERS = tuple(source for source, _ in REPLACEMENTS)


class AnonymizeError(RuntimeError):
    """工作簿结构或验证失败。"""


def configure_console() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


def text(value: Any) -> str:
    return str(value or "").strip()


def replace_identifiers(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    updated = value
    for source, target in REPLACEMENTS:
        updated = updated.replace(source, target)
    return updated


def find_columns(sheet) -> dict[str, int]:
    headers = [text(cell.value) for cell in sheet[1]]
    required = (SOURCE_ROW_HEADER, FINAL_AUDITOR_HEADER, FINAL_REASON_HEADER)
    missing = [header for header in required if header not in headers]
    if missing:
        raise AnonymizeError(f"缺少必要列：{missing}")
    return {header: headers.index(header) + 1 for header in required}


def detail_rows(sheet, source_row_column: int) -> list[int]:
    rows = [
        row
        for row in range(2, sheet.max_row + 1)
        if isinstance(sheet.cell(row, source_row_column).value, int)
    ]
    if not rows:
        raise AnonymizeError("工作簿中没有可识别的明细行")
    return rows


def numeric_snapshot(sheet) -> dict[tuple[int, int], Any]:
    return {
        (row, column): cell.value
        for row in range(1, sheet.max_row + 1)
        for column in range(1, sheet.max_column + 1)
        if isinstance((cell := sheet.cell(row, column)).value, (int, float)) and not isinstance(cell.value, bool)
    }


def verify_anonymized(path: Path, expected_numbers: dict[tuple[int, int], Any]) -> tuple[int, int]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        columns = find_columns(sheet)
        rows = detail_rows(sheet, columns[SOURCE_ROW_HEADER])
        current_numbers = numeric_snapshot(sheet)
        if current_numbers != expected_numbers:
            raise AnonymizeError("匿名化操作意外修改了数值单元格")

        headers = [text(cell.value) for cell in sheet[1]]
        expected_prefixes = ("测评员1 ", "测评员2 ", "测评员3 ")
        for prefix in expected_prefixes:
            if not any(header.startswith(prefix) for header in headers):
                raise AnonymizeError(f"缺少匿名化后的评审列：{prefix}")

        display_cells: list[Any] = list(sheet[1])
        for row in rows:
            display_cells.append(sheet.cell(row, columns[FINAL_AUDITOR_HEADER]))
            display_cells.append(sheet.cell(row, columns[FINAL_REASON_HEADER]))
        detail_end = max(rows)
        for row in range(detail_end + 1, sheet.max_row + 1):
            display_cells.extend(sheet[row])
        remaining = [
            identifier
            for identifier in ORIGINAL_IDENTIFIERS
            if any(identifier in text(cell.value) for cell in display_cells)
        ]
        if remaining:
            raise AnonymizeError(f"仍存在未匿名化的模型标识：{remaining}")
        return len(rows), detail_end
    finally:
        workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="匿名化三模型人工抽检工作簿中的评审员名称")
    parser.add_argument("--input", default=str(TARGET_FILE), help="待匿名化的 Excel 路径")
    parser.add_argument("--dry-run", action="store_true", help="只检查工作簿结构，不修改文件")
    return parser.parse_args()


def main() -> None:
    configure_console()
    args = parse_args()
    target = Path(args.input)
    if not target.is_file():
        raise AnonymizeError(f"找不到工作簿：{target}")
    workbook = load_workbook(target)
    try:
        sheet = workbook.active
        columns = find_columns(sheet)
        rows = detail_rows(sheet, columns[SOURCE_ROW_HEADER])
        numbers = numeric_snapshot(sheet)
        print(f"明细记录数：{len(rows)}")
        if args.dry_run:
            return

        for cell in sheet[1]:
            cell.value = replace_identifiers(cell.value)
        for row in rows:
            for column in (columns[FINAL_AUDITOR_HEADER], columns[FINAL_REASON_HEADER]):
                cell = sheet.cell(row, column)
                cell.value = replace_identifiers(cell.value)
        detail_end = max(rows)
        for row in range(detail_end + 1, sheet.max_row + 1):
            for cell in sheet[row]:
                cell.value = replace_identifiers(cell.value)
        workbook.save(target)
    finally:
        workbook.close()

    verified_rows, _ = verify_anonymized(target, numbers)
    print(f"已覆盖匿名化：{target}")
    print(f"已验证明细记录数：{verified_rows}")


if __name__ == "__main__":
    try:
        main()
    except AnonymizeError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(2)
