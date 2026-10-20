# -*- coding: utf-8 -*-
"""Shared configuration and utilities for the DeepSeek audit commands."""

from __future__ import annotations

import importlib
import os
import sys
from pathlib import Path
from typing import Any

from openpyxl.styles import Font, PatternFill


INPUT_DIR = Path(r"D:\测试用例模板\新题库\测试结果\成品题库已回答")
WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SEED = 20260727
DEFAULT_WORKERS = 4
RISK_HEADER = "风险大类/场景大类"
QUESTION_HEADER = "测试问题"
REPLY_HEADER = "模型回复"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
FAIL_FONT = Font(color="9C0006")
SUMMARY_FILL = PatternFill("solid", fgColor="D9EAF7")


class AuditError(RuntimeError):
    pass


def configure_console() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


def text(value: Any) -> str:
    return str(value or "").strip()


def deepseek_config(model_override: str | None) -> tuple[str, str, str]:
    api_key = os.getenv("DEEPSEEK_API_KEY", "").strip()
    base_url = os.getenv("DEEPSEEK_API_URL", "").strip()
    model = (model_override or os.getenv("DEEPSEEK_PRO_MODEL", "").strip() or "deepseek-v4-pro").strip()
    if not api_key or not base_url:
        if str(WORKSPACE_ROOT) not in sys.path:
            sys.path.insert(0, str(WORKSPACE_ROOT))
        for module_name in ("config11", "config"):
            try:
                module = importlib.import_module(module_name)
                configured_key = str(getattr(module, "DEEPSEEK_API_KEY", "")).strip()
                configured_url = str(getattr(module, "DEEPSEEK_API_URL", "")).strip()
            except ImportError:
                continue
            if configured_key and configured_url:
                api_key = api_key or configured_key
                base_url = base_url or configured_url
                break
    if not api_key or not base_url:
        raise AuditError("DeepSeek API 配置为空")
    return api_key, base_url.rstrip("/"), model


def remove_secret(value: str, api_key: str) -> str:
    return value.replace(api_key, "***") if api_key else value
