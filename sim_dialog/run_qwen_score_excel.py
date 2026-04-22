from __future__ import annotations

import json
import re
import sys
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:
    from . import score_dialog_config as score_config
except ImportError:
    import score_dialog_config as score_config


WINDOWS_INVALID_FILENAME_CHARS = '<>:"/\\|?*'
REQUEST_RETRY_TIMES = 3
REQUEST_RETRY_BACKOFF_SEC = 1.5
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


def sanitize_filename_part(name: str) -> str:
    if not name:
        return "output"
    sanitized = "".join("_" if ch in WINDOWS_INVALID_FILENAME_CHARS else ch for ch in name)
    sanitized = sanitized.strip().strip(".")
    return sanitized or "output"


def resolve_config_path(raw_path: str) -> Path:
    path = Path(raw_path).expanduser()
    if not path.is_absolute():
        path = (PROJECT_ROOT / path).resolve()
    return path


def resolve_input_dir() -> Path:
    manual_path = str(getattr(score_config, "ACTIVE_INPUT_PATH", "") or "").strip()
    if not manual_path:
        raise ValueError(
            "ACTIVE_INPUT_PATH is empty in score_dialog_config.py. "
            "Please fill in score_dialog_config.py -> ACTIVE_INPUT_PATH with an Excel folder."
        )

    path = resolve_config_path(manual_path)
    if not path.exists():
        raise FileNotFoundError(f"Configured ACTIVE_INPUT_PATH does not exist: {path}")
    if not path.is_dir():
        raise NotADirectoryError(f"Configured ACTIVE_INPUT_PATH must be an Excel folder: {path}")
    return path


def natural_sort_key(path: Path) -> tuple[Any, ...]:
    return tuple(
        int(part) if part.isdigit() else part.casefold()
        for part in re.split(r"(\d+)", path.name)
    )


def collect_input_excels(input_dir: Path) -> list[Path]:
    return sorted(
        (
            path
            for path in input_dir.iterdir()
            if path.is_file()
            and path.suffix.casefold() in EXCEL_SUFFIXES
            and not path.name.startswith("~$")
        ),
        key=natural_sort_key,
    )


def build_batch_output_dir(timestamp: str) -> Path:
    output_root = str(getattr(score_config, "SCORE_OUTPUT_ROOT", "outputs/score_outputs") or "").strip()
    if not output_root:
        raise ValueError(
            "SCORE_OUTPUT_ROOT is empty in score_dialog_config.py. "
            "Please fill in score_dialog_config.py -> SCORE_OUTPUT_ROOT."
        )
    return resolve_config_path(output_root) / timestamp


def build_output_path(input_path: Path, output_dir: Path) -> Path:
    base_name = sanitize_filename_part(str(getattr(score_config, "OUTPUT_BASENAME", "qwen_score")))
    source_stem = sanitize_filename_part(input_path.stem)
    timestamp = sanitize_filename_part(output_dir.name)
    return output_dir / f"{base_name}_{source_stem}_{timestamp}.xlsx"


def validate_dimensions() -> list[dict[str, Any]]:
    dimensions = getattr(score_config, "DIMENSIONS", None)
    if not isinstance(dimensions, list):
        raise ValueError("DIMENSIONS must be a list.")
    if len(dimensions) != 3:
        raise ValueError(f"DIMENSIONS must contain exactly 3 items, got {len(dimensions)}.")

    validated: list[dict[str, Any]] = []
    default_model = str(getattr(score_config, "DEFAULT_QWEN_MODEL", "qwen-plus")).strip()
    for idx, raw_dimension in enumerate(dimensions, start=1):
        if not isinstance(raw_dimension, dict):
            raise ValueError(f"DIMENSIONS[{idx}] must be a dict.")

        name = raw_dimension.get("name")
        model = raw_dimension.get("model", default_model)
        system_prompt = raw_dimension.get("system_prompt")
        score_field = raw_dimension.get("score_field", "final_score")

        if not isinstance(name, str) or not name.strip():
            raise ValueError(f"DIMENSIONS[{idx}].name must be a non-empty string.")
        if not isinstance(model, str) or not model.strip():
            raise ValueError(f"DIMENSIONS[{idx}].model must be a non-empty string.")
        if not isinstance(system_prompt, str) or not system_prompt.strip():
            raise ValueError(f"DIMENSIONS[{idx}].system_prompt must be a non-empty string.")
        if not isinstance(score_field, str) or not score_field.strip():
            raise ValueError(f"DIMENSIONS[{idx}].score_field must be a non-empty string.")

        dimension = dict(raw_dimension)
        dimension["name"] = name.strip()
        dimension["system_prompt"] = system_prompt.strip()
        dimension["model"] = model.strip()
        dimension["score_field"] = score_field.strip()
        validated.append(dimension)

    return validated


def validate_shared_prompt_config() -> tuple[str, str]:
    assistant_prompt = getattr(score_config, "ASSISTANT_SYSTEM_PROMPT", None)
    packed_template = getattr(score_config, "PACKED_USER_MESSAGE_TEMPLATE", None)

    if not isinstance(assistant_prompt, str) or not assistant_prompt.strip():
        raise ValueError("ASSISTANT_SYSTEM_PROMPT must be a non-empty string.")
    if not isinstance(packed_template, str) or not packed_template.strip():
        raise ValueError("PACKED_USER_MESSAGE_TEMPLATE must be a non-empty string.")
    if "{assistant_prompt}" not in packed_template:
        raise ValueError("PACKED_USER_MESSAGE_TEMPLATE must contain '{assistant_prompt}'.")
    if "{dialogue_text}" not in packed_template:
        raise ValueError("PACKED_USER_MESSAGE_TEMPLATE must contain '{dialogue_text}'.")

    return assistant_prompt.strip(), packed_template.strip()


def build_packed_user_message(
    assistant_prompt: str,
    dialogue_text: str,
    template: str,
) -> str:
    return template.format(
        assistant_prompt=assistant_prompt,
        dialogue_text=dialogue_text,
    ).strip()


def select_worksheet_name(workbook: Any) -> str:
    preferred_name = str(getattr(score_config, "WORKSHEET_NAME", "dialog_jsonl") or "").strip()
    if preferred_name and preferred_name in workbook.sheetnames:
        return preferred_name
    return workbook.sheetnames[0]


def normalize_message_content(content: Any) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, str):
                text = item.strip()
                if text:
                    parts.append(text)
                continue
            if isinstance(item, dict):
                text = item.get("text")
                if isinstance(text, str) and text.strip():
                    parts.append(text.strip())
                    continue
            text = str(item).strip()
            if text:
                parts.append(text)
        return "\n".join(parts).strip()
    if content is None:
        return ""
    return str(content).strip()


def load_sessions_from_excel(input_path: Path) -> tuple[str, OrderedDict[Any, list[dict[str, str]]]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheet_name = select_worksheet_name(workbook)
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError(f"Input Excel is empty: {input_path}")

        header_map = {str(col).strip(): idx for idx, col in enumerate(header) if col is not None}
        for required_col in ("session_id", "json_line"):
            if required_col not in header_map:
                raise ValueError(
                    f"Missing required column '{required_col}' in sheet '{sheet_name}'. "
                    f"Available columns: {list(header_map.keys())}"
                )

        session_idx = header_map["session_id"]
        json_idx = header_map["json_line"]
        conversations: OrderedDict[Any, list[dict[str, str]]] = OrderedDict()

        for row_number, row in enumerate(rows, start=2):
            if row is None or all(cell in (None, "") for cell in row):
                continue

            session_id = row[session_idx] if session_idx < len(row) else None
            json_line = row[json_idx] if json_idx < len(row) else None

            if session_id in (None, ""):
                raise ValueError(f"Row {row_number}: session_id is empty.")
            if not isinstance(json_line, str) or not json_line.strip():
                raise ValueError(f"Row {row_number}: json_line must be a non-empty JSON string.")

            try:
                payload = json.loads(json_line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Row {row_number}: invalid json_line JSON: {exc}") from exc

            if not isinstance(payload, dict):
                raise ValueError(f"Row {row_number}: json_line must decode to an object.")

            role = payload.get("role")
            content = normalize_message_content(payload.get("content"))
            if not isinstance(role, str) or not role.strip():
                raise ValueError(f"Row {row_number}: message.role must be a non-empty string.")
            if not content:
                raise ValueError(f"Row {row_number}: message.content is empty after normalization.")

            conversations.setdefault(session_id, []).append(
                {
                    "role": role.strip(),
                    "content": content,
                }
            )

        if not conversations:
            raise ValueError(f"No valid conversation rows found in: {input_path}")
        return sheet_name, conversations
    finally:
        workbook.close()


def render_dialogue_text(messages: list[dict[str, str]]) -> str:
    lines: list[str] = []
    for index, message in enumerate(messages, start=1):
        role = message["role"]
        content = message["content"].replace("\r\n", "\n").strip()
        lines.append(f"[{index}] {role}: {content}")
    return "\n".join(lines)


def extract_response_text(response: Any) -> str:
    message = response.choices[0].message
    content = getattr(message, "content", None)
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            text = getattr(item, "text", None)
            if isinstance(text, str) and text.strip():
                parts.append(text.strip())
        merged = "\n".join(parts).strip()
        if merged:
            return merged
    refusal = getattr(message, "refusal", None)
    if isinstance(refusal, str) and refusal.strip():
        return refusal.strip()
    return ""


def strip_code_fence(text: str) -> str:
    stripped = text.strip()
    if stripped.startswith("```") and stripped.endswith("```"):
        stripped = re.sub(r"^```[a-zA-Z0-9_-]*\s*", "", stripped, count=1)
        stripped = re.sub(r"\s*```$", "", stripped, count=1)
    return stripped.strip()


def extract_json_object_text(text: str) -> str:
    stripped = strip_code_fence(text)
    start = stripped.find("{")
    end = stripped.rfind("}")
    if start == -1 or end == -1 or end <= start:
        return stripped
    return stripped[start : end + 1].strip()


def repair_invalid_json_escapes(text: str) -> str:
    valid_escape_starts = {'"', "\\", "/", "b", "f", "n", "r", "t", "u"}
    repaired: list[str] = []
    in_string = False
    escape = False

    for ch in text:
        if not in_string:
            repaired.append(ch)
            if ch == '"':
                in_string = True
            continue

        if escape:
            if ch in valid_escape_starts:
                repaired.append(ch)
            else:
                # Preserve the original backslash as a literal inside the JSON string.
                repaired.append("\\")
                repaired.append(ch)
            escape = False
            continue

        if ch == "\\":
            repaired.append(ch)
            escape = True
            continue

        repaired.append(ch)
        if ch == '"':
            in_string = False

    return "".join(repaired)


def normalize_score_value(value: Any) -> int | float:
    if isinstance(value, bool):
        raise ValueError("Score value cannot be bool.")
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            raise ValueError("Score value is empty.")
        if re.fullmatch(r"-?\d+", stripped):
            return int(stripped)
        if re.fullmatch(r"-?\d+\.\d+", stripped):
            return float(stripped)
    raise ValueError(f"Unsupported score value: {value!r}")


def parse_score_response(raw_text: str, score_field: str) -> dict[str, Any]:
    if not raw_text.strip():
        raise ValueError("Model response is empty.")

    candidate = extract_json_object_text(raw_text)
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError as exc:
        repaired_candidate = repair_invalid_json_escapes(candidate)
        if repaired_candidate != candidate:
            try:
                parsed = json.loads(repaired_candidate)
            except json.JSONDecodeError as repaired_exc:
                raise ValueError(f"Failed to parse JSON response: {repaired_exc}") from repaired_exc
        else:
            raise ValueError(f"Failed to parse JSON response: {exc}") from exc

    if not isinstance(parsed, dict):
        raise ValueError("Parsed response is not a JSON object.")
    if score_field not in parsed:
        raise KeyError(f"Missing score field '{score_field}' in parsed response.")

    parsed["_normalized_score"] = normalize_score_value(parsed[score_field])
    return parsed


def build_client() -> OpenAI:
    base_url = str(getattr(score_config, "QWEN_BASE_URL", "") or "").strip()
    api_key = str(getattr(score_config, "QWEN_API_KEY", "") or "").strip()
    timeout = float(getattr(score_config, "REQUEST_TIMEOUT_SEC", 120.0) or 120.0)

    if not base_url:
        raise ValueError("QWEN_BASE_URL is empty in score_dialog_config.py")
    if not api_key:
        raise ValueError("QWEN_API_KEY is empty in score_dialog_config.py")

    return OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=timeout,
        # Disable SDK-level retries so the script's own retry loop remains the only retry layer.
        max_retries=0,
    )


def should_use_streaming() -> bool:
    use_streaming = getattr(score_config, "USE_STREAMING", False)
    if not isinstance(use_streaming, bool):
        raise ValueError("USE_STREAMING must be a boolean in score_dialog_config.py")
    return use_streaming


def extract_stream_response_text(stream: Any) -> str:
    content_parts: list[str] = []
    reasoning_parts: list[str] = []

    try:
        for chunk in stream:
            choices = getattr(chunk, "choices", None) or []
            if not choices:
                continue

            delta = getattr(choices[0], "delta", None)
            if delta is None:
                continue

            content_piece = getattr(delta, "content", None)
            reasoning_piece = getattr(delta, "reasoning_content", None)

            if content_piece:
                content_parts.append(str(content_piece))
            if reasoning_piece:
                reasoning_parts.append(str(reasoning_piece))
    finally:
        if stream is not None and hasattr(stream, "close"):
            stream.close()

    merged_content = "".join(content_parts).strip()
    if merged_content:
        return merged_content

    merged_reasoning = "".join(reasoning_parts).strip()
    if merged_reasoning:
        return merged_reasoning

    return ""


def request_dimension_score(
    client: OpenAI,
    dimension: dict[str, Any],
    packed_user_message: str,
) -> dict[str, Any]:
    dim_name = dimension["name"]
    use_streaming = should_use_streaming()
    messages: list[dict[str, str]] = [
        {"role": "system", "content": dimension["system_prompt"]},
        {"role": "user", "content": packed_user_message},
    ]

    request: dict[str, Any] = {
        "model": dimension["model"],
        "messages": messages,
    }

    if "temperature" in dimension:
        request["temperature"] = dimension["temperature"]
    if "max_tokens" in dimension:
        request["max_tokens"] = dimension["max_tokens"]

    extra_body = dimension.get("extra_body")
    if isinstance(extra_body, dict) and extra_body:
        request["extra_body"] = extra_body
    if use_streaming:
        request["stream"] = True
        request["stream_options"] = {"include_usage": True}

    last_error: Exception | None = None
    raw_text = ""
    latency_ms = 0.0
    for attempt in range(1, REQUEST_RETRY_TIMES + 1):
        started = time.perf_counter()
        print(
            f"  [DIM] {dim_name}: attempt={attempt}/{REQUEST_RETRY_TIMES} sending request...",
            flush=True,
        )
        try:
            response = client.chat.completions.create(**request)
            latency_ms = round((time.perf_counter() - started) * 1000, 2)
            if use_streaming:
                raw_text = extract_stream_response_text(response)
            else:
                raw_text = extract_response_text(response)
            parsed = parse_score_response(raw_text, dimension["score_field"])
            clean_parsed = {key: value for key, value in parsed.items() if key != "_normalized_score"}

            return {
                "status": "success",
                "error_message": "",
                "latency_ms": latency_ms,
                "raw_response_text": raw_text,
                "parsed": clean_parsed,
            }
        except KeyError as exc:
            last_error = exc
            latency_ms = round((time.perf_counter() - started) * 1000, 2)
            return {
                "status": "missing_score_field",
                "error_message": str(exc),
                "latency_ms": latency_ms,
                "raw_response_text": raw_text,
                "parsed": None,
            }
        except ValueError as exc:
            last_error = exc
            latency_ms = round((time.perf_counter() - started) * 1000, 2)
            return {
                "status": "parse_error",
                "error_message": str(exc),
                "latency_ms": latency_ms,
                "raw_response_text": raw_text,
                "parsed": None,
            }
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            latency_ms = round((time.perf_counter() - started) * 1000, 2)
            if attempt < REQUEST_RETRY_TIMES:
                backoff_sec = round(REQUEST_RETRY_BACKOFF_SEC * attempt, 2)
                print(
                    f"  [DIM] {dim_name}: attempt={attempt}/{REQUEST_RETRY_TIMES} failed, "
                    f"latency_ms={latency_ms}, error={exc}. Retrying in {backoff_sec}s...",
                    flush=True,
                )
                time.sleep(backoff_sec)
                continue
            print(
                f"  [DIM] {dim_name}: attempt={attempt}/{REQUEST_RETRY_TIMES} failed, "
                f"latency_ms={latency_ms}, error={exc}. No more retries.",
                flush=True,
            )

    return {
        "status": "request_error",
        "error_message": str(last_error) if last_error else "Unknown request error.",
        "latency_ms": latency_ms,
        "raw_response_text": "",
        "parsed": None,
    }


def serialize_excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    return value


def build_sheet_name(name: str, used_names: set[str]) -> str:
    sanitized = re.sub(r"[\\/*?:\[\]]", "_", name).strip()
    if not sanitized:
        sanitized = "sheet"
    sanitized = sanitized[:31]

    candidate = sanitized
    suffix = 1
    while candidate in used_names:
        tail = f"_{suffix}"
        candidate = f"{sanitized[: 31 - len(tail)]}{tail}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def collect_columns(rows: list[dict[str, Any]]) -> list[str]:
    columns = ["session_id"]
    seen = {"session_id"}
    for row in rows:
        for key in row.keys():
            if key not in seen:
                columns.append(key)
                seen.add(key)
    return columns


def build_summary_columns(
    dimensions: list[dict[str, Any]],
    summary_field_order: dict[str, list[str]],
) -> list[str]:
    columns = ["session_id"]
    for dimension in dimensions:
        columns.extend(summary_field_order.get(dimension["name"], []))
    return columns


def write_excel(
    output_path: Path,
    summary_rows: list[dict[str, Any]],
    summary_columns: list[str],
    dimensions: list[dict[str, Any]],
    dimension_rows: dict[str, list[dict[str, Any]]],
    error_rows: list[dict[str, Any]],
) -> None:
    used_sheet_names: set[str] = {"summary"}
    sheet_name_map = {
        dimension["name"]: build_sheet_name(dimension["name"], used_sheet_names)
        for dimension in dimensions
    }
    if error_rows:
        error_sheet_name = build_sheet_name("errors", used_sheet_names)
    else:
        error_sheet_name = ""

    try:
        import pandas as pd  # type: ignore

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            summary_df = pd.DataFrame(summary_rows, columns=summary_columns)
            summary_df.to_excel(writer, sheet_name="summary", index=False)

            for dimension in dimensions:
                dim_name = dimension["name"]
                rows = dimension_rows.get(dim_name, [])
                columns = collect_columns(rows)
                df = pd.DataFrame(rows, columns=columns)
                df.to_excel(writer, sheet_name=sheet_name_map[dim_name], index=False)

            if error_rows:
                error_columns = collect_columns(error_rows)
                error_df = pd.DataFrame(error_rows, columns=error_columns)
                error_df.to_excel(writer, sheet_name=error_sheet_name, index=False)
        return
    except ModuleNotFoundError:
        pass

    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = "summary"
    summary_ws.append(summary_columns)
    for row in summary_rows:
        summary_ws.append([row.get(column, "") for column in summary_columns])

    for dimension in dimensions:
        dim_name = dimension["name"]
        rows = dimension_rows.get(dim_name, [])
        columns = collect_columns(rows)

        worksheet = workbook.create_sheet(sheet_name_map[dim_name])

        worksheet.append(columns)
        for row in rows:
            worksheet.append([row.get(column, "") for column in columns])

    if error_rows:
        error_columns = collect_columns(error_rows)
        error_ws = workbook.create_sheet(error_sheet_name)
        error_ws.append(error_columns)
        for row in error_rows:
            error_ws.append([row.get(column, "") for column in error_columns])

    workbook.save(output_path)


def process_one_excel(
    input_path: Path,
    output_path: Path,
    client: OpenAI,
    dimensions: list[dict[str, Any]],
    assistant_prompt: str,
    packed_template: str,
) -> None:
    print(f"[INFO] Input Excel: {input_path}")
    print(f"[INFO] Output Excel: {output_path}")

    selected_sheet, conversations = load_sessions_from_excel(input_path)
    print(f"[INFO] Worksheet: {selected_sheet}")
    print(f"[INFO] Sessions found: {len(conversations)}")

    summary_rows: list[dict[str, Any]] = []
    summary_field_order: dict[str, list[str]] = {dimension["name"]: [] for dimension in dimensions}
    dimension_rows: dict[str, list[dict[str, Any]]] = {dimension["name"]: [] for dimension in dimensions}
    error_rows: list[dict[str, Any]] = []

    total_sessions = len(conversations)
    for session_index, (session_id, messages) in enumerate(conversations.items(), start=1):
        dialogue_text = render_dialogue_text(messages)
        packed_user_message = build_packed_user_message(
            assistant_prompt=assistant_prompt,
            dialogue_text=dialogue_text,
            template=packed_template,
        )
        summary_row: dict[str, Any] = {"session_id": session_id}

        print(f"[INFO] Session {session_index}/{total_sessions} -> session_id={session_id}")
        for dimension in dimensions:
            result = request_dimension_score(
                client=client,
                dimension=dimension,
                packed_user_message=packed_user_message,
            )

            dim_name = dimension["name"]
            parsed = result.get("parsed")
            if isinstance(parsed, dict):
                row: dict[str, Any] = {"session_id": session_id}
                for key, value in parsed.items():
                    serialized_value = serialize_excel_value(value)
                    row[key] = serialized_value
                    summary_key = f"{dim_name}_{key}"
                    summary_row[summary_key] = serialized_value
                    if summary_key not in summary_field_order[dim_name]:
                        summary_field_order[dim_name].append(summary_key)
                dimension_rows[dim_name].append(row)

            if result["status"] != "success":
                error_rows.append(
                    {
                        "session_id": session_id,
                        "dimension_name": dim_name,
                        "status": result["status"],
                        "error_message": result["error_message"],
                        "raw_response_text": result["raw_response_text"],
                    }
                )

            print(
                f"  [DIM] {dim_name}: status={result['status']}, "
                f"latency_ms={result['latency_ms']}"
            )

        summary_rows.append(summary_row)

    summary_columns = build_summary_columns(dimensions, summary_field_order)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_excel(output_path, summary_rows, summary_columns, dimensions, dimension_rows, error_rows)
    print(f"[DONE] Score Excel saved to: {output_path}")


def main() -> int:
    input_dir = resolve_input_dir()
    input_excels = collect_input_excels(input_dir)
    if not input_excels:
        print(f"[ERROR] No .xlsx/.xlsm files found in input folder: {input_dir}")
        return 1

    dimensions = validate_dimensions()
    assistant_prompt, packed_template = validate_shared_prompt_config()
    use_streaming = should_use_streaming()
    batch_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = build_batch_output_dir(batch_timestamp)
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"[INFO] Input folder: {input_dir}")
    print(f"[INFO] Output folder: {output_dir}")
    print(f"[INFO] Excel files found: {len(input_excels)}")
    print(f"[INFO] Active dimensions: {[dimension['name'] for dimension in dimensions]}")
    print(f"[INFO] Request mode: {'streaming' if use_streaming else 'non-stream'}")

    client = build_client()
    failed_files: list[tuple[Path, str]] = []
    total_files = len(input_excels)
    for file_index, input_path in enumerate(input_excels, start=1):
        output_path = build_output_path(input_path, output_dir)
        print(
            f"[INFO] File {file_index}/{total_files}: {input_path.name} -> {output_path.name}"
        )
        try:
            process_one_excel(
                input_path=input_path,
                output_path=output_path,
                client=client,
                dimensions=dimensions,
                assistant_prompt=assistant_prompt,
                packed_template=packed_template,
            )
        except Exception as exc:  # noqa: BLE001
            failed_files.append((input_path, str(exc)))
            print(f"[ERROR] Failed to score {input_path}: {exc}")
            continue

    succeeded_count = total_files - len(failed_files)
    print(
        f"[DONE] Batch scoring finished. succeeded={succeeded_count}, "
        f"failed={len(failed_files)}, output_folder={output_dir}"
    )
    if failed_files:
        for input_path, error_message in failed_files:
            print(f"[ERROR] Failed file: {input_path} | {error_message}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
