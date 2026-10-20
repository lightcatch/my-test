from collections import defaultdict
from pathlib import Path
import csv
import random
import shutil
import wave

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_ROOT = Path(r"\\Desktop-p5q0of7\公开数据集\cv-corpus-24.0-2025-12-05")
OUT_ROOT = Path("downloaded_audio") / "integrated_dialects"
SAMPLE_SIZE = 200
RANDOM_SEED = 20260625

LOCALES = {
    "yue": "\u7ca4\u8bed",
    "en": "\u82f1\u8bed",
}

COL_ID = "\u7f16\u53f7"
COL_LANG = "\u8bed\u8a00"
COL_TEXT = "\u6587\u672c"
COL_FILE = "\u6587\u4ef6\u540d"
COL_DURATION = "\u65f6\u957f"

FIELDNAMES = [
    COL_ID,
    COL_LANG,
    COL_TEXT,
    "saved_path",
    COL_FILE,
    COL_DURATION,
    "id",
    "speaker_id",
    "sample_index",
    "dialect",
]


def wav_duration(path):
    try:
        with wave.open(str(path), "rb") as wav_file:
            return round(wav_file.getnframes() / wav_file.getframerate(), 3)
    except wave.Error:
        return ""


def load_validated_rows(locale):
    locale_root = SOURCE_ROOT / locale
    wav_dir = locale_root / "clips_wav_16k"
    validated_path = locale_root / "validated.tsv"

    wav_by_name = {path.name: path for path in wav_dir.glob("*.wav")}
    rows_by_speaker = defaultdict(list)

    with validated_path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            wav_name = Path(row["path"]).with_suffix(".wav").name
            source_path = wav_by_name.get(wav_name)
            if source_path is None:
                continue

            speaker_id = row.get("client_id", "")
            rows_by_speaker[speaker_id].append(
                {
                    "source_path": source_path,
                    "speaker_id": speaker_id,
                    "id": source_path.stem,
                    "text": row.get("sentence", ""),
                    "dialect": locale,
                }
            )

    return rows_by_speaker


def balanced_sample(rows_by_speaker, locale):
    rng = random.Random(f"{RANDOM_SEED}:{locale}")
    pools = {}
    for speaker_id, rows in rows_by_speaker.items():
        shuffled = list(rows)
        rng.shuffle(shuffled)
        pools[speaker_id] = shuffled

    selected = []
    speaker_ids = list(pools)
    rng.shuffle(speaker_ids)

    while len(selected) < SAMPLE_SIZE:
        progressed = False
        for speaker_id in speaker_ids:
            pool = pools[speaker_id]
            if not pool:
                continue
            selected.append(pool.pop())
            progressed = True
            if len(selected) >= SAMPLE_SIZE:
                break
        if not progressed:
            break

    if len(selected) < SAMPLE_SIZE:
        raise RuntimeError(f"{locale} only has {len(selected)} usable rows")

    return selected


def write_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(FIELDNAMES)

    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [14, 12, 42, 36, 34, 10, 28, 24, 14, 16]
    for column_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=column_index).column_letter].width = width
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([row.get(field, "") for field in FIELDNAMES])

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def build_locale(locale, language):
    rows_by_speaker = load_validated_rows(locale)
    selected = balanced_sample(rows_by_speaker, locale)

    locale_out = OUT_ROOT / locale
    wav_out = locale_out / "wav"
    if locale_out.exists():
        shutil.rmtree(locale_out)
    wav_out.mkdir(parents=True, exist_ok=True)

    output_rows = []
    for index, item in enumerate(selected, start=1):
        source_path = item["source_path"]
        dst_name = f"{locale}_{index:04d}_{item['id']}{source_path.suffix or '.wav'}"
        dst_path = wav_out / dst_name
        shutil.copy2(source_path, dst_path)

        saved_path = dst_path.relative_to(locale_out).as_posix()
        output_rows.append(
            {
                COL_ID: f"{locale}_{index:04d}",
                COL_LANG: language,
                COL_TEXT: item["text"],
                "saved_path": saved_path,
                COL_FILE: dst_path.name,
                COL_DURATION: wav_duration(dst_path),
                "id": item["id"],
                "speaker_id": item["speaker_id"],
                "sample_index": index,
                "dialect": item["dialect"],
            }
        )

    write_excel(locale_out / f"{language}_\u6620\u5c04\u8868.xlsx", output_rows)
    with (locale_out / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(output_rows)

    speaker_counts = defaultdict(int)
    for row in output_rows:
        speaker_counts[row["speaker_id"]] += 1

    print(f"{language} ({locale}): {len(output_rows)}")
    print(f"speakers: {len(speaker_counts)}")
    print("top speaker counts:", sorted(speaker_counts.values(), reverse=True)[:10])


def main():
    for locale, language in LOCALES.items():
        build_locale(locale, language)
    print(f"output: {OUT_ROOT.resolve()}")


if __name__ == "__main__":
    main()
