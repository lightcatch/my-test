from collections import defaultdict
from pathlib import Path
import csv
import random
import shutil
import wave

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_ROOT = Path("downloaded_audio") / "sichuan"
OUT_ROOT = Path("downloaded_audio") / "integrated_dialects" / "sichuan"
WAV_DIR = OUT_ROOT / "wav"
SAMPLE_SIZE = 200
RANDOM_SEED = 20260625

LANGUAGE = "\u56db\u5ddd\u8bdd"
DIALECT = "Sichuan"

COL_ID = "\u7f16\u53f7"
COL_LANG = "\u8bed\u8a00"
COL_TEXT = "\u6587\u672c"
COL_FILE = "\u6587\u4ef6\u540d"
COL_DURATION = "\u65f6\u957f"

FIELDNAMES = [
    COL_ID,
    COL_LANG,
    COL_TEXT,
    "saved_path",
    COL_FILE,
    COL_DURATION,
    "id",
    "speaker_id",
    "sample_index",
    "dialect",
]


def wav_duration(path):
    try:
        with wave.open(str(path), "rb") as wav_file:
            return round(wav_file.getnframes() / wav_file.getframerate(), 3)
    except wave.Error:
        return ""


def read_manifest():
    rows_by_speaker = defaultdict(list)
    manifest_path = SOURCE_ROOT / "list.txt"

    with manifest_path.open("r", encoding="utf-8-sig", newline="") as f:
        for line in f:
            parts = line.rstrip("\r\n").split("\t")
            if len(parts) < 5:
                continue

            wav_name = parts[1]
            speaker_id = parts[2]
            text = parts[-1]
            source_path = SOURCE_ROOT / "wav" / speaker_id / wav_name
            if not source_path.exists():
                continue

            rows_by_speaker[speaker_id].append(
                {
                    "source_path": source_path,
                    "speaker_id": speaker_id,
                    "id": source_path.stem,
                    "text": text,
                }
            )

    return rows_by_speaker


def balanced_sample(rows_by_speaker):
    rng = random.Random(RANDOM_SEED)
    speaker_ids = sorted(rows_by_speaker)
    per_speaker = SAMPLE_SIZE // len(speaker_ids)
    remainder = SAMPLE_SIZE % len(speaker_ids)

    selected = []
    for index, speaker_id in enumerate(speaker_ids):
        rows = list(rows_by_speaker[speaker_id])
        rng.shuffle(rows)
        take_count = per_speaker + int(index < remainder)
        selected.extend(rows[:take_count])

    rng.shuffle(selected)
    return selected


def write_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(FIELDNAMES)

    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [14, 12, 42, 36, 34, 10, 28, 16, 14, 16]
    for column_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=column_index).column_letter].width = width
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([row.get(field, "") for field in FIELDNAMES])

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def main():
    rows_by_speaker = read_manifest()
    selected = balanced_sample(rows_by_speaker)
    if len(selected) != SAMPLE_SIZE:
        raise RuntimeError(f"expected {SAMPLE_SIZE} samples, got {len(selected)}")

    if OUT_ROOT.exists():
        shutil.rmtree(OUT_ROOT)
    WAV_DIR.mkdir(parents=True, exist_ok=True)

    output_rows = []
    for index, item in enumerate(selected, start=1):
        source_path = item["source_path"]
        dst_name = f"sichuan_{index:04d}_{item['id']}{source_path.suffix or '.wav'}"
        dst_path = WAV_DIR / dst_name
        shutil.copy2(source_path, dst_path)

        saved_path = dst_path.relative_to(OUT_ROOT).as_posix()
        output_rows.append(
            {
                COL_ID: f"sichuan_{index:04d}",
                COL_LANG: LANGUAGE,
                COL_TEXT: item["text"],
                "saved_path": saved_path,
                COL_FILE: dst_path.name,
                COL_DURATION: wav_duration(dst_path),
                "id": item["id"],
                "speaker_id": item["speaker_id"],
                "sample_index": index,
                "dialect": DIALECT,
            }
        )

    write_excel(OUT_ROOT / ("\u56db\u5ddd\u8bdd_\u6620\u5c04\u8868.xlsx"), output_rows)
    with (OUT_ROOT / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(output_rows)

    speaker_counts = defaultdict(int)
    for row in output_rows:
        speaker_counts[row["speaker_id"]] += 1

    print(f"done: {len(output_rows)}")
    for speaker_id, count in sorted(speaker_counts.items()):
        print(f"{speaker_id}: {count}")
    print(OUT_ROOT.resolve())


if __name__ == "__main__":
    main()
