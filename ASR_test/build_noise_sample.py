from collections import defaultdict
from pathlib import Path
import csv
import random
import shutil
import wave

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_ROOT = Path("downloaded_audio") / "noises"
OUT_ROOT = Path("downloaded_audio") / "integrated_dialects" / "noises"
WAV_DIR = OUT_ROOT / "wav"
SAMPLE_SIZE = 500
RANDOM_SEED = 20260625
LANGUAGE = "\u566a\u97f3"

COL_ID = "\u7f16\u53f7"
COL_LANG = "\u8bed\u8a00"
COL_TEXT = "\u6587\u672c"
COL_FILE = "\u6587\u4ef6\u540d"
COL_DURATION = "\u65f6\u957f"

FIELDNAMES = [
    COL_ID,
    COL_LANG,
    COL_TEXT,
    "saved_path",
    COL_FILE,
    COL_DURATION,
    "id",
    "speaker_id",
    "sample_index",
    "dialect",
]


def wav_duration(path):
    try:
        with wave.open(str(path), "rb") as wav_file:
            return round(wav_file.getnframes() / wav_file.getframerate(), 3)
    except wave.Error:
        return ""


def load_rows():
    rows_by_type = defaultdict(list)
    with (SOURCE_ROOT / "metadata.csv").open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            saved_path = row.get("saved_path") or ""
            if not saved_path:
                continue

            source_path = SOURCE_ROOT / saved_path
            if not source_path.exists():
                continue

            audio_type = row.get("audio_type") or saved_path.split("/", 1)[0]
            rows_by_type[audio_type].append(
                {
                    "source_path": source_path,
                    "audio_type": audio_type,
                    "text": row.get("text", ""),
                    "duration": row.get("duration_seconds", ""),
                    "id": row.get("name") or Path(saved_path).stem,
                }
            )

    return rows_by_type


def balanced_sample(rows_by_type):
    rng = random.Random(RANDOM_SEED)
    audio_types = sorted(rows_by_type)
    per_type = SAMPLE_SIZE // len(audio_types)
    remainder = SAMPLE_SIZE % len(audio_types)

    selected = []
    for index, audio_type in enumerate(audio_types):
        rows = list(rows_by_type[audio_type])
        rng.shuffle(rows)
        take_count = per_type + int(index < remainder)
        selected.extend(rows[:take_count])

    rng.shuffle(selected)
    if len(selected) != SAMPLE_SIZE:
        raise RuntimeError(f"expected {SAMPLE_SIZE} samples, got {len(selected)}")
    return selected


def write_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(FIELDNAMES)

    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [14, 12, 42, 36, 34, 10, 28, 16, 14, 24]
    for column_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=column_index).column_letter].width = width
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([row.get(field, "") for field in FIELDNAMES])

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def main():
    rows_by_type = load_rows()
    selected = balanced_sample(rows_by_type)

    if OUT_ROOT.exists():
        shutil.rmtree(OUT_ROOT)
    WAV_DIR.mkdir(parents=True, exist_ok=True)

    output_rows = []
    for index, item in enumerate(selected, start=1):
        source_path = item["source_path"]
        dst_name = f"noise_{index:04d}_{item['audio_type']}_{source_path.stem}{source_path.suffix or '.wav'}"
        dst_path = WAV_DIR / dst_name
        shutil.copy2(source_path, dst_path)

        saved_path = dst_path.relative_to(OUT_ROOT).as_posix()
        output_rows.append(
            {
                COL_ID: f"noise_{index:04d}",
                COL_LANG: LANGUAGE,
                COL_TEXT: item["text"],
                "saved_path": saved_path,
                COL_FILE: dst_path.name,
                COL_DURATION: item["duration"] or wav_duration(dst_path),
                "id": item["id"],
                "speaker_id": "",
                "sample_index": index,
                "dialect": item["audio_type"],
            }
        )

    write_excel(OUT_ROOT / ("\u566a\u97f3_\u6620\u5c04\u8868.xlsx"), output_rows)
    with (OUT_ROOT / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(output_rows)

    counts = defaultdict(int)
    for row in output_rows:
        counts[row["dialect"]] += 1

    print(f"done: {len(output_rows)}")
    for audio_type, count in sorted(counts.items()):
        print(f"{audio_type}: {count}")
    print(OUT_ROOT.resolve())


if __name__ == "__main__":
    main()
