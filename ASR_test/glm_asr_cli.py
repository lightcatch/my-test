from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from scipy.io import wavfile
from scipy.signal import resample_poly
from transformers import AutoProcessor, GlmAsrForConditionalGeneration


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_MODEL_ID = "zai-org/GLM-ASR-Nano-2512"
DEFAULT_REPETITION_PENALTY = 1.1
DEFAULT_TRANSCRIPTION_PROMPT = "Transcribe the input speech."
CUDA_REINSTALL_COMMAND = (
    r"uv pip install --reinstall --python .\.venv\Scripts\python.exe "
    r"--torch-backend cu128 torch==2.10.0"
)


@dataclass(frozen=True)
class SmokeTestSample:
    audio_file: str
    reference_text: str


SMOKE_TEST_SAMPLES = [
    SmokeTestSample(
        audio_file="3356_bdae8b45c8f045698678cd701468ca5e.wav",
        reference_text="你们这边就是就说他的呃工作的成果。",
    ),
    SmokeTestSample(
        audio_file="3359_1f2a0d5a4f7f4df0b1b06ed7cdf71571.wav",
        reference_text="其他本身他也要是相。",
    ),
    SmokeTestSample(
        audio_file="3361_8061d2ab884b48b0b6218ed9224fd365.wav",
        reference_text="对，这就是我知道啊，那才是你们提的这个意见是吧。",
    ),
]


class CliError(Exception):
    """Raised for user-facing errors that should not print a traceback."""


class HelpFormatter(argparse.ArgumentDefaultsHelpFormatter, argparse.RawTextHelpFormatter):
    """Show defaults while preserving manual line breaks in help text."""


def build_parser() -> argparse.ArgumentParser:
    common = argparse.ArgumentParser(add_help=False, formatter_class=HelpFormatter)
    common.add_argument(
        "--model-id",
        default=DEFAULT_MODEL_ID,
        help="Model id on Hugging Face.",
    )
    common.add_argument(
        "--device",
        choices=("auto", "cuda", "cpu"),
        default="auto",
        help=(
            "Execution device.\n"
            "- auto: prefer CUDA when available, otherwise CPU\n"
            "- cuda: require CUDA and fail fast if unavailable\n"
            "- cpu: force CPU inference"
        ),
    )
    common.add_argument(
        "--max-new-tokens",
        type=int,
        default=500,
        help="Maximum number of tokens to generate for each transcription.",
    )

    parser = argparse.ArgumentParser(
        description=(
            "Run local ASR inference with zai-org/GLM-ASR-Nano-2512.\n\n"
            "Common workflows:\n"
            "1. Transcribe one local WAV file\n"
            "2. Run the fixed three-sample smoke test\n"
            "3. Fill a new Excel file with ASR text for rows that contain audio links"
        ),
        epilog=(
            "Examples:\n"
            "  uv run python ASR_test\\glm_asr_cli.py transcribe --audio ASR_test\\audio\\4858_xxx.wav --device cuda\n"
            "  uv run python ASR_test\\glm_asr_cli.py smoke-test --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda\n"
            "  uv run python ASR_test\\glm_asr_cli.py fill-excel --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda\n"
            "  uv run python ASR_test\\glm_asr_cli.py fill-excel --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda --limit 5"
        ),
        formatter_class=HelpFormatter,
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    transcribe = subparsers.add_parser(
        "transcribe",
        parents=[common],
        help="Transcribe a single local audio file.",
        description=(
            "Transcribe one local WAV file and print the recognized text to stdout.\n"
            "Exit code is 0 when the transcription is non-empty."
        ),
        epilog=(
            "Example:\n"
            "  uv run python ASR_test\\glm_asr_cli.py transcribe --audio ASR_test\\audio\\4858_8b3e8d457c6e4fe8b4223f2910317243.wav --device cuda"
        ),
        formatter_class=HelpFormatter,
    )
    transcribe.add_argument(
        "--audio",
        type=Path,
        required=True,
        help="Path to a local WAV file.",
    )

    smoke_test = subparsers.add_parser(
        "smoke-test",
        parents=[common],
        help="Run the fixed three-sample GPU smoke test.",
        description=(
            "Run the fixed three-sample smoke test against known local audio files.\n"
            "Writes an Excel report with reference text, prediction, elapsed time and status."
        ),
        epilog=(
            "Example:\n"
            "  uv run python ASR_test\\glm_asr_cli.py smoke-test --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda"
        ),
        formatter_class=HelpFormatter,
    )
    smoke_test.add_argument(
        "--excel",
        type=Path,
        default=SCRIPT_DIR / "chat_history.xlsx",
        help="Path to the chat history Excel file used to resolve the three fixed samples.",
    )
    smoke_test.add_argument(
        "--audio-dir",
        type=Path,
        default=SCRIPT_DIR / "audio",
        help="Directory containing exported WAV files.",
    )
    smoke_test.add_argument(
        "--sample-count",
        type=int,
        default=3,
        help="Number of fixed smoke-test samples to run. Maximum is 3.",
    )
    smoke_test.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Excel report output path.",
    )

    fill_excel = subparsers.add_parser(
        "fill-excel",
        parents=[common],
        help="Fill a new Excel file with ASR output for rows that contain audio links.",
        description=(
            "Read an Excel workbook, process the selected worksheet row by row, and append one new column\n"
            "containing ASR text for rows that contain an audio hyperlink or audio_file reference.\n\n"
            "The source workbook is never modified unless you explicitly point --output at the same path.\n"
            "Rows without audio stay blank in the ASR output column.\n"
            "If any row fails, the command still saves the output file but returns a non-zero exit code."
        ),
        epilog=(
            "Examples:\n"
            "  uv run python ASR_test\\glm_asr_cli.py fill-excel --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda\n"
            "  uv run python ASR_test\\glm_asr_cli.py fill-excel --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda --limit 5\n"
            "  uv run python ASR_test\\glm_asr_cli.py fill-excel --excel ASR_test\\chat_history.xlsx --audio-dir ASR_test\\audio --device cuda --output ASR_test\\results\\my_asr.xlsx"
        ),
        formatter_class=HelpFormatter,
    )
    fill_excel.add_argument(
        "--excel",
        type=Path,
        default=SCRIPT_DIR / "chat_history.xlsx",
        help="Path to the source Excel file.",
    )
    fill_excel.add_argument(
        "--audio-dir",
        type=Path,
        default=SCRIPT_DIR / "audio",
        help="Directory containing exported WAV files referenced by the workbook.",
    )
    fill_excel.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Excel output path.",
    )
    fill_excel.add_argument(
        "--sheet-index",
        type=int,
        default=0,
        help="Zero-based worksheet index to process.",
    )
    fill_excel.add_argument(
        "--asr-column",
        default="ASR识别文字",
        help="Column name to append at the end of the worksheet for ASR output.",
    )
    fill_excel.add_argument(
        "--skip-existing",
        action="store_true",
        help="Skip rows that already have a value in the ASR output column.",
    )
    fill_excel.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Only process the first N resolvable audio rows. Useful for spot checks.",
    )

    return parser


def resolve_device(requested_device: str) -> str:
    if requested_device == "auto":
        return "cuda" if torch.cuda.is_available() else "cpu"
    if requested_device == "cuda" and not torch.cuda.is_available():
        details = [
            "CUDA was requested, but torch.cuda.is_available() is False.",
            f"Current torch version: {torch.__version__}",
        ]
        if "+cpu" in torch.__version__ or not torch.backends.cuda.is_built():
            details.append("The current environment appears to be using a CPU-only torch build.")
        else:
            details.append("CUDA support is not available in the current Python environment.")
        details.append("Reinstall torch with CUDA support using:")
        details.append(CUDA_REINSTALL_COMMAND)
        raise CliError("\n".join(details))
    return requested_device


def get_default_smoke_test_output_path() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return SCRIPT_DIR / "results" / f"glm_asr_smoke_test_{timestamp}.xlsx"


def get_default_fill_excel_output_path() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return SCRIPT_DIR / "results" / f"chat_history_with_asr_{timestamp}.xlsx"


def get_model_dtype(device: str) -> torch.dtype:
    return torch.float16 if device == "cuda" else torch.float32


def load_processor_and_model(model_id: str, device: str) -> tuple[Any, GlmAsrForConditionalGeneration]:
    processor = AutoProcessor.from_pretrained(model_id)
    dtype = get_model_dtype(device)
    model = GlmAsrForConditionalGeneration.from_pretrained(
        model_id,
        torch_dtype=dtype,
        use_safetensors=True,
    )
    model = model.to(device)
    model.eval()
    return processor, model


def read_wav_as_float32(audio_path: Path, target_sampling_rate: int) -> np.ndarray:
    sampling_rate, audio = wavfile.read(audio_path)
    if audio.ndim > 1:
        audio = audio.mean(axis=1)

    if np.issubdtype(audio.dtype, np.integer):
        info = np.iinfo(audio.dtype)
        scale = float(max(abs(info.min), abs(info.max)))
        audio = audio.astype(np.float32) / scale
    elif np.issubdtype(audio.dtype, np.floating):
        audio = audio.astype(np.float32, copy=False)
    else:
        raise CliError(f"Unsupported WAV dtype for {audio_path}: {audio.dtype}")

    if sampling_rate != target_sampling_rate:
        audio = resample_poly(audio, target_sampling_rate, sampling_rate).astype(np.float32)

    return np.clip(audio, -1.0, 1.0).astype(np.float32, copy=False)


def prepare_inputs_for_model(inputs: Any, model: GlmAsrForConditionalGeneration) -> dict[str, Any]:
    prepared: dict[str, Any] = {}
    for name, value in inputs.items():
        if isinstance(value, torch.Tensor):
            if torch.is_floating_point(value):
                prepared[name] = value.to(device=model.device, dtype=model.dtype)
            else:
                prepared[name] = value.to(device=model.device)
        else:
            prepared[name] = value
    return prepared


def maybe_synchronize(device: str) -> None:
    if device == "cuda":
        torch.cuda.synchronize()


def transcribe_audio(
    processor: Any,
    model: GlmAsrForConditionalGeneration,
    audio_path: Path,
    max_new_tokens: int,
) -> str:
    target_sr = processor.feature_extractor.sampling_rate
    audio = read_wav_as_float32(audio_path, target_sr)
    # An explicit prompt is more stable than the processor default on short audio.
    inputs = processor.apply_transcription_request(audio, prompt=DEFAULT_TRANSCRIPTION_PROMPT)
    prepared = prepare_inputs_for_model(inputs, model)

    input_ids = prepared.get("input_ids")
    if not isinstance(input_ids, torch.Tensor):
        raise CliError("Processor output is missing input_ids; cannot decode generated text.")

    maybe_synchronize(model.device.type)
    with torch.inference_mode():
        outputs = model.generate(
            **prepared,
            do_sample=False,
            max_new_tokens=max_new_tokens,
            repetition_penalty=DEFAULT_REPETITION_PENALTY,
            eos_token_id=model.generation_config.eos_token_id,
            pad_token_id=processor.tokenizer.pad_token_id,
        )
    maybe_synchronize(model.device.type)

    decoded = processor.batch_decode(
        outputs[:, input_ids.shape[1] :],
        skip_special_tokens=True,
    )
    return decoded[0].strip() if decoded else ""


def validate_transcribe_args(args: argparse.Namespace) -> None:
    if not args.audio.exists():
        raise CliError(f"Audio file does not exist: {args.audio}")
    if args.max_new_tokens <= 0:
        raise CliError("--max-new-tokens must be a positive integer.")


def run_transcribe(args: argparse.Namespace) -> int:
    validate_transcribe_args(args)
    device = resolve_device(args.device)
    processor, model = load_processor_and_model(args.model_id, device)
    text = transcribe_audio(processor, model, args.audio, args.max_new_tokens)
    print(text)
    return 0 if text else 1


def select_smoke_test_rows(excel_path: Path, audio_dir: Path, sample_count: int) -> list[dict[str, Any]]:
    if sample_count < 1:
        raise CliError("--sample-count must be at least 1.")
    if sample_count > len(SMOKE_TEST_SAMPLES):
        raise CliError(f"--sample-count cannot exceed {len(SMOKE_TEST_SAMPLES)} for this smoke test.")
    if not excel_path.exists():
        raise CliError(f"Excel file does not exist: {excel_path}")
    if not audio_dir.exists():
        raise CliError(f"Audio directory does not exist: {audio_dir}")

    df = pd.read_excel(excel_path, sheet_name=0)
    required_columns = {"seq", "content", "audio_file"}
    missing = required_columns - set(df.columns)
    if missing:
        raise CliError(f"Excel file is missing required columns: {sorted(missing)}")

    indexed = df.set_index("audio_file", drop=False)
    selected_rows: list[dict[str, Any]] = []

    for sample in SMOKE_TEST_SAMPLES[:sample_count]:
        if sample.audio_file not in indexed.index:
            raise CliError(f"Smoke-test sample was not found in Excel: {sample.audio_file}")

        row = indexed.loc[sample.audio_file]
        if isinstance(row, pd.DataFrame):
            row = row.iloc[0]

        audio_path = audio_dir / sample.audio_file
        if not audio_path.exists():
            raise CliError(f"Smoke-test audio file does not exist: {audio_path}")

        reference_text = str(row["content"]).strip() if pd.notna(row["content"]) else sample.reference_text
        selected_rows.append(
            {
                "seq": int(row["seq"]) if pd.notna(row["seq"]) else None,
                "audio_file": sample.audio_file,
                "reference_text": reference_text or sample.reference_text,
                "audio_path": audio_path,
            }
        )

    return selected_rows


def run_smoke_test(args: argparse.Namespace) -> int:
    if args.max_new_tokens <= 0:
        raise CliError("--max-new-tokens must be a positive integer.")

    device = resolve_device(args.device)
    output_path = args.output or get_default_smoke_test_output_path()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    selected_rows = select_smoke_test_rows(args.excel, args.audio_dir, args.sample_count)
    processor, model = load_processor_and_model(args.model_id, device)

    results: list[dict[str, Any]] = []
    success_count = 0

    for row in selected_rows:
        prediction = ""
        error = ""
        status = "success"
        started = time.perf_counter()

        try:
            prediction = transcribe_audio(
                processor=processor,
                model=model,
                audio_path=row["audio_path"],
                max_new_tokens=args.max_new_tokens,
            )
            if not prediction:
                status = "error"
                error = "Prediction is empty."
            else:
                success_count += 1
        except Exception as exc:  # noqa: BLE001 - keep smoke test running and write a report.
            status = "error"
            error = str(exc)

        elapsed_sec = round(time.perf_counter() - started, 3)
        results.append(
            {
                "seq": row["seq"],
                "audio_file": row["audio_file"],
                "reference_text": row["reference_text"],
                "prediction": prediction,
                "prediction_non_empty": bool(prediction),
                "device": device,
                "torch_version": torch.__version__,
                "model_id": args.model_id,
                "elapsed_sec": elapsed_sec,
                "status": status,
                "error": error,
            }
        )

    result_df = pd.DataFrame(
        results,
        columns=[
            "seq",
            "audio_file",
            "reference_text",
            "prediction",
            "prediction_non_empty",
            "device",
            "torch_version",
            "model_id",
            "elapsed_sec",
            "status",
            "error",
        ],
    )
    result_df.to_excel(output_path, index=False)

    print(f"Device: {device}")
    print(f"Torch: {torch.__version__}")
    print(f"Processed: {len(results)}")
    print(f"Succeeded: {success_count}")
    print(f"Failed: {len(results) - success_count}")
    print(f"Result file: {output_path}")

    all_good = all(
        row["status"] == "success" and row["prediction_non_empty"]
        for row in results
    )
    return 0 if all_good else 1


def get_header_map(worksheet: Any) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for col in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=col).value
        if value not in (None, ""):
            header_map[str(value)] = col
    return header_map


def get_or_create_output_column(worksheet: Any, column_name: str) -> tuple[int, bool]:
    header_map = get_header_map(worksheet)
    if column_name in header_map:
        return header_map[column_name], False

    output_col = worksheet.max_column + 1
    worksheet.cell(row=1, column=output_col, value=column_name)
    return output_col, True


def cell_has_text(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def resolve_audio_path_for_excel_row(
    worksheet: Any,
    row_idx: int,
    header_map: dict[str, int],
    excel_dir: Path,
    audio_dir: Path,
) -> tuple[Path | None, str | None]:
    audio_file_col = header_map.get("audio_file")
    audio_link_col = header_map.get("audio_link")
    audio_file = None
    if audio_file_col is not None:
        raw_audio_file = worksheet.cell(row=row_idx, column=audio_file_col).value
        if cell_has_text(raw_audio_file):
            audio_file = str(raw_audio_file).strip()

    if audio_link_col is not None:
        link_cell = worksheet.cell(row=row_idx, column=audio_link_col)
        if link_cell.hyperlink and cell_has_text(link_cell.hyperlink.target):
            hyperlink_target = str(link_cell.hyperlink.target).strip()
            target_path = Path(hyperlink_target)
            candidates = []
            if target_path.is_absolute():
                candidates.append(target_path)
            else:
                candidates.append((excel_dir / target_path).resolve())
            if audio_file:
                candidates.append((audio_dir / audio_file).resolve())
            candidates.append((audio_dir / target_path.name).resolve())

            for candidate in candidates:
                if candidate.exists():
                    return candidate, None

            return None, f"Audio file not found for hyperlink target: {hyperlink_target}"

    if audio_file:
        candidate = (audio_dir / audio_file).resolve()
        if candidate.exists():
            return candidate, None
        return None, f"Audio file not found: {candidate}"

    return None, None


def validate_fill_excel_args(args: argparse.Namespace) -> None:
    if not args.excel.exists():
        raise CliError(f"Excel file does not exist: {args.excel}")
    if not args.audio_dir.exists():
        raise CliError(f"Audio directory does not exist: {args.audio_dir}")
    if args.max_new_tokens <= 0:
        raise CliError("--max-new-tokens must be a positive integer.")
    if args.sheet_index < 0:
        raise CliError("--sheet-index must be zero or greater.")
    if args.limit is not None and args.limit <= 0:
        raise CliError("--limit must be a positive integer when provided.")
    if not cell_has_text(args.asr_column):
        raise CliError("--asr-column must not be empty.")


def run_fill_excel(args: argparse.Namespace) -> int:
    validate_fill_excel_args(args)
    device = resolve_device(args.device)
    output_path = args.output or get_default_fill_excel_output_path()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(args.excel)
    if args.sheet_index >= len(workbook.worksheets):
        raise CliError(
            f"--sheet-index {args.sheet_index} is out of range for workbook with {len(workbook.worksheets)} sheet(s)."
        )

    worksheet = workbook.worksheets[args.sheet_index]
    header_map = get_header_map(worksheet)
    output_col, created_output_col = get_or_create_output_column(worksheet, args.asr_column)
    processor, model = load_processor_and_model(args.model_id, device)

    attempted = 0
    succeeded = 0
    failed = 0
    skipped_existing = 0
    failed_rows: list[tuple[int, str, str]] = []
    excel_dir = args.excel.resolve().parent
    audio_dir = args.audio_dir.resolve()

    for row_idx in range(2, worksheet.max_row + 1):
        if args.limit is not None and attempted >= args.limit:
            break

        output_cell = worksheet.cell(row=row_idx, column=output_col)
        if args.skip_existing and cell_has_text(output_cell.value):
            skipped_existing += 1
            continue

        audio_path, resolution_error = resolve_audio_path_for_excel_row(
            worksheet=worksheet,
            row_idx=row_idx,
            header_map=header_map,
            excel_dir=excel_dir,
            audio_dir=audio_dir,
        )
        if audio_path is None:
            if resolution_error:
                failed += 1
                failed_rows.append((row_idx, "", resolution_error))
                print(f"Row {row_idx}: {resolution_error}", file=sys.stderr)
            continue

        attempted += 1
        print(f"[{attempted}] Row {row_idx}: {audio_path.name}")

        try:
            text = transcribe_audio(
                processor=processor,
                model=model,
                audio_path=audio_path,
                max_new_tokens=args.max_new_tokens,
            )
            if text:
                output_cell.value = text
                succeeded += 1
            else:
                failed += 1
                failed_rows.append((row_idx, str(audio_path), "Prediction is empty."))
                print(f"Row {row_idx}: Prediction is empty for {audio_path}", file=sys.stderr)
        except Exception as exc:  # noqa: BLE001 - continue processing remaining rows.
            failed += 1
            failed_rows.append((row_idx, str(audio_path), str(exc)))
            print(f"Row {row_idx}: {audio_path} -> {exc}", file=sys.stderr)

    workbook.save(output_path)

    print(f"Device: {device}")
    print(f"Torch: {torch.__version__}")
    print(f"Sheet: {worksheet.title}")
    print(f"ASR column: {args.asr_column}")
    print(f"Created new column: {created_output_col}")
    print(f"Attempted audio rows: {attempted}")
    print(f"Succeeded: {succeeded}")
    print(f"Failed: {failed}")
    print(f"Skipped existing: {skipped_existing}")
    print(f"Result file: {output_path}")

    return 0 if failed == 0 else 1


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        if args.command == "transcribe":
            return run_transcribe(args)
        if args.command == "smoke-test":
            return run_smoke_test(args)
        if args.command == "fill-excel":
            return run_fill_excel(args)
        raise CliError(f"Unsupported command: {args.command}")
    except CliError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("Interrupted by user.", file=sys.stderr)
        return 130
    except Exception as exc:  # noqa: BLE001 - provide a concise user-facing failure.
        print(f"Unexpected error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
