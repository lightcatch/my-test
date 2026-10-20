from collections import defaultdict
from pathlib import Path
import csv
import random
import shutil
import wave

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_ROOT = Path(r"\\Desktop-p5q0of7\公开数据集\zh_audit_test")
TRANS_PATH = SOURCE_ROOT / "TRANS.txt"
OUT_ROOT = Path("downloaded_audio") / "integrated_dialects" / "mandarin"
WAV_DIR = OUT_ROOT / "wav"
SAMPLE_SIZE = 400
LANGUAGE = "普通话"
RANDOM_SEED = 20260625


FIELDNAMES = [
    "编号",
    "语言",
    "文本",
    "saved_path",
    "时长",
    "id",
    "speaker_id",
    "source_path",
    "source_filename",
    "source_transcription",
    "sample_index",
]


def wav_duration(path):
    try:
        with wave.open(str(path), "rb") as wav_file:
            return wav_file.getnframes() / wav_file.getframerate()
    except wave.Error:
        return None


def read_transcriptions():
    mapping = {}
    with TRANS_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            mapping[row["UtteranceID"]] = row
    return mapping


def collect_audio(transcriptions):
    by_speaker = defaultdict(list)
    for speaker_dir in sorted(SOURCE_ROOT.iterdir()):
        if not speaker_dir.is_dir():
            continue

        speaker_id = speaker_dir.name
        for wav_path in sorted(speaker_dir.glob("*.wav")):
            trans = transcriptions.get(wav_path.name, {})
            by_speaker[speaker_id].append(
                {
                    "source_path": wav_path,
                    "speaker_id": speaker_id,
                    "source_filename": wav_path.name,
                    "text": trans.get("Transcription", ""),
                }
            )

    return by_speaker


def balanced_sample(by_speaker):
    rng = random.Random(RANDOM_SEED)
    pools = {}
    for speaker_id, items in by_speaker.items():
        shuffled = list(items)
        rng.shuffle(shuffled)
        pools[speaker_id] = shuffled

    selected = []
    speaker_ids = list(pools)
    rng.shuffle(speaker_ids)

    while len(selected) < SAMPLE_SIZE:
        progressed = False
        for speaker_id in speaker_ids:
            pool = pools[speaker_id]
            if not pool:
                continue

            selected.append(pool.pop())
            progressed = True
            if len(selected) >= SAMPLE_SIZE:
                break

        if not progressed:
            break

    return selected


def write_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(FIELDNAMES)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(field, "") for field in FIELDNAMES])

    freeze_cell = "A2"
    ws.freeze_panes = freeze_cell
    widths = {
        "A": 12,
        "B": 12,
        "C": 36,
        "D": 32,
        "E": 10,
        "F": 28,
        "G": 14,
        "H": 78,
        "I": 34,
        "J": 36,
        "K": 14,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

    wb.save(path)


def main():
    if not SOURCE_ROOT.exists():
        raise FileNotFoundError(f"source not found: {SOURCE_ROOT}")

    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    WAV_DIR.mkdir(parents=True, exist_ok=True)

    transcriptions = read_transcriptions()
    by_speaker = collect_audio(transcriptions)
    selected = balanced_sample(by_speaker)

    if len(selected) < SAMPLE_SIZE:
        raise RuntimeError(f"only found {len(selected)} audio files, expected {SAMPLE_SIZE}")

    rows = []
    for index, item in enumerate(selected, start=1):
        source_path = item["source_path"]
        file_id = source_path.stem
        suffix = source_path.suffix or ".wav"
        dst_name = f"mandarin_{index:04d}_{file_id}{suffix}"
        dst_path = WAV_DIR / dst_name
        shutil.copy2(source_path, dst_path)

        saved_path = dst_path.relative_to(OUT_ROOT).as_posix()
        duration = wav_duration(dst_path)

        rows.append(
            {
                "编号": f"mandarin_{index:04d}",
                "语言": LANGUAGE,
                "文本": item["text"],
                "saved_path": saved_path,
                "时长": "" if duration is None else round(duration, 3),
                "id": file_id,
                "speaker_id": item["speaker_id"],
                "source_path": str(source_path),
                "source_filename": item["source_filename"],
                "source_transcription": item["text"],
                "sample_index": index,
            }
        )

    write_excel(OUT_ROOT / "普通话_映射表.xlsx", rows)
    write_excel(OUT_ROOT / "普通话_excel模版.xlsx", [])

    csv_path = OUT_ROOT / "metadata.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    speaker_counts = defaultdict(int)
    for row in rows:
        speaker_counts[row["speaker_id"]] += 1

    print(f"done: copied {len(rows)} files")
    print(f"speakers: {len(speaker_counts)}")
    print(f"output: {OUT_ROOT.resolve()}")
    print(f"excel: {(OUT_ROOT / '普通话_映射表.xlsx').resolve()}")
    for speaker_id, count in sorted(speaker_counts.items()):
        print(f"{speaker_id}: {count}")


if __name__ == "__main__":
    main()
