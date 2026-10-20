# -*- coding: utf-8 -*-
"""定位 FireRedASR 在 noises_en 上为何 127/255 识别为空 -- 看是模型能力、超长截断还是异常。"""
from __future__ import annotations

import csv
import statistics
from pathlib import Path

import openpyxl

META = Path(r"d:/PyCode/test/downloaded_audio/integrated_dialects/noises_en/metadata.csv")
XLSX = Path(r"d:/PyCode/test/downloaded_audio/integrated_dialects/打分结果/fire_red_ASR/noises_en_fire_red_ASR.xlsx")

rows = list(csv.DictReader(open(META, encoding="utf-8-sig")))
by_code = {r["编号"]: r for r in rows}

wb = openpyxl.load_workbook(XLSX)
ws = wb.active

empty = []
good = []
for r in range(2, ws.max_row + 1):
    code = ws.cell(row=r, column=1).value
    hyp = ws.cell(row=r, column=11).value or ""
    if not code:
        continue
    rec = by_code.get(code)
    if not rec:
        continue
    dur = float(rec["时长"])
    text = rec["文本"]
    (empty if not hyp else good).append((code, dur, text, hyp))

print(f"total rows: {ws.max_row - 1}")
print(f"empty hyp : {len(empty)}")
print(f"good hyp  : {len(good)}")
print()
print("时长统计 (秒):")
print(f"  empty avg={statistics.mean(d for _, d, _, _ in empty):.2f}  max={max(d for _, d, _, _ in empty):.2f}  min={min(d for _, d, _, _ in empty):.2f}")
print(f"  good  avg={statistics.mean(d for _, d, _, _ in good):.2f}  max={max(d for _, d, _, _ in good):.2f}  min={min(d for _, d, _, _ in good):.2f}")
print()

# 按时长分桶看比例
print("失败按时长分桶:")
buckets = [(0, 3), (3, 5), (5, 8), (8, 12), (12, 100)]
for lo, hi in buckets:
    cnt_e = sum(1 for _, d, _, _ in empty if lo <= d < hi)
    cnt_g = sum(1 for _, d, _, _ in good if lo <= d < hi)
    total = cnt_e + cnt_g
    rate = cnt_e / total * 100 if total else 0
    print(f"  {lo:>2}-{hi:<3}s  失败 {cnt_e:>3} / 总 {total:>3}   ({rate:5.1f}%)")

print()
print("失败样本前 5 条 (含时长/参考文本):")
for code, d, text, _ in empty[:5]:
    print(f"  {code}  {d:.2f}s  ref={text!r}")
