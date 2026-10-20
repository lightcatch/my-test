# -*- coding: utf-8 -*-
"""把 FireRedASR 的失败/成功状态和音频采样率交叉对比。"""
from __future__ import annotations

import contextlib
import csv
import wave
from collections import Counter
from pathlib import Path

import openpyxl

META = Path(r"d:/PyCode/test/downloaded_audio/integrated_dialects/noises_en/metadata.csv")
XLSX = Path(r"d:/PyCode/test/downloaded_audio/integrated_dialects/打分结果/fire_red_ASR/noises_en_fire_red_ASR.xlsx")

rows = {r["编号"]: r for r in csv.DictReader(open(META, encoding="utf-8-sig"))}
wb = openpyxl.load_workbook(XLSX)
ws = wb.active

by_sr = {16000: Counter(), 48000: Counter()}
for r in range(2, ws.max_row + 1):
    code = ws.cell(row=r, column=1).value
    if not code:
        continue
    hyp = ws.cell(row=r, column=11).value or ""
    rec = rows.get(code)
    if not rec:
        continue
    wav = (META.parent / rec["saved_path"]).resolve()
    with contextlib.closing(wave.open(str(wav), "rb")) as wf:
        sr = wf.getframerate()
    by_sr.setdefault(sr, Counter())[
        "empty" if not hyp else "good"
    ] += 1

print(f"{'sr':<8}{'good':>6}{'empty':>8}{'fail %':>9}")
for sr in sorted(by_sr):
    c = by_sr[sr]
    total = c["good"] + c["empty"]
    rate = c["empty"] / total * 100 if total else 0
    print(f"{sr:<8}{c['good']:>6}{c['empty']:>8}{rate:>8.1f}%")
