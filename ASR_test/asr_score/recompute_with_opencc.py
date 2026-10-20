# -*- coding: utf-8 -*-
"""把已生成的打分 Excel 中"识别出的文字"列做 OpenCC 简->繁 (s2t) 转换, 重算 CER 后回写。

主要场景: FireRedASR 在粤语场景输出简体普通话, 与繁体粤字参考差距大被 CER 严重惩罚。
做 s2t 转换后 CER 会回归合理水平。

用法:
  .venv\\Scripts\\python.exe -X utf8 ASR_test\\asr_score\\recompute_with_opencc.py ^
    "打分结果\\fire_red_ASR\\yue_fire_red_ASR.xlsx" ^
    "打分结果\\fire_red_ASR\\mixed_noise_yue_fire_red_ASR.xlsx"
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
PARENT_DIR = SCRIPT_DIR.parent  # ASR_test
if str(PARENT_DIR) not in sys.path:
    sys.path.insert(0, str(PARENT_DIR))

from asr_score.cer import compute_cer  # noqa: E402


def s2t_text(converter, text: str) -> str:
    if not text:
        return text
    return converter.convert(text)


def recompute_one(xlsx_path: Path, converter) -> tuple[float | None, float | None]:
    """覆盖原 Excel: 把识别文字 s2t 转换, 重算 S/D/I/N。返回 (原 avg_cer, 新 avg_cer)。"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # 读表头, 定位列
    header = {c.value: c.column for c in ws[1]}
    needed = ["文本", "识别出的文字", "替换错误数S", "删除错误数D",
              "插入错误数I", "参考总字数N"]
    for k in needed:
        if k not in header:
            raise RuntimeError(f"{xlsx_path.name} 表头缺列: {k}")

    col_ref = header["文本"]
    col_hyp = header["识别出的文字"]
    col_s = header["替换错误数S"]
    col_d = header["删除错误数D"]
    col_i = header["插入错误数I"]
    col_n = header["参考总字数N"]

    old_cers, new_cers = [], []
    for r in range(2, ws.max_row + 1):
        ref = ws.cell(row=r, column=col_ref).value
        hyp = ws.cell(row=r, column=col_hyp).value
        if not isinstance(ref, str) or not ref.strip():
            continue
        old_n = ws.cell(row=r, column=col_n).value or 0
        old_s = ws.cell(row=r, column=col_s).value or 0
        old_d = ws.cell(row=r, column=col_d).value or 0
        old_i = ws.cell(row=r, column=col_i).value or 0
        if isinstance(old_n, (int, float)) and old_n > 0:
            old_cers.append((float(old_s) + float(old_d) + float(old_i)) / float(old_n) * 100)

        if not isinstance(hyp, str):
            hyp = "" if hyp is None else str(hyp)

        new_hyp = s2t_text(converter, hyp)
        if new_hyp != hyp:
            ws.cell(row=r, column=col_hyp, value=new_hyp)

        c = compute_cer(ref, new_hyp)
        ws.cell(row=r, column=col_s, value=c.s)
        ws.cell(row=r, column=col_d, value=c.d)
        ws.cell(row=r, column=col_i, value=c.i)
        ws.cell(row=r, column=col_n, value=c.n)
        if c.n > 0:
            new_cers.append(c.cer)

    wb.save(xlsx_path)

    old_avg = round(sum(old_cers) / len(old_cers), 2) if old_cers else None
    new_avg = round(sum(new_cers) / len(new_cers), 2) if new_cers else None
    return old_avg, new_avg


def subset_from_path(xlsx_path: Path) -> str:
    """文件名 xxx_<DisplayName>.xlsx -> 子集名 xxx"""
    name = xlsx_path.stem
    parent_dir = xlsx_path.parent.name  # 例如 fire_red_ASR
    suffix = "_" + parent_dir
    if name.endswith(suffix):
        return name[: -len(suffix)]
    return name


def update_summary(xlsx_path: Path, new_avg_cer: float | None) -> None:
    """同步更新该 ASR 目录下 _summary.csv 中匹配 name 的行的 avg_cer。"""
    if new_avg_cer is None:
        return
    summary = xlsx_path.parent / "_summary.csv"
    if not summary.exists():
        print(f"  [skip summary] {summary} 不存在")
        return
    target_name = subset_from_path(xlsx_path)
    rows: list[dict] = []
    fields: list[str] = []
    with open(summary, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fields = list(reader.fieldnames or [])
        for r in reader:
            if r.get("name") == target_name:
                r["avg_cer"] = str(new_avg_cer)
            rows.append(r)
    with open(summary, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    targets = [Path(a) for a in sys.argv[1:]]
    for p in targets:
        if not p.exists():
            print(f"  ERROR: 文件不存在: {p}")
            return 1

    try:
        from opencc import OpenCC  # type: ignore
    except ImportError as e:
        print(f"OpenCC 没装: {e}")
        print("请: uv pip install --python <venv-python> opencc-python-reimplemented")
        return 1

    converter = OpenCC("s2t")  # 简体 -> 繁体

    print(f"OpenCC s2t 转换 + 重算 CER，共 {len(targets)} 个文件\n")
    for p in targets:
        print(f">>> {p}")
        old, new = recompute_one(p, converter)
        update_summary(p, new)
        msg = f"  原 avg_cer={old}%  ->  新 avg_cer={new}%"
        if old is not None and new is not None:
            delta = round(new - old, 2)
            msg += f"  (delta {delta:+.2f}%)"
        print(msg)
    print("\n完成。可再执行 build_overall_summary.py 刷新总表。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
