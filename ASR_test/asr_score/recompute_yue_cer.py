# -*- coding: utf-8 -*-
"""一次性工具: 把粤语子集的参考文本与识别文本简繁归一(繁->简)后重算 CER。

背景: 粤语子集(yue / mixed_noise_yue)参考文本是繁体, FireRedASR2 识别输出是简体,
内容其实识别对了, 但 CER 因简繁字形不一致虚高(~34%)。本脚本用 opencc t2s 把两边
都转简体后重算, 看真实识别能力。

- 不修改主流程 cer.py, 复用 asr_score.cer.compute_cer。
- 默认只打印对比表; 加 --write 时生成副本 <name>_..._简繁归一.xlsx, 原文件不动。
- 用项目根 .venv 跑(有 opencc + openpyxl):
    .venv/Scripts/python.exe -X utf8 ASR_test/asr_score/recompute_yue_cer.py [--write]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

# 让脚本能 import asr_score.cer(ASR_test 加入 path)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from asr_score.cer import compute_cer  # noqa: E402

RESULT_DIR = (
    Path(__file__).resolve().parents[2]
    / "downloaded_audio" / "integrated_dialects" / "打分结果" / "FireRed2_ONNX_INT8"
)
SUBSETS = ["yue", "mixed_noise_yue"]
SUFFIX = "FireRed2_ONNX_INT8"

# 列索引(1-based), 与 打分模版 一致
COL_REF = 3    # 文本(参考)
COL_HYP = 11   # 识别出的文字
COL_S, COL_D, COL_I, COL_N = 12, 13, 14, 15


def main() -> int:
    parser = argparse.ArgumentParser(description="粤语子集简繁归一后重算 CER。")
    parser.add_argument("--write", action="store_true",
                        help="生成简繁归一后的 Excel 副本(原文件不动)。")
    args = parser.parse_args()

    import opencc  # type: ignore
    from openpyxl import load_workbook  # type: ignore

    t2s = opencc.OpenCC("t2s")

    print(f"{'子集':<20}{'条数':>6}{'原始行均':>10}{'归一行均':>10}"
          f"{'原始micro':>11}{'归一micro':>11}")
    print("-" * 68)

    for name in SUBSETS:
        xlsx = RESULT_DIR / f"{name}_{SUFFIX}.xlsx"
        if not xlsx.is_file():
            print(f"{name}: 找不到 {xlsx}")
            continue
        wb = load_workbook(xlsx)
        ws = wb["metadata"]

        rows = 0
        old_cer_sum = new_cer_sum = 0.0
        old_e = old_n = new_e = new_n = 0

        for r in range(2, ws.max_row + 1):
            ref = ws.cell(r, COL_REF).value
            hyp = ws.cell(r, COL_HYP).value
            if not ref:
                continue
            ref = str(ref)
            hyp = str(hyp) if hyp is not None else ""

            r_old = compute_cer(ref, hyp)
            ref_s, hyp_s = t2s.convert(ref), t2s.convert(hyp)
            r_new = compute_cer(ref_s, hyp_s)

            rows += 1
            if r_old.n > 0:
                old_cer_sum += r_old.cer
                new_cer_sum += r_new.cer
            old_e += r_old.total_errors
            old_n += r_old.n
            new_e += r_new.total_errors
            new_n += r_new.n

            if args.write:
                ws.cell(r, COL_S).value = r_new.s
                ws.cell(r, COL_D).value = r_new.d
                ws.cell(r, COL_I).value = r_new.i
                ws.cell(r, COL_N).value = r_new.n
                # 第16列 CER 是 Excel 公式, 会自动按 S/D/I/N 重算, 不用改

        old_avg = old_cer_sum / rows if rows else 0.0
        new_avg = new_cer_sum / rows if rows else 0.0
        old_micro = old_e / old_n * 100 if old_n else 0.0
        new_micro = new_e / new_n * 100 if new_n else 0.0
        print(f"{name:<20}{rows:>6}{old_avg:>9.2f}%{new_avg:>9.2f}%"
              f"{old_micro:>10.2f}%{new_micro:>10.2f}%")

        if args.write:
            out = RESULT_DIR / f"{name}_{SUFFIX}_简繁归一.xlsx"
            wb.save(out)
            print(f"    已写副本 -> {out}")

    print("\n说明: 行均=每行 CER 的算术平均(与 _summary 的 avg_cer 口径一致); "
          "micro=总错误/总N。简繁归一即参考与识别都用 opencc t2s 转简体后再算。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
