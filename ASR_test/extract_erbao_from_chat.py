from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_ROOT = Path(
    "D:/\u4e0b\u8f7d/\u98de\u4e66\u4e0b\u8f7d/"
    "YB01202606110200166D649549_2026-06-15_chat_history"
)
CHAT_HISTORY = SOURCE_ROOT / "chat_history.xlsx"
SOURCE_AUDIO_DIR = SOURCE_ROOT / "audio"
OUTPUT_ROOT = Path("downloaded_audio") / "integrated_dialects" / "erbao"
OUTPUT_WAV_DIR = OUTPUT_ROOT / "wav"
CORRECT_TEXTS_FILE = Path("ASR_test") / "erbao_correct_texts.txt"

START_AUDIO_ID = "2066356405645963266"
START_TEXT = "\u5988\u670b\u6709\u9001\u8bde\u6211"
TAKE_USER_MESSAGES = 20

COLS = {
    "record_id": "\u8bb0\u5f55ID",
    "role": "\u89d2\u8272",
    "content": "\u5185\u5bb9",
    "audio_id": "\u97f3\u9891ID",
    "duration": "\u97f3\u9891\u65f6\u957f(\u79d2)",
    "record_time": "\u8bb0\u5f55\u65f6\u95f4",
}

OUTPUT_COLUMNS = [
    "\u7f16\u53f7",
    "\u8bed\u8a00",
    "\u6587\u672c",
    "saved_path",
    "\u6587\u4ef6\u540d",
    "\u65f6\u957f",
    "id",
    "speaker_id",
    "sample_index",
    "dialect",
]


def find_start_index(df: pd.DataFrame) -> int:
    id_mask = df.apply(
        lambda col: col.astype(str).str.contains(START_AUDIO_ID, na=False)
    ).any(axis=1)
    text_mask = df[COLS["content"]].astype(str).str.contains(START_TEXT, na=False)
    matches = df.index[id_mask | text_mask].tolist()
    if not matches:
        raise RuntimeError(f"Start row not found: {START_TEXT}/{START_AUDIO_ID}")
    return int(matches[0])


def find_audio_file(audio_id: str) -> Path:
    preferred = SOURCE_AUDIO_DIR / f"{audio_id}_2.wav"
    if preferred.exists():
        return preferred

    matches = sorted(SOURCE_AUDIO_DIR.glob(f"{audio_id}*.wav"))
    if matches:
        return matches[0]

    raise FileNotFoundError(f"Audio file not found for audio_id={audio_id}")


def load_correct_texts() -> list[str]:
    if not CORRECT_TEXTS_FILE.exists():
        raise FileNotFoundError(
            f"Correct text file is missing: {CORRECT_TEXTS_FILE}\n"
            f"Put the {TAKE_USER_MESSAGES} correct texts in it, one per line, "
            "in the same order as the screenshot."
        )

    texts = [
        line.strip()
        for line in CORRECT_TEXTS_FILE.read_text(encoding="utf-8-sig").splitlines()
        if line.strip()
    ]
    if len(texts) != TAKE_USER_MESSAGES:
        raise RuntimeError(
            f"Expected {TAKE_USER_MESSAGES} correct texts in {CORRECT_TEXTS_FILE}, "
            f"got {len(texts)}"
        )
    return texts


def write_mapping_xlsx(rows: list[dict[str, str]]) -> Path:
    xlsx_path = OUTPUT_ROOT / "\u4e8c\u5b9d_\u6620\u5c04\u8868.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "\u4e8c\u5b9d"
    ws.append(OUTPUT_COLUMNS)

    for row in rows:
        ws.append([row[column] for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [14, 10, 42, 36, 36, 10, 24, 14, 14, 14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return xlsx_path


def write_metadata_csv(rows: list[dict[str, str]]) -> Path:
    csv_path = OUTPUT_ROOT / "metadata.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row[column] for column in OUTPUT_COLUMNS})
    return csv_path


def main() -> None:
    if not CHAT_HISTORY.exists():
        raise FileNotFoundError(CHAT_HISTORY)
    if not SOURCE_AUDIO_DIR.exists():
        raise FileNotFoundError(SOURCE_AUDIO_DIR)

    OUTPUT_WAV_DIR.mkdir(parents=True, exist_ok=True)
    correct_texts = load_correct_texts()

    df = pd.read_excel(CHAT_HISTORY, sheet_name=0, dtype=str)
    start_index = find_start_index(df)

    user_rows = (
        df.iloc[start_index:]
        .loc[lambda x: x[COLS["role"]].eq("USER")]
        .head(TAKE_USER_MESSAGES)
        .copy()
    )
    if len(user_rows) != TAKE_USER_MESSAGES:
        raise RuntimeError(
            f"Only found {len(user_rows)} USER rows after index {start_index}; "
            f"expected {TAKE_USER_MESSAGES}"
        )

    rows: list[dict[str, str]] = []
    for sample_index, (source_index, source_row) in enumerate(user_rows.iterrows(), 1):
        audio_id = str(source_row[COLS["audio_id"]]).strip()
        audio_path = find_audio_file(audio_id)
        sample_id = f"erbao_{sample_index:04d}"
        filename = f"{sample_id}_{audio_id}.wav"
        target_path = OUTPUT_WAV_DIR / filename
        shutil.copy2(audio_path, target_path)

        rows.append(
            {
                "\u7f16\u53f7": sample_id,
                "\u8bed\u8a00": "\u4e8c\u5b9d",
                "\u6587\u672c": correct_texts[sample_index - 1],
                "saved_path": f"wav/{filename}",
                "\u6587\u4ef6\u540d": filename,
                "\u65f6\u957f": str(source_row[COLS["duration"]]),
                "id": audio_id,
                "speaker_id": "erbao",
                "sample_index": str(sample_index),
                "dialect": "erbao",
            }
        )

    mapping_path = write_mapping_xlsx(rows)
    metadata_path = write_metadata_csv(rows)

    print(f"start_index={start_index} excel_row={start_index + 2}")
    print(f"copied_wav={len(rows)}")
    print(f"mapping={mapping_path}")
    print(f"metadata={metadata_path}")


if __name__ == "__main__":
    main()
