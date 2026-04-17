import argparse
import base64
import json
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook


APP_ID = "6314392664"
ACCESS_TOKEN = "REDACTED_CREDENTIAL"
RESOURCE_ID = "volc.bigasr.auc_turbo"
MODEL_NAME = "bigmodel"
RECOGNIZE_URL = "https://openspeech.bytedance.com/api/v3/auc/bigmodel/recognize/flash"

SCRIPT_DIR = Path(__file__).resolve().parent
ASR_DIR = SCRIPT_DIR.parents[1]
DEFAULT_EXCEL = ASR_DIR / "chat_history.xlsx"
DEFAULT_AUDIO_DIR = ASR_DIR / "audio"
DEFAULT_RESULTS_DIR = ASR_DIR / "results"
DEFAULT_ASR_COLUMN = "AUC Flash\u8bc6\u522b\u6587\u5b57"
TEXT_KEYS = {"text", "transcript", "sentence", "utterance", "result_text", "final_text"}


class AucFlashError(RuntimeError):
    pass


def file_to_base64(file_path: str | Path) -> str:
    with open(file_path, "rb") as file:
        return base64.b64encode(file.read()).decode("utf-8")


def recognize_task(
    file_url: str | None = None,
    file_path: str | Path | None = None,
    *,
    appid: str = APP_ID,
    token: str = ACCESS_TOKEN,
    timeout: int = 90,
    verbose: bool = True,
) -> requests.Response:
    headers = {
        "X-Api-App-Key": appid,
        "X-Api-Access-Key": token,
        "X-Api-Resource-Id": RESOURCE_ID,
        "X-Api-Request-Id": str(uuid.uuid4()),
        "X-Api-Sequence": "-1",
    }

    audio_data: dict[str, str] | None = None
    if file_url:
        audio_data = {"url": file_url}
    elif file_path:
        audio_data = {"data": file_to_base64(file_path)}

    if not audio_data:
        raise ValueError("Either file_url or file_path is required.")

    request = {
        "user": {
            "uid": appid,
        },
        "audio": audio_data,
        "request": {
            "model_name": MODEL_NAME,
            "enable_itn": True,
            "enable_punc": True,
            "enable_ddc": True,
        },
    }

    response = requests.post(RECOGNIZE_URL, json=request, headers=headers, timeout=timeout)
    if verbose:
        print_response_summary(response)
    return response


def print_response_summary(response: requests.Response) -> None:
    status_code = response.headers.get("X-Api-Status-Code", "")
    message = response.headers.get("X-Api-Message", "")
    logid = response.headers.get("X-Tt-Logid", "")
    print(f"X-Api-Status-Code: {status_code}")
    print(f"X-Api-Message: {message}")
    print(f"X-Tt-Logid: {logid}")
    try:
        print(f"Response JSON: {json.dumps(response.json(), ensure_ascii=False)}")
    except ValueError:
        print(f"Response Text: {response.text}")


def response_json(response: requests.Response) -> Any:
    try:
        return response.json()
    except ValueError as exc:
        raise AucFlashError(f"Response is not valid JSON: {response.text[:300]}") from exc


def collect_texts(payload: Any) -> list[str]:
    texts: list[str] = []
    if isinstance(payload, dict):
        for key, value in payload.items():
            if key in TEXT_KEYS and isinstance(value, str):
                text = value.strip()
                if text:
                    texts.append(text)
            else:
                texts.extend(collect_texts(value))
        return texts
    if isinstance(payload, list):
        for item in payload:
            texts.extend(collect_texts(item))
    return texts


def extract_transcript(payload: Any) -> str:
    if isinstance(payload, dict):
        result = payload.get("result")
        if isinstance(result, dict):
            text = result.get("text")
            if isinstance(text, str) and text.strip():
                return text.strip()
            utterances = result.get("utterances")
            if isinstance(utterances, list):
                texts = [
                    item.get("text", "").strip()
                    for item in utterances
                    if isinstance(item, dict) and isinstance(item.get("text"), str) and item.get("text", "").strip()
                ]
                if texts:
                    return "".join(texts).strip()
        if isinstance(result, list):
            texts = [
                item.get("text", "").strip()
                for item in result
                if isinstance(item, dict) and isinstance(item.get("text"), str) and item.get("text", "").strip()
            ]
            if texts:
                return "".join(texts).strip()

    texts = collect_texts(payload)
    return "".join(texts).strip()


def recognize_text(file_path: Path, *, timeout: int, verbose: bool = False) -> str:
    response = recognize_task(file_path=file_path, timeout=timeout, verbose=verbose)
    status_code = response.headers.get("X-Api-Status-Code", "")
    message = response.headers.get("X-Api-Message", "")
    if status_code != "20000000":
        raise AucFlashError(f"AUC Flash failed: code={status_code}, message={message}")

    transcript = extract_transcript(response_json(response))
    if not transcript:
        raise AucFlashError("AUC Flash returned an empty transcript.")
    return transcript


def get_default_output_path() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return DEFAULT_RESULTS_DIR / f"chat_history_with_auc_flash_{timestamp}.xlsx"


def get_header_map(worksheet: Any) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for col_idx in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=col_idx).value
        if value is not None and str(value).strip():
            header_map[str(value).strip()] = col_idx
    return header_map


def get_or_create_column(worksheet: Any, column_name: str) -> tuple[int, bool]:
    header_map = get_header_map(worksheet)
    if column_name in header_map:
        return header_map[column_name], False
    col_idx = worksheet.max_column + 1
    worksheet.cell(row=1, column=col_idx, value=column_name)
    return col_idx, True


def has_text(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def audio_candidates(raw_target: str, excel_dir: Path, audio_dir: Path) -> list[Path]:
    target = Path(raw_target.replace("\\", "/"))
    if target.is_absolute():
        return [target]
    return [
        (excel_dir / target).resolve(),
        (audio_dir / target.name).resolve(),
        (audio_dir / raw_target).resolve(),
    ]


def resolve_audio_path_for_row(
    worksheet: Any,
    row_idx: int,
    header_map: dict[str, int],
    excel_dir: Path,
    audio_dir: Path,
) -> tuple[Path | None, str | None]:
    audio_file_col = header_map.get("audio_file")
    audio_link_col = header_map.get("audio_link")
    audio_file = ""

    if audio_file_col is not None:
        value = worksheet.cell(row=row_idx, column=audio_file_col).value
        if has_text(value):
            audio_file = str(value).strip()

    if audio_link_col is not None:
        link_cell = worksheet.cell(row=row_idx, column=audio_link_col)
        hyperlink = link_cell.hyperlink.target if link_cell.hyperlink else None
        if has_text(hyperlink):
            for candidate in audio_candidates(str(hyperlink).strip(), excel_dir, audio_dir):
                if candidate.exists():
                    return candidate, None
            if audio_file:
                fallback = (audio_dir / audio_file).resolve()
                if fallback.exists():
                    return fallback, None
            return None, f"Audio file not found for hyperlink target: {hyperlink}"

    if audio_file:
        candidate = (audio_dir / audio_file).resolve()
        if candidate.exists():
            return candidate, None
        return None, f"Audio file not found: {candidate}"

    return None, None


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="AUC Flash ASR demo and Excel fill tool.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    transcribe = subparsers.add_parser("transcribe", help="Transcribe one local audio file.")
    transcribe.add_argument("--file", type=Path, required=True, help="Local audio file path.")
    transcribe.add_argument("--timeout", type=int, default=90, help="HTTP request timeout in seconds.")

    fill_excel = subparsers.add_parser("fill-excel", help="Fill a new Excel file with AUC Flash output.")
    fill_excel.add_argument("--excel", type=Path, default=DEFAULT_EXCEL, help="Source Excel path.")
    fill_excel.add_argument("--audio-dir", type=Path, default=DEFAULT_AUDIO_DIR, help="Directory containing audio files.")
    fill_excel.add_argument("--output", type=Path, default=None, help="Output Excel path.")
    fill_excel.add_argument("--sheet-index", type=int, default=0, help="Zero-based worksheet index.")
    fill_excel.add_argument("--asr-column", default=DEFAULT_ASR_COLUMN, help="Output ASR column name.")
    fill_excel.add_argument("--limit", type=int, default=None, help="Only process the first N audio rows.")
    fill_excel.add_argument("--skip-existing", action="store_true", help="Skip rows with existing output values.")
    fill_excel.add_argument("--timeout", type=int, default=90, help="HTTP request timeout in seconds.")
    return parser


def validate_transcribe_args(args: argparse.Namespace) -> None:
    if not args.file.exists():
        raise AucFlashError(f"Audio file does not exist: {args.file}")
    if args.timeout <= 0:
        raise AucFlashError("--timeout must be positive.")


def run_transcribe(args: argparse.Namespace) -> int:
    validate_transcribe_args(args)
    text = recognize_text(args.file.resolve(), timeout=args.timeout, verbose=True)
    print(text)
    return 0


def validate_fill_excel_args(args: argparse.Namespace) -> None:
    if not args.excel.exists():
        raise AucFlashError(f"Excel file does not exist: {args.excel}")
    if not args.audio_dir.exists() or not args.audio_dir.is_dir():
        raise AucFlashError(f"Audio directory does not exist: {args.audio_dir}")
    if args.sheet_index < 0:
        raise AucFlashError("--sheet-index must be >= 0.")
    if args.limit is not None and args.limit <= 0:
        raise AucFlashError("--limit must be positive when provided.")
    if not args.asr_column.strip():
        raise AucFlashError("--asr-column must not be empty.")
    if args.timeout <= 0:
        raise AucFlashError("--timeout must be positive.")


def run_fill_excel(args: argparse.Namespace) -> int:
    validate_fill_excel_args(args)
    output_path = args.output or get_default_output_path()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(args.excel)
    if args.sheet_index >= len(workbook.worksheets):
        raise AucFlashError(f"--sheet-index {args.sheet_index} is out of range.")

    worksheet = workbook.worksheets[args.sheet_index]
    header_map = get_header_map(worksheet)
    output_col, created_output_col = get_or_create_column(worksheet, args.asr_column)
    excel_dir = args.excel.resolve().parent
    audio_dir = args.audio_dir.resolve()

    attempted = 0
    succeeded = 0
    failed = 0
    skipped_existing = 0
    failed_rows: list[tuple[int, str, str]] = []

    for row_idx in range(2, worksheet.max_row + 1):
        if args.limit is not None and attempted >= args.limit:
            break

        output_cell = worksheet.cell(row=row_idx, column=output_col)
        if args.skip_existing and has_text(output_cell.value):
            skipped_existing += 1
            continue

        audio_path, resolution_error = resolve_audio_path_for_row(
            worksheet=worksheet,
            row_idx=row_idx,
            header_map=header_map,
            excel_dir=excel_dir,
            audio_dir=audio_dir,
        )
        if audio_path is None:
            if resolution_error:
                failed += 1
                failed_rows.append((row_idx, "", resolution_error))
                print(f"Row {row_idx}: {resolution_error}", file=sys.stderr)
            continue

        attempted += 1
        try:
            print(f"[{attempted}] Row {row_idx}: {audio_path}")
            text = recognize_text(audio_path, timeout=args.timeout)
            output_cell.value = text
            succeeded += 1
            print(f"Row {row_idx}: {text}")
        except Exception as exc:  # noqa: BLE001 - keep filling remaining rows.
            failed += 1
            failed_rows.append((row_idx, str(audio_path), str(exc)))
            print(f"Row {row_idx}: {audio_path} -> {exc}", file=sys.stderr)

    workbook.save(output_path)

    print(f"Sheet: {worksheet.title}")
    print(f"ASR column: {args.asr_column}")
    print(f"Created new column: {created_output_col}")
    print(f"Attempted audio rows: {attempted}")
    print(f"Succeeded: {succeeded}")
    print(f"Failed: {failed}")
    print(f"Skipped existing: {skipped_existing}")
    print(f"Result file: {output_path}")
    if failed_rows:
        print("Failed rows:")
        for row_idx, audio, reason in failed_rows:
            print(f"  row={row_idx} audio={audio} reason={reason}")

    return 0 if failed == 0 else 1


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        if args.command == "transcribe":
            return run_transcribe(args)
        if args.command == "fill-excel":
            return run_fill_excel(args)
        raise AucFlashError(f"Unsupported command: {args.command}")
    except AucFlashError as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
