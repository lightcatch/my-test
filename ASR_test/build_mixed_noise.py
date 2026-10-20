from collections import Counter, defaultdict
from math import gcd
from pathlib import Path
import csv
import random
import shutil
import wave

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.io import wavfile
from scipy.signal import resample_poly


RANDOM_SEED = 20260625
SAMPLE_SIZE_PER_REGION = 100
SNR_DB = 10.0
TARGET_SR = 16000

INTEGRATED_ROOT = Path("downloaded_audio") / "integrated_dialects"
OUT_ROOT = Path("downloaded_audio") / "mixed_noise"
KESPEECH_ROOT = Path("downloaded_audio") / "kespeech"
SICHUAN_ROOT = Path("downloaded_audio") / "sichuan"
COMMON_VOICE_ROOT = Path(r"\\Desktop-p5q0of7\公开数据集\cv-corpus-24.0-2025-12-05")
ESC50_ROOT = Path(r"\\Desktop-p5q0of7\公开数据集\ESC-50_audio")

REGION_LANGUAGE = {
    "sichuan": "\u56db\u5ddd\u8bdd",
    "yue": "\u7ca4\u8bed",
    "henan": "\u6cb3\u5357\u8bdd",
    "dongbei": "\u4e1c\u5317\u8bdd",
    "shandong": "\u5c71\u4e1c\u8bdd",
}

NOISE_TYPES = [
    "\u96e8\u58f0",
    "\u98ce\u58f0",
    "\u811a\u6b65\u58f0",
    "\u952e\u76d8\u6253\u5b57\u58f0",
    "\u9f20\u6807\u70b9\u51fb\u58f0",
    "\u6572\u95e8\u58f0",
    "\u95e8\u5431\u5440\u58f0",
    "\u5438\u5c18\u5668\u58f0",
    "\u6d17\u8863\u673a\u58f0",
    "\u6c7d\u8f66\u5587\u53ed\u58f0",
]

COL_ID = "\u7f16\u53f7"
COL_LANG = "\u8bed\u8a00"
COL_TEXT = "\u6587\u672c"
COL_FILE = "\u6587\u4ef6\u540d"
COL_DURATION = "\u65f6\u957f"

FIELDNAMES = [
    COL_ID,
    COL_LANG,
    COL_TEXT,
    "saved_path",
    COL_FILE,
    COL_DURATION,
    "id",
    "speaker_id",
    "sample_index",
    "dialect",
    "noise_type",
    "noise_file",
    "snr_db",
]


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def existing_ids(region):
    path = INTEGRATED_ROOT / region / "metadata.csv"
    return {row["id"] for row in read_csv(path) if row.get("id")}


def wav_duration(path):
    with wave.open(str(path), "rb") as wav_file:
        return round(wav_file.getnframes() / wav_file.getframerate(), 3)


def load_sichuan_candidates():
    seen_ids = existing_ids("sichuan")
    candidates = []
    with (SICHUAN_ROOT / "list.txt").open("r", encoding="utf-8-sig", newline="") as f:
        for line in f:
            parts = line.rstrip("\r\n").split("\t")
            if len(parts) < 5:
                continue
            wav_name = parts[1]
            speaker_id = parts[2]
            file_id = Path(wav_name).stem
            if file_id in seen_ids:
                continue
            wav_path = SICHUAN_ROOT / "wav" / speaker_id / wav_name
            if not wav_path.exists():
                continue
            candidates.append(
                {
                    "source_path": wav_path,
                    "id": file_id,
                    "speaker_id": speaker_id,
                    "text": parts[-1],
                    "dialect": "Sichuan",
                }
            )
    return candidates


def load_common_voice_candidates(locale):
    seen_ids = existing_ids(locale)
    locale_root = COMMON_VOICE_ROOT / locale
    wav_dir = locale_root / "clips_wav_16k"
    wav_by_name = {path.name: path for path in wav_dir.glob("*.wav")}
    candidates = []
    with (locale_root / "validated.tsv").open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            wav_name = Path(row["path"]).with_suffix(".wav").name
            source_path = wav_by_name.get(wav_name)
            if source_path is None:
                continue
            file_id = source_path.stem
            if file_id in seen_ids:
                continue
            candidates.append(
                {
                    "source_path": source_path,
                    "id": file_id,
                    "speaker_id": row.get("client_id", ""),
                    "text": row.get("sentence", ""),
                    "dialect": locale,
                }
            )
    return candidates


def load_kespeech_candidates(region):
    seen_ids = existing_ids(region)
    candidates = []
    for row in read_csv(KESPEECH_ROOT / region / "metadata.csv"):
        file_id = row.get("id", "")
        if not file_id or file_id in seen_ids:
            continue
        source_path = KESPEECH_ROOT / row["saved_path"]
        if not source_path.exists():
            continue
        candidates.append(
            {
                "source_path": source_path,
                "id": file_id,
                "speaker_id": file_id.split("_", 1)[0],
                "text": row.get("text", ""),
                "dialect": row.get("dialect", region),
            }
        )
    return candidates


def balanced_sample(candidates, region, count):
    rng = random.Random(f"{RANDOM_SEED}:{region}:speech")
    pools = defaultdict(list)
    for item in candidates:
        pools[item["speaker_id"]].append(item)
    for pool in pools.values():
        rng.shuffle(pool)

    speaker_ids = list(pools)
    rng.shuffle(speaker_ids)
    selected = []
    while len(selected) < count:
        progressed = False
        for speaker_id in speaker_ids:
            if not pools[speaker_id]:
                continue
            selected.append(pools[speaker_id].pop())
            progressed = True
            if len(selected) >= count:
                break
        if not progressed:
            break
    if len(selected) < count:
        raise RuntimeError(f"{region} only has {len(selected)} unselected usable samples")
    return selected


def esc50_label(path):
    stem = path.stem
    return stem.split("_")[-1] if "_" in stem else stem


def load_noise_files():
    files_by_type = defaultdict(list)
    for path in ESC50_ROOT.glob("*.wav"):
        label = esc50_label(path)
        if label in NOISE_TYPES:
            files_by_type[label].append(path)
    for label in NOISE_TYPES:
        if len(files_by_type[label]) < 10:
            raise RuntimeError(f"not enough ESC-50 noise files for {label}: {len(files_by_type[label])}")
        files_by_type[label].sort()
    return files_by_type


def noise_assignments(region, noise_files_by_type):
    rng = random.Random(f"{RANDOM_SEED}:{region}:noise")
    assignments = []
    for noise_type in NOISE_TYPES:
        files = list(noise_files_by_type[noise_type])
        rng.shuffle(files)
        for path in files[:10]:
            assignments.append((noise_type, path))
    rng.shuffle(assignments)
    return assignments


def to_float_mono(data):
    if data.ndim == 2:
        data = data.mean(axis=1)
    if np.issubdtype(data.dtype, np.integer):
        max_abs = float(np.iinfo(data.dtype).max)
        return data.astype(np.float32) / max_abs
    return data.astype(np.float32)


def resample_if_needed(samples, source_sr, target_sr):
    if source_sr == target_sr:
        return samples
    factor = gcd(source_sr, target_sr)
    return resample_poly(samples, target_sr // factor, source_sr // factor).astype(np.float32)


def load_audio(path):
    sr, data = wavfile.read(path)
    samples = to_float_mono(data)
    samples = resample_if_needed(samples, sr, TARGET_SR)
    return samples


def fit_noise_length(noise, target_length, rng):
    if len(noise) == target_length:
        return noise
    if len(noise) < target_length:
        repeats = int(np.ceil(target_length / len(noise)))
        return np.tile(noise, repeats)[:target_length]
    start = rng.randint(0, len(noise) - target_length)
    return noise[start : start + target_length]


def mix_with_snr(speech, noise, snr_db):
    speech_rms = float(np.sqrt(np.mean(np.square(speech))) + 1e-12)
    noise_rms = float(np.sqrt(np.mean(np.square(noise))) + 1e-12)
    target_noise_rms = speech_rms / (10 ** (snr_db / 20))
    mixed = speech + noise * (target_noise_rms / noise_rms)
    peak = float(np.max(np.abs(mixed)) + 1e-12)
    if peak > 0.99:
        mixed = mixed * (0.99 / peak)
    return mixed


def write_wav(path, samples):
    pcm = np.clip(samples, -1.0, 1.0)
    wavfile.write(path, TARGET_SR, (pcm * 32767.0).astype(np.int16))


def write_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(FIELDNAMES)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [14, 12, 42, 40, 38, 10, 30, 24, 14, 18, 18, 42, 10]
    for column_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=column_index).column_letter].width = width
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append([row.get(field, "") for field in FIELDNAMES])
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_metadata(region_dir, rows, language):
    with (region_dir / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)
    write_excel(region_dir / f"{language}_\u6df7\u5408\u566a\u97f3\u6620\u5c04\u8868.xlsx", rows)


def build_region(region, candidates, noise_files_by_type):
    language = REGION_LANGUAGE[region]
    selected = balanced_sample(candidates, region, SAMPLE_SIZE_PER_REGION)
    assignments = noise_assignments(region, noise_files_by_type)
    region_dir = OUT_ROOT / region
    wav_dir = region_dir / "wav"
    wav_dir.mkdir(parents=True, exist_ok=True)

    rows = []
    for index, (speech_item, (noise_type, noise_path)) in enumerate(zip(selected, assignments), start=1):
        rng = random.Random(f"{RANDOM_SEED}:{region}:mix:{index}")
        speech = load_audio(speech_item["source_path"])
        noise = fit_noise_length(load_audio(noise_path), len(speech), rng)
        mixed = mix_with_snr(speech, noise, SNR_DB)

        file_id = speech_item["id"]
        dst_name = f"{region}_mixed_{index:04d}_{file_id}_{noise_type}.wav"
        dst_path = wav_dir / dst_name
        write_wav(dst_path, mixed)

        saved_path = dst_path.relative_to(region_dir).as_posix()
        rows.append(
            {
                COL_ID: f"{region}_mixed_{index:04d}",
                COL_LANG: language,
                COL_TEXT: speech_item["text"],
                "saved_path": saved_path,
                COL_FILE: dst_path.name,
                COL_DURATION: wav_duration(dst_path),
                "id": file_id,
                "speaker_id": speech_item["speaker_id"],
                "sample_index": index,
                "dialect": speech_item["dialect"],
                "noise_type": noise_type,
                "noise_file": noise_path.name,
                "snr_db": int(SNR_DB) if SNR_DB.is_integer() else SNR_DB,
            }
        )

    write_metadata(region_dir, rows, language)
    return rows


def load_candidates():
    return {
        "sichuan": load_sichuan_candidates(),
        "yue": load_common_voice_candidates("yue"),
        "henan": load_kespeech_candidates("henan"),
        "dongbei": load_kespeech_candidates("dongbei"),
        "shandong": load_kespeech_candidates("shandong"),
    }


def validate_outputs(all_rows):
    for region, rows in all_rows.items():
        if len(rows) != SAMPLE_SIZE_PER_REGION:
            raise AssertionError(f"{region} rows != {SAMPLE_SIZE_PER_REGION}")
        region_dir = OUT_ROOT / region
        wavs = list((region_dir / "wav").glob("*.wav"))
        if len(wavs) != SAMPLE_SIZE_PER_REGION:
            raise AssertionError(f"{region} wav count != {SAMPLE_SIZE_PER_REGION}")
        csv_rows = read_csv(region_dir / "metadata.csv")
        if len(csv_rows) != SAMPLE_SIZE_PER_REGION:
            raise AssertionError(f"{region} csv rows != {SAMPLE_SIZE_PER_REGION}")
        wb = load_workbook(next(region_dir.glob("*_*.xlsx")), read_only=True)
        ws = wb.active
        if ws.max_row != SAMPLE_SIZE_PER_REGION + 1:
            raise AssertionError(f"{region} excel rows != {SAMPLE_SIZE_PER_REGION + 1}")
        noise_counts = Counter(row["noise_type"] for row in rows)
        for noise_type in NOISE_TYPES:
            if noise_counts[noise_type] != 10:
                raise AssertionError(f"{region} noise count mismatch: {noise_type}={noise_counts[noise_type]}")
        used_ids = existing_ids(region)
        overlap = used_ids.intersection(row["id"] for row in rows)
        if overlap:
            raise AssertionError(f"{region} contains previously selected ids: {sorted(overlap)[:5]}")
        for wav_path in wavs[:5]:
            with wave.open(str(wav_path), "rb") as wav_file:
                if wav_file.getframerate() != TARGET_SR or wav_file.getnchannels() != 1:
                    raise AssertionError(f"bad audio format: {wav_path}")


def main():
    if OUT_ROOT.exists():
        resolved = OUT_ROOT.resolve()
        workspace = Path.cwd().resolve()
        if workspace not in resolved.parents:
            raise RuntimeError(f"refusing to delete unexpected path: {resolved}")
        shutil.rmtree(OUT_ROOT)
    OUT_ROOT.mkdir(parents=True, exist_ok=True)

    candidates = load_candidates()
    for region, rows in candidates.items():
        print(f"{region} unselected candidates: {len(rows)}")

    noise_files_by_type = load_noise_files()
    all_rows = {}
    for region in ["sichuan", "yue", "henan", "dongbei", "shandong"]:
        rows = build_region(region, candidates[region], noise_files_by_type)
        all_rows[region] = rows
        print(f"{region}: {len(rows)} mixed files")

    validate_outputs(all_rows)
    print(f"done: {sum(len(rows) for rows in all_rows.values())} mixed files")
    print(OUT_ROOT.resolve())


if __name__ == "__main__":
    main()
