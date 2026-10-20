from pathlib import Path
import csv
import random
import shutil

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


SOURCE_ROOT = Path("downloaded_audio") / "kespeech"
OUT_ROOT = Path("downloaded_audio") / "integrated_dialects"
SAMPLE_SIZE_PER_REGION = 200
RANDOM_SEED = 20260625

REGION_LANGUAGE = {
    "dongbei": "东北话",
    "henan": "河南话",
    "shandong": "山东话",
}

FIELDNAMES = [
    "编号",
    "语言",
    "文本",
    "saved_path",
    "时长",
    "id",
    "speaker_id",
    "source_path",
    "source_filename",
    "source_transcription",
    "sample_index",
    "dialect",
    "source_wav_path",
]


def write_excel(path, rows, fieldnames=None):
    fieldnames = fieldnames or FIELDNAMES
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(fieldnames)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(field, "") for field in fieldnames])

    ws.freeze_panes = "A2"
    widths = {
        "A": 14,
        "B": 12,
        "C": 40,
        "D": 32,
        "E": 10,
        "F": 28,
        "G": 16,
        "H": 72,
        "I": 30,
        "J": 40,
        "K": 14,
        "L": 16,
        "M": 48,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def read_region_rows(region):
    metadata_path = SOURCE_ROOT / region / "metadata.csv"
    with metadata_path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def build_region(region, rows):
    rng = random.Random(f"{RANDOM_SEED}:{region}")
    candidates = [row for row in rows if row.get("saved_path")]
    if len(candidates) < SAMPLE_SIZE_PER_REGION:
        raise RuntimeError(f"{region} only has {len(candidates)} usable rows")

    selected = rng.sample(candidates, SAMPLE_SIZE_PER_REGION)
    region_out = OUT_ROOT / region
    wav_dir = region_out / "wav"
    wav_dir.mkdir(parents=True, exist_ok=True)

    output_rows = []
    for index, row in enumerate(selected, start=1):
        source_path = SOURCE_ROOT / row["saved_path"]
        file_id = row.get("id") or Path(row["saved_path"]).stem
        suffix = source_path.suffix or ".wav"
        dst_name = f"{region}_{index:04d}_{file_id}{suffix}"
        dst_path = wav_dir / dst_name
        shutil.copy2(source_path, dst_path)

        saved_path = dst_path.relative_to(region_out).as_posix()
        speaker_id = str(file_id).split("_")[0] if file_id else ""
        output_rows.append(
            {
                "编号": f"{region}_{index:04d}",
                "语言": REGION_LANGUAGE[region],
                "文本": row.get("text", ""),
                "saved_path": saved_path,
                "时长": row.get("duration_seconds", ""),
                "id": file_id,
                "speaker_id": speaker_id,
                "source_path": str(source_path),
                "source_filename": source_path.name,
                "source_transcription": row.get("text", ""),
                "sample_index": index,
                "dialect": row.get("dialect", ""),
                "source_wav_path": row.get("source_wav_path", ""),
            }
        )

    write_excel(region_out / f"{REGION_LANGUAGE[region]}_映射表.xlsx", output_rows)
    write_excel(region_out / f"{REGION_LANGUAGE[region]}_excel模版.xlsx", [])

    with (region_out / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(output_rows)

    return output_rows


def main():
    all_rows = []
    for region in REGION_LANGUAGE:
        rows = read_region_rows(region)
        output_rows = build_region(region, rows)
        all_rows.extend(
            {"region": region, "region_name": REGION_LANGUAGE[region], **row}
            for row in output_rows
        )
        print(f"{REGION_LANGUAGE[region]} ({region}): {len(output_rows)}")

    combined_fieldnames = ["region", "region_name", *FIELDNAMES]
    with (OUT_ROOT / "kespeech_三方言_映射表.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=combined_fieldnames)
        writer.writeheader()
        writer.writerows(all_rows)

    write_excel(OUT_ROOT / "kespeech_三方言_映射表.xlsx", all_rows, combined_fieldnames)
    print(f"done: {len(all_rows)} files")
    print(f"output: {OUT_ROOT.resolve()}")


if __name__ == "__main__":
    main()
